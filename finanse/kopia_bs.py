"""Pakiet 6: jednorazowa kopia z Budget Studio jako dane własne apki, kontrola sum co do grosza i porównanie
każdego kolejnego eksportu BS z planami apki (okres równoległy 1.11 do 31.12.2026).

Kopia (decyzja Mata 2026-10-19: cały 2026 z historią) to model planu jak w pakiecie 2, nie zrzut płatności:
  - arkusze z zakresami (rozszerzonymi tak, by mieściły wszystkie płatności),
  - wiersze z pełnymi ustawieniami wyprowadzonymi z płatności planowanych (reguła dnia wystawienia, dni na płatność,
    stawka, metoda, Księgowane, W cash flow); trzy wiersze umowy o pracę z eksportu BS („- dla wykonawcy”, „- PIT”,
    „- ZUS”) scalone w jeden wiersz „Umowa o pracę” z proporcjami z płatności,
  - komórki z kwotami; plany są pochodną komórki (generator z plan.py), a to, czym plan BS różni się od reguły wiersza
    (inna data, aktualność, komentarz), przechodzi jako ręczna zmiana pojedynczego planu (nadpisanie),
  - podział komórki wprost z eksportu: worek (kwoty tygodni, także zwinięty miniony miesiąc i bieżący podzielony na
    pozostałe tygodnie) i umowa, gdy proporcje wiersza dają inny grosz,
  - płatności rzeczywiste 1:1 z datami i cechami; komórka, której nie da się odtworzyć generatorem co do grosza,
    zostaje jako plany z eksportu (jawne, jak w imporcie).
Kopia trafia do pliku Finanse/Aplikacja/kopia_bs_RRRR-MM-DD.json (600), a do dziennika zmian jeden wpis „kopia_bs”
z nazwą pliku i jego sumą kontrolną. Od tej chwili kopia jest źródłem budżetu; zdarzenia „bud” sprzed kopii nie
obowiązują, późniejsze nakładają się na nią, a kolejne eksporty BS tylko się z nią porównuje.

Kontrola sum: arkusz × miesiąc × rodzaj × typ (plan, rzeczywista): liczba, netto, VAT, brutto w BS i w apce, plus suma
cash flow (brutto wg miesiąca dnia realnej płatności) i wyniku (netto wg miesiąca księgowania) per miesiąc; każda
różnica z przyczyną z porównania płatność po płatności (po ID z BS). Kopię da się przyjąć tylko przy 0 różnic.
"""
import copy
import csv
import hashlib
import io
import json
import re
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from . import budzet, konfig, plan

WERSJA = 1
ZRODLO = 'kopia BS'
_SKLADOWA = re.compile(r'(Umowa o pracę|Umowa zlecenie|Umowa o dzieło) - (dla wykonawcy|dla pracownika|PIT|ZUS)')
_PLIK = re.compile(r'kopia_bs_\d{4}-\d{2}-\d{2}(?:_\d{1,3})?\.json')
POLA_DAT = ('wystawienie', 'termin', 'realny', 'ksiegowanie')
POLA_CECH = ('metoda', 'ksiegowane', 'w_cashflow', 'komentarz', 'aktualna', 'kontrahent')
MAKS_PRZYCZYN = 500


def skladowa_bs(podstawa: str):
    """„Umowa o pracę - PIT” → ('Umowa o pracę', 'us'); inne podstawy → None."""
    m = _SKLADOWA.fullmatch((podstawa or '').strip())
    return (m[1], {'PIT': 'us', 'ZUS': 'zus'}.get(m[2], 'osoba')) if m else None


def aliasy_umow(wiersze: dict, platnosci: list[dict]) -> dict:
    """Wiersze składowych umowy z eksportu BS → wiersz „dla wykonawcy” tej samej umowy (składowe jednej wypłaty mają to
    samo ID płatności w BS). {id wiersza PIT albo ZUS: id wiersza osoby}."""
    rodzic = {}

    def korzen(x):
        while rodzic.get(x, x) != x:
            x = rodzic[x]
        return x
    po_bs = defaultdict(set)
    for p in platnosci:
        if skladowa_bs(wiersze[p['wiersz']]['podstawa']):
            po_bs[p['bs_id']].add(p['wiersz'])
    for ws in po_bs.values():
        ws = sorted(ws)
        for x in ws[1:]:
            a, b = korzen(ws[0]), korzen(x)
            if a != b:
                rodzic[max(a, b)] = min(a, b)
    grupy = defaultdict(list)
    for wid in {w for ws in po_bs.values() for w in ws}:
        grupy[korzen(wid)].append(wid)
    aliasy = {}
    for ws in grupy.values():
        osoby = sorted(w for w in ws if skladowa_bs(wiersze[w]['podstawa'])[1] == 'osoba')
        cel = osoby[0] if osoby else min(ws)
        aliasy.update({w: cel for w in ws if w != cel})
    return dict(sorted(aliasy.items()))


def eksporty() -> list[Path]:
    from .importy.wspolne import pliki
    return pliki(konfig.katalog_dane(), 'Budget Studio/budget_studio_*.csv')


def _arkusze_json() -> Path:
    return konfig.katalog_finanse() / 'Aplikacja' / 'arkusze_bs.json'


def _wiersze_csv(p: Path) -> dict[int, dict]:
    with open(p, encoding='utf-8-sig', newline='') as f:
        return {nr: x for nr, x in enumerate(csv.DictReader(f, delimiter='\t'), start=2)}


def _sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ---------- Ustawienia wiersza z płatności ----------

def _regula_z_dat(probki) -> str:
    """Reguła dnia wystawienia, która odtwarza najwięcej dat z próbek (miesiąc budżetowy, wystawienie, księgowanie).
    Remis: bez przesunięcia, potem „Ostatni”, potem wcześniejszy dzień."""
    kandydaci = set()
    for m, wyst, ks in probki:
        if not wyst:
            continue
        o = budzet._roznica_mies(m, wyst[:7])
        if abs(o) > 12:
            continue
        d = int(wyst[8:10])
        b = bool(ks) and ks[:7] == m and o != 0 and d <= 15
        dni = [d] + (['ostatni'] if d == plan.ostatni_dzien(wyst[:7]) else [])
        for dz in dni:
            kandydaci.add(plan.regula_tekst({'dzien': dz, 'przes': o, 'budzet': b and dz != 'ostatni'}))
    if not kandydaci:
        return 'Ostatni'

    def ocena(t):
        trafione = sum(1 for m, wyst, ks in probki if wyst and plan.wystawienie_z_reguly(m, t) == (wyst, f'{(ks or wyst)[:7]}-01'))
        r = plan.regula(t)
        return (trafione, r['przes'] == 0, r['dzien'] == 'ostatni', -(r['dzien'] if isinstance(r['dzien'], int) else 0))
    return max(sorted(kandydaci), key=ocena)


def _najczestszy(wartosci, domyslna):
    c = Counter(v for v in wartosci if v is not None)
    return c.most_common(1)[0][0] if c else domyslna


def _ustawienia(w: dict, plany: list[dict], wszystkie: list[dict]):
    """Pełne ustawienia wiersza z płatności: najpierw z planów (BS generuje je z ustawień wiersza), gdy brak: z rzeczywistych."""
    pod = w['podstawa']
    probki = plany or wszystkie
    osoby = [p for p in probki if p.get('skladowa', 'osoba') == 'osoba'] or probki
    if pod in plan.WOREK:
        w['dzien_wystawienia'] = '(auto)'
    else:
        w['dzien_wystawienia'] = _regula_z_dat([(p['miesiac'], p['wystawienie'], p['ksiegowanie']) for p in osoby])
    roznice = [(date.fromisoformat(p['termin']) - date.fromisoformat(p['wystawienie'])).days
               for p in osoby if p.get('wystawienie') and p.get('termin')]
    w['dni_na_platnosc'] = _najczestszy([d for d in roznice if 0 <= d <= 99], 0)
    if plan.ma_vat(pod):
        st = [budzet.stawka_proc(p['stawka']) for p in probki if p['netto']]
        w['stawka_vat'] = _najczestszy(st, None) if any(x is not None for x in st) else _najczestszy(
            [budzet.stawka_proc(p['stawka']) for p in wszystkie], None)
    else:
        w['stawka_vat'] = None
    w['metoda'] = _najczestszy([p.get('metoda') or '' for p in probki], '')
    w['ksiegowane'] = True if plan.jest_umowa(pod) else bool(_najczestszy([bool(p.get('ksiegowane', True)) for p in probki], True))
    w['w_cashflow'] = bool(_najczestszy([bool(p.get('w_cashflow', True)) for p in probki], True))
    w['split'] = False
    w['cykl'] = 'brak'
    w['zrodlo'] = ZRODLO
    w.pop('plan', None)


# ---------- Budowa kopii ----------

def _st(w: dict, platnosci: list[dict], tydzien: int) -> dict:
    """Stan jak w budzet._stan dla jednego wiersza (generator z pakietu 2 na tych samych zasadach co w apce)."""
    return {'arkusz': {}, 'wiersz': {w['id']: w}, 'platnosc': {p['id']: dict(p) for p in platnosci}, 'komorka': {},
            'nadpisane': {}, 'parametry': {'tydzien': tydzien}, 'podzialy': {}}


def _generuj(w, m, rzeczywiste, kwota, podzial, nadpisane, tydzien) -> list[dict]:
    st = _st(w, rzeczywiste, tydzien)
    st['komorka'][(w['id'], m)] = kwota
    st['nadpisane'] = nadpisane
    if podzial is not None:
        st['podzialy'][f"{w['id']}:{m}"] = podzial
    budzet._odswiez(st, w['id'], {m})
    gen = [p for p in st['platnosc'].values() if p.get('generowany')]
    kolej = {'osoba': 0, 'us': 1, 'zus': 2}

    def nr(pid):
        m = re.search(r':t(\d+)$', pid)
        return int(m[1]) if m else 0
    return sorted(gen, key=lambda p: (kolej.get(p.get('skladowa'), 9), nr(p['id']), p['id']))


def _cechy_bs(p: dict, x: dict) -> dict:
    """Wartości pól planu z eksportu BS w zapisie generatora (kontrahent z wiersza eksportu)."""
    return {'netto': p['netto'], 'vat': p['vat'], 'brutto': p['brutto'], 'stawka': budzet.stawka_proc(p['stawka']),
            'wystawienie': p['wystawienie'], 'termin': p['termin'], 'realny': p['realny'], 'ksiegowanie': p['ksiegowanie'],
            'metoda': p.get('metoda') or '', 'ksiegowane': bool(p.get('ksiegowane', True)), 'w_cashflow': bool(p.get('w_cashflow', True)),
            'komentarz': p.get('komentarz') or '', 'aktualna': bool(p.get('aktualna', True)), 'kontrahent': (x.get('Kontrahent') or '').strip()}


def _cechy_gen(g: dict, w: dict) -> dict:
    return {'netto': g['netto'], 'vat': g['vat'], 'brutto': g['brutto'], 'stawka': g.get('stawka'),
            'wystawienie': g.get('wystawienie'), 'termin': g.get('termin'), 'realny': g.get('realny'), 'ksiegowanie': g.get('ksiegowanie'),
            'metoda': g.get('metoda') or '', 'ksiegowane': bool(g.get('ksiegowane', True)), 'w_cashflow': bool(g.get('w_cashflow', True)),
            'komentarz': g.get('komentarz') or '', 'aktualna': bool(g.get('aktualna', True)),
            'kontrahent': g.get('kontrahent') or w.get('kontrahent', '')}


def _roznice(g: dict, b: dict, pod: str) -> list[str]:
    pola = ['netto', 'vat', 'brutto', *POLA_DAT, *POLA_CECH] + (['stawka'] if plan.ma_vat(pod) else [])
    return [k for k in pola if g.get(k) != b.get(k)]


def _komorka(w, m, plany, rzeczywiste, csv_po_nr, tydzien):
    """Komórka z planami BS jako (kwota, podział, nadpisania) albo None, gdy generator nie odtworzy jej co do grosza.
    Zwraca też opis zgodności z regułą (worek, umowa) do raportu."""
    pod = w['podstawa']
    worek, umowa = pod in plan.WOREK, plan.jest_umowa(pod)
    if not worek and rzeczywiste:
        return None, 'komórka z planem i płatnością rzeczywistą (nie worek)', None
    if umowa:
        po = {p.get('skladowa'): p for p in plany}
        if len(po) != len(plany) or set(po) != set(plan.UMOWY[pod]):
            return None, 'niepełna umowa (brak którejś składowej)', None
        bs = [po[k] for k in plan.UMOWY[pod]]
    elif worek:
        bs = sorted(plany, key=lambda p: (p['wystawienie'] or '', p['id']))
    else:
        if len(plany) != 1:
            return None, f'{len(plany)} plany w komórce (nie worek)', None
        bs = plany
    suma_planow = sum(p['netto'] for p in bs)
    kwota = suma_planow + (sum(p['netto'] for p in rzeczywiste) if worek else 0)
    zgodnosc = None
    if worek:
        podzial = [[p['wystawienie'], p['netto']] for p in bs]
        regula = plan.podziel_worek(suma_planow, m, tydzien) if suma_planow else [(f'{m}-01', 0)]
        zgodnosc = ('worek', [tuple(x) for x in podzial] == [tuple(x) for x in regula], _rodzaj_worka(podzial, m))
    elif umowa:
        bs_kw = {k: po[k]['netto'] for k in plan.UMOWY[pod]}
        rg = plan.podziel_umowe(kwota, w.get('proporcje'), plan.UMOWY[pod]) if kwota else {'osoba': 0}
        zgodne = rg == bs_kw
        podzial = None if zgodne else bs_kw
        zgodnosc = ('umowa', zgodne, max(abs(rg.get(k, 0) - v) for k, v in bs_kw.items()))
    else:
        podzial = None
    cechy = [_cechy_bs(p, csv_po_nr.get(p['wiersz_pliku'], {})) for p in bs]
    nadpisane = {}
    for pole in (*POLA_DAT, None):   # daty po kolei (zmiana wystawienia przesuwa termin), potem reszta cech naraz
        gen = _generuj(w, m, rzeczywiste, kwota, podzial, nadpisane, tydzien)
        if len(gen) != len(bs):
            return None, 'inna liczba planów niż w BS', zgodnosc
        for g, b in zip(gen, cechy):
            gc = _cechy_gen(g, w)
            for k in ([pole] if pole else [k for k in POLA_CECH if k != 'kontrahent'] + ['kontrahent']):
                if gc[k] != b[k]:
                    if k == 'kontrahent' and not b[k]:
                        continue
                    nadpisane.setdefault(g['id'], {})[k] = b[k]
        if pole is None:
            for g, b in zip(gen, cechy):
                if plan.ma_vat(pod) and g.get('stawka') != b['stawka']:
                    nadpisane.setdefault(g['id'], {})['stawka'] = b['stawka']
    gen = _generuj(w, m, rzeczywiste, kwota, podzial, nadpisane, tydzien)
    for g, b in zip(gen, cechy):
        r = _roznice(_cechy_gen(g, w), b, pod)
        if r:
            return None, 'generator nie odtwarza: ' + ', '.join(sorted(set(r))), zgodnosc
    pochodzenie = {g['id']: {'bs_pid': p['id'], 'bs_id': p['bs_id'], 'etap_bs': p.get('etap_bs', ''), 'tytul': p.get('tytul', ''),
                             'zmodyfikowano': p.get('zmodyfikowano', ''), 'wiersz_pliku': p.get('wiersz_pliku')} for g, p in zip(gen, bs)}
    return {'kwota': kwota, 'podzial': podzial if (worek or podzial) else None, 'nadpisane': nadpisane, 'pochodzenie': pochodzenie}, None, zgodnosc


def _rodzaj_worka(podzial, m) -> str:
    if len(podzial) == 1 and podzial[0][0] == f'{m}-{plan.ostatni_dzien(m):02d}':
        return 'miniony miesiąc zwinięty do jednego planu z ostatniego dnia'
    if podzial and podzial[0][0] > f'{m}-01':
        return 'bieżący miesiąc: tylko pozostałe tygodnie'
    return 'inny podział'


def zbuduj_kopie(csv_path: Path, arkusze_json: Path | None = None, mapa: dict | None = None, tydzien: int | None = None) -> dict:
    """Kopia modelu z eksportu BS (bez zapisu). mapa: mapa stałych ID (uzupełniana w miejscu; do podglądu podaj kopię)."""
    csv_path = Path(csv_path)
    tydzien = budzet.parametry_planu()['tydzien'] if tydzien is None else tydzien
    imp = budzet.generuj_import(csv_path, arkusze_json if arkusze_json is not None else _arkusze_json(),
                                mapa if mapa is not None else budzet.nowa_mapa())
    csv_po_nr = _wiersze_csv(csv_path)
    wiersze = {w['id']: dict(w) for w in imp['wiersze']}
    platnosci = [dict(p) for p in imp['platnosci']]

    # 1. Umowy: trzy wiersze składowych z eksportu (ta sama płatność BS) w jeden wiersz umowy, składowa i spinacz w płatności.
    aliasy = aliasy_umow(wiersze, platnosci)
    po_bs = defaultdict(list)
    for p in platnosci:
        s = skladowa_bs(wiersze[p['wiersz']]['podstawa'])
        if s:
            p['skladowa'] = s[1]
            po_bs[p['bs_id']].append(p)
    for wid in aliasy:
        wiersze.pop(wid)
    for wid, w in wiersze.items():
        s = skladowa_bs(w['podstawa'])
        if s:
            w['podstawa'], w['rodzaj'] = s[0], -1
    for p in platnosci:
        p['wiersz'] = aliasy.get(p['wiersz'], p['wiersz'])
        p['zrodlo'] = ZRODLO
    for ps in po_bs.values():
        osoba = next((p for p in ps if p['skladowa'] == 'osoba'), None)
        for p in ps:
            if osoba and p is not osoba and p['typ'] != 'plan':
                p['spinacz'] = osoba['id']

    # 2. Komórki (wiersz × miesiąc) i ustawienia wierszy
    komorki = defaultdict(list)
    for p in platnosci:
        komorki[(p['wiersz'], p['miesiac'])].append(p)
    po_wierszu = defaultdict(list)
    for p in platnosci:
        po_wierszu[p['wiersz']].append(p)
    for wid, w in wiersze.items():
        ps = po_wierszu.get(wid, [])
        _ustawienia(w, [p for p in ps if p['typ'] == 'plan'], ps)
        if plan.jest_umowa(w['podstawa']):
            w['proporcje'] = _proporcje(w, ps)

    # 3. Każda komórka z planami: kwota + podział + nadpisania albo plany jawne
    model_komorki, nadpisane, podzialy, pochodzenie, jawne = [], {}, {}, {}, []
    wer = {'worek': Counter(), 'worek_inne': Counter(), 'umowa': Counter(), 'umowa_maks_gr': 0, 'jawne': Counter(), 'nadpisania': Counter()}
    for (wid, m), ps in sorted(komorki.items()):
        w = wiersze[wid]
        plany = [p for p in ps if p['typ'] == 'plan']
        rz = [p for p in ps if p['typ'] != 'plan']
        jawne += rz
        if not plany:
            continue
        wynik, powod, zg = _komorka(w, m, plany, rz, csv_po_nr, tydzien)
        if zg:
            if zg[0] == 'worek':
                wer['worek']['zgodne' if zg[1] else 'inne'] += 1
                if not zg[1]:
                    wer['worek_inne'][zg[2]] += 1
            else:
                wer['umowa']['zgodne' if zg[1] else 'inne'] += 1
                wer['umowa_maks_gr'] = max(wer['umowa_maks_gr'], zg[2])
        if not wynik:
            wer['jawne'][powod] += len(plany)
            jawne += plany
            continue
        model_komorki.append({'wiersz': wid, 'miesiac': m, 'kwota': wynik['kwota']})
        if wynik['podzial'] is not None:
            podzialy[f'{wid}:{m}'] = wynik['podzial']
        for pid, n in wynik['nadpisane'].items():
            if n:
                nadpisane[pid] = n
                wer['nadpisania'].update(n.keys())
        pochodzenie.update(wynik['pochodzenie'])

    # 4. Arkusze: zakres obejmuje wszystkie płatności (cały 2026 z historią i plany na 2027)
    miesiace_ark = defaultdict(set)
    for p in platnosci:
        miesiace_ark[wiersze[p['wiersz']]['arkusz']].add(p['miesiac'])
    arkusze, rozszerzone = [], 0
    for a in imp['arkusze']:
        a = {**a, 'zrodlo': ZRODLO}
        ms = miesiace_ark.get(a['id'])
        if ms and (min(ms) < a['od'] or max(ms) > a['do']):
            a['od'], a['do'] = min(a['od'], min(ms)), max(a['do'], max(ms))
            rozszerzone += 1
        arkusze.append(a)
    wer['arkusze_rozszerzone'] = rozszerzone
    wer['tydzien'] = tydzien
    return {'wersja': WERSJA, 'utworzono': date.today().isoformat(),
            'zrodlo': {**imp['zrodlo'], 'eksport': imp['zrodlo']['plik']},
            'arkusze': arkusze, 'wiersze': sorted(wiersze.values(), key=lambda w: w['id']),
            'komorki': model_komorki, 'nadpisane': dict(sorted(nadpisane.items())), 'podzialy': dict(sorted(podzialy.items())),
            'platnosci': sorted(jawne, key=lambda p: p['id']), 'pochodzenie': dict(sorted(pochodzenie.items())),
            'wiersze_bs': dict(sorted(aliasy.items())),
            'weryfikacja': {k: (dict(v) if isinstance(v, Counter) else v) for k, v in wer.items()}}


def _proporcje(w, ps) -> dict | None:
    """Kwoty referencyjne umowy: składowe z najpóźniejszej komórki z planami (gdy brak: z rzeczywistymi), w groszach."""
    po = defaultdict(dict)
    for p in ps:
        po[(p['typ'] == 'plan', p['miesiac'])][p.get('skladowa', 'osoba')] = p['netto']
    skladowe = plan.UMOWY[w['podstawa']]
    for (_, m), x in sorted(po.items(), reverse=True):
        if set(x) == set(skladowe) and sum(x.values()):
            return {k: x[k] for k in skladowe}
    return None


# ---------- Model z kopii ----------

def rozwin(kopia: dict) -> dict:
    """Model budżetu z kopii (format jak z importu): plany komórek liczone generatorem z kwot, ustawień wierszy,
    podziałów i nadpisań, plus płatności jawne. Pochodzenie z BS (ID, etap w BS, tytuł) dopisane do planów."""
    model = {'zrodlo': {**kopia['zrodlo'], 'plik': kopia.get('plik', kopia['zrodlo'].get('plik'))},
             'arkusze': copy.deepcopy(kopia['arkusze']), 'wiersze': copy.deepcopy(kopia['wiersze']),
             'platnosci': copy.deepcopy(kopia['platnosci']), 'komorki': copy.deepcopy(kopia['komorki']),
             'nadpisane': copy.deepcopy(kopia['nadpisane']), 'podzialy': copy.deepcopy(kopia.get('podzialy', {})), 'stare_id': {}}
    st = budzet._stan(model)
    for wid in sorted({w for (w, _) in st['komorka']}):
        if wid in st['wiersz']:
            budzet._odswiez(st, wid)
    poch = kopia.get('pochodzenie', {})
    for pid, p in st['platnosc'].items():
        x = poch.get(pid)
        if x and p.get('generowany'):
            p['bs_id'], p['etap_bs'], p['wiersz_pliku'] = x['bs_id'], x.get('etap_bs', p.get('etap_bs')), x.get('wiersz_pliku')
            if not p.get('tytul'):
                p['tytul'] = x.get('tytul', '')
    model['platnosci'] = sorted(st['platnosc'].values(), key=lambda p: p['id'])
    return model


def aktywna(wpisy: list[dict]) -> dict | None:
    """Ostatni nie cofnięty wpis „kopia_bs” (kopię da się powtórzyć albo cofnąć jak każdą decyzję)."""
    cofniete = {w['dane'].get('wpis') for w in wpisy if w.get('typ') == 'cofnij'}
    k = [w for w in wpisy if w.get('typ') == 'kopia_bs' and w.get('id') not in cofniete]
    return k[-1] if k else None


def po_kopii(wpisy: list[dict], wpis: dict) -> list[dict]:
    """Wpisy nakładane na kopię: wszystkie poza zdarzeniami budżetu sprzed kopii."""
    granica = (wpis['czas'], wpis['id'])
    return [w for w in wpisy if w.get('typ') != 'bud' or (w.get('czas', ''), w.get('id', '')) > granica]


def plik_kopii(nazwa: str) -> Path:
    if not _PLIK.fullmatch(nazwa or ''):
        raise ValueError('zła nazwa pliku kopii')
    return konfig.katalog_finanse() / 'Aplikacja' / nazwa


def wczytaj(wpis: dict) -> dict:
    """Plik kopii wskazany wpisem; ValueError, gdy go nie ma (np. iCloud jeszcze nie pobrał) albo ma inną treść."""
    p = plik_kopii(wpis['dane']['plik'])
    if not p.exists():
        raise ValueError(f'brak pliku kopii {p.name} (może iCloud jeszcze go nie pobrał); budżet pokazuje import z eksportu BS')
    tresc = p.read_bytes()
    if hashlib.sha256(tresc).hexdigest() != wpis['dane']['sha256']:
        raise ValueError(f'plik kopii {p.name} ma inną treść niż przy przyjęciu; budżet pokazuje import z eksportu BS')
    kopia = json.loads(tresc)
    kopia['plik'] = p.name
    return kopia


# ---------- Kontrola sum ----------

def _rekord(p: dict, w: dict, bs_pid) -> dict:
    return {'id': p['id'], 'bs_pid': bs_pid, 'arkusz': w['arkusz'], 'wiersz': w['id'], 'miesiac': p['miesiac'], 'rodzaj': w['rodzaj'],
            'typ': 'plan' if p['typ'] == 'plan' else 'rzeczywista', 'netto': int(p['netto']), 'vat': int(p.get('vat') or 0),
            'brutto': int(p.get('brutto') if p.get('brutto') is not None else p['netto']),
            'cf': (p.get('realny') or p.get('termin') or p['miesiac'])[:7], 'wynik': (p.get('ksiegowanie') or p['miesiac'])[:7],
            'wystawienie': p.get('wystawienie') or '', 'realny': p.get('realny') or p.get('termin') or '',
            'w_cashflow': bool(p.get('w_cashflow', True)), 'ksiegowane': bool(p.get('ksiegowane', True)), 'aktualna': bool(p.get('aktualna', True))}


def _podatkowy_bs(x: dict) -> bool:
    """Wiersz eksportu BS z podstawą podatkową (VAT, podatek dochodowy): w apce zastępuje go arkusz systemowy (pakiet 8)."""
    return budzet._podstawa(x.get('Podstawa płatności')) in plan.PODATKI


def rekordy_bs(imp: dict, csv_po_nr: dict[int, dict], bez_podatkow: bool = False) -> list[dict]:
    """Strona BS kontroli: kwoty, daty i przełączniki wprost z wierszy pliku eksportu (nie z przetworzonego importu);
    z importu tylko stałe ID płatności, wiersza i arkusza. bez_podatkow: bez wierszy VAT i podatku dochodowego (porównanie
    z bazą, w której te wiersze zastępuje arkusz systemowy Opłaty urzędowe; zgodność danin pokazuje ekran Podatki)."""
    from .format import grosze
    w = {x['id']: x for x in imp['wiersze']}
    ark = {(a['folder'], a['kategoria'], a['nazwa']): a['id'] for a in imp['arkusze']}
    wynik = []
    for p in imp['platnosci']:
        x = csv_po_nr[p['wiersz_pliku']]
        if bez_podatkow and _podatkowy_bs(x):
            continue
        aid = ark.get((x['Folder'].strip(), x['Kategoria'].strip(), x['Arkusz'].strip()), w[p['wiersz']]['arkusz'])
        dok = bool((x.get('Numer wewnętrzny') or '').strip() or (x.get('Numer dokumentu') or '').strip())
        surowa = {'id': p['id'], 'miesiac': (budzet._data(x['Miesiąc budżetowy']) or '')[:7],
                  'typ': 'rzeczywista' if dok or x['Etap płatności'] in ('Opłacona', 'Niewymagana') else 'plan',
                  'netto': abs(grosze(x['Netto (wartość)'])), 'vat': abs(grosze(x['VAT'] or '0')), 'brutto': abs(grosze(x['Brutto (wartość)'])),
                  'realny': budzet._data(x['Dzień płatności - realny']), 'termin': budzet._data(x['Dzień płatności']),
                  'wystawienie': budzet._data(x['Dzień wystawienia']),
                  'ksiegowanie': budzet._data(x['Miesiąc księgowania']), 'w_cashflow': x['W cash flow'] == 'Tak',
                  'ksiegowane': x['Księgowane'] == 'Tak', 'aktualna': x['Status aktualności'] == 'Aktualne'}
        wynik.append(_rekord(surowa, {**w[p['wiersz']], 'arkusz': aid, 'rodzaj': int(x['Rodzaj'])}, p['id']))
    return wynik


def rekordy_modelu(model: dict, pochodzenie: dict | None = None) -> list[dict]:
    """Płatności modelu apki; bs_pid: ID płatności BS (jawne: własne ID, plany z kopii: z pochodzenia, dopóki ich nie
    przeliczono zmianą w apce)."""
    pochodzenie = pochodzenie or {}
    w = {x['id']: x for x in model['wiersze']}
    wynik = []
    for p in model['platnosci']:
        if p['wiersz'] not in w:
            continue
        bs_pid = p['id'] if p['id'].startswith('bs:') else (pochodzenie.get(p['id'], {}).get('bs_pid') if p.get('bs_id') else None)
        wynik.append(_rekord(p, w[p['wiersz']], bs_pid))
    return wynik


def rekordy_bazy(con, pochodzenie: dict | None = None) -> list[dict]:
    """To samo z bazy (to, co apka pokazuje po odbudowie), bez arkusza systemowego Opłaty urzędowe (pakiet 8: liczy go apka,
    nie ma odpowiednika wiersz w wiersz w BS)."""
    pochodzenie = pochodzenie or {}
    wynik = []
    for r in con.execute("SELECT p.*, w.rodzaj AS w_rodzaj, w.arkusz AS w_arkusz FROM bud_platnosc p JOIN bud_wiersz w ON w.id = p.wiersz"
                         " WHERE w.zrodlo != 'system' AND w.id NOT LIKE 'sys:%'"):
        p = dict(r)
        bs_pid = p['id'] if p['id'].startswith('bs:') else (pochodzenie.get(p['id'], {}).get('bs_pid') if p.get('bs_id') else None)
        wynik.append(_rekord(p, {'id': p['wiersz'], 'arkusz': p['w_arkusz'], 'rodzaj': p['w_rodzaj']}, bs_pid))
    return wynik


def _suma(xs) -> dict:
    return {'liczba': len(xs), 'netto': sum(x['netto'] for x in xs), 'vat': sum(x['vat'] for x in xs), 'brutto': sum(x['brutto'] for x in xs)}


PRZYCZYNY = {'arkusz': 'inny arkusz', 'miesiac': 'inny miesiąc budżetowy', 'rodzaj': 'inny rodzaj (wpływ albo obciążenie)',
             'typ': 'inny typ (plan albo rzeczywista)', 'netto': 'inne netto', 'vat': 'inny VAT', 'brutto': 'inne brutto',
             'cf': 'inny miesiąc dnia realnej płatności (cash flow)', 'wynik': 'inny miesiąc księgowania (wynik)',
             'wystawienie': 'inny dzień wystawienia', 'realny': 'inny dzień realnej płatności',
             'w_cashflow': 'inne „W cash flow”', 'ksiegowane': 'inne „Księgowane”', 'aktualna': 'inna aktualność'}


def kontrola(bs: list[dict], apka: list[dict], arkusze: dict, liczba_csv: int | None = None, z_aktualnoscia: bool = True) -> dict:
    """Raport kontroli sum. bs, apka: rekordy płatności (_rekord); arkusze: {id: {folder, kategoria, nazwa}}.
    z_aktualnoscia=True (kopia): aktualność porównywana płatność po płatności, cash flow i wynik z płatności aktualnych.
    False (porównanie po kopii): BS sam przestawia plan na nieaktualny, gdy minie jego data, więc aktualność nie jest
    różnicą w planach; cash flow i wynik ze wszystkich płatności."""
    grupy = defaultdict(lambda: ([], []))
    for i, strona in enumerate((bs, apka)):
        for x in strona:
            grupy[(x['arkusz'], x['miesiac'], x['rodzaj'], x['typ'])][i].append(x)
    # płatność po płatności (po ID z BS): przyczyny różnic
    po_bs = {x['bs_pid']: x for x in bs}
    po_apce = defaultdict(list)
    for x in apka:
        po_apce[x['bs_pid']].append(x)
    przyczyny = []   # (rekord bs albo None, rekord apki albo None, lista przyczyn)
    for pid, b in po_bs.items():
        a_lista = po_apce.get(pid, [])
        if not a_lista:
            przyczyny.append((b, None, ['brak w apce']))
            continue
        for a in a_lista[1:]:
            przyczyny.append((None, a, ['ta sama płatność BS dwa razy w apce']))
        r = [PRZYCZYNY[k] for k in PRZYCZYNY if b[k] != a_lista[0][k] and (z_aktualnoscia or k != 'aktualna')]
        if r:
            przyczyny.append((b, a_lista[0], r))
    for x in apka:
        if not x['bs_pid'] or x['bs_pid'] not in po_bs:
            przyczyny.append((None, x, ['jest tylko w apce (nowa albo zmieniona w apce, albo usunięta w BS)']))

    def opis(klucz_f):
        out = []
        for b, a, r in przyczyny:
            if (b and klucz_f(b)) or (a and klucz_f(a)):
                out.append(' i '.join(r) + f" ({(b or a)['bs_pid'] or a['id']})")
        return out[:20]

    komorki = []
    for k in sorted(grupy, key=lambda k: (k[0], k[1], -k[2], k[3])):
        sb, sa = _suma(grupy[k][0]), _suma(grupy[k][1])
        if sb != sa:
            komorki.append({'arkusz': k[0], 'miesiac': k[1], 'rodzaj': k[2], 'typ': k[3], 'bs': sb, 'apka': sa,
                            'przyczyny': opis(lambda x, k=k: (x['arkusz'], x['miesiac'], x['rodzaj'], x['typ']) == k)})
    znak = lambda xs, pole: sum(x[pole] * x['rodzaj'] for x in xs)
    po_ark = defaultdict(lambda: ([], []))
    for i, strona in enumerate((bs, apka)):
        for x in strona:
            po_ark[x['arkusz']][i].append(x)
    lista_ark = []
    for aid in sorted(po_ark, key=lambda a: tuple((arkusze.get(a) or {}).get(k, '') for k in ('folder', 'kategoria', 'nazwa'))):
        b, a = po_ark[aid]
        sb = {'liczba': len(b), 'netto': znak(b, 'netto'), 'vat': znak(b, 'vat'), 'brutto': znak(b, 'brutto')}
        sa = {'liczba': len(a), 'netto': znak(a, 'netto'), 'vat': znak(a, 'vat'), 'brutto': znak(a, 'brutto')}
        lista_ark.append({'arkusz': aid, **{k: (arkusze.get(aid) or {}).get(k, '') for k in ('folder', 'kategoria', 'nazwa')},
                          'bs': sb, 'apka': sa, 'zgodny': sb == sa})
    razem = {s: {'liczba': len(xs), 'netto': znak(xs, 'netto'), 'vat': znak(xs, 'vat'), 'brutto': znak(xs, 'brutto')}
             for s, xs in (('bs', bs), ('apka', apka))}
    miesiace = defaultdict(lambda: {'cf_bs': 0, 'cf_apka': 0, 'wynik_bs': 0, 'wynik_apka': 0})
    for s, xs in (('bs', bs), ('apka', apka)):
        for x in xs:
            if z_aktualnoscia and not x['aktualna']:
                continue
            if x['w_cashflow']:
                miesiace[x['cf']][f'cf_{s}'] += x['brutto'] * x['rodzaj']
            if x['ksiegowane']:
                miesiace[x['wynik']][f'wynik_{s}'] += x['netto'] * x['rodzaj']
    lista_mies = []
    for m in sorted(miesiace):
        v = miesiace[m]
        zg = v['cf_bs'] == v['cf_apka'] and v['wynik_bs'] == v['wynik_apka']
        r = {'miesiac': m, **v, 'zgodny': zg}
        if not zg:
            r['przyczyny'] = opis(lambda x, m=m: (x['cf'] == m or x['wynik'] == m) and (x['aktualna'] or not z_aktualnoscia))
        lista_mies.append(r)
    roznice_mies = sum(1 for x in lista_mies if not x['zgodny'])
    zla_liczba = liczba_csv is not None and liczba_csv != len(bs)
    platnosci = [{'bs_pid': (b or {}).get('bs_pid'), 'apka_id': (a or {}).get('id'), 'arkusz': (b or a)['arkusz'], 'miesiac': (b or a)['miesiac'],
                  'typ': (b or a)['typ'], 'przyczyny': r,
                  'bs': {k: b[k] for k in ('netto', 'vat', 'brutto')} if b else None,
                  'apka': {k: a[k] for k in ('netto', 'vat', 'brutto')} if a else None} for b, a, r in przyczyny[:MAKS_PRZYCZYN]]
    return {'roznic': len(komorki) + roznice_mies + len(przyczyny) + int(zla_liczba), 'roznic_komorek': len(komorki),
            'roznic_miesiecy': roznice_mies, 'roznic_platnosci': len(przyczyny), 'zla_liczba': zla_liczba,
            'liczba_csv': liczba_csv, 'liczba_bs': len(bs), 'liczba_apka': len(apka), 'komorki': komorki, 'arkusze': lista_ark,
            'razem': razem, 'miesiace': lista_mies, 'platnosci': platnosci, 'platnosci_razem': len(przyczyny), 'z_aktualnoscia': z_aktualnoscia}


def liczba_w_csv(p: Path, bez_podatkow: bool = False) -> int:
    """Liczba wierszy eksportu z rodzajem 1 albo -1, liczona wprost z pliku (niezależnie od importu)."""
    return sum(1 for x in _wiersze_csv(p).values() if x.get('Rodzaj') in ('1', '-1') and not (bez_podatkow and _podatkowy_bs(x)))


def _nazwy_arkuszy(*zrodla) -> dict:
    wynik = {}
    for arkusze in zrodla:
        for a in arkusze:
            wynik.setdefault(a['id'], {k: a.get(k, '') for k in ('folder', 'kategoria', 'nazwa')})
    return wynik


def podglad(eksport: Path | None = None, wpisy: list[dict] | None = None) -> dict:
    """Kopia z najnowszego eksportu BS i jej kontrola sum, bez żadnego zapisu (mapa stałych ID w kopii roboczej).
    zdarzen_przed: ile obowiązujących dziś zmian planów (wpisy „bud”) przestanie obowiązywać po przyjęciu kopii."""
    lista = [eksport] if eksport else eksporty()
    if not lista:
        raise ValueError('brak eksportu Budget Studio w Raporty/Dane/Budget Studio')
    csv_path = lista[0]
    mapa = copy.deepcopy(budzet.wczytaj_mape())
    kopia = zbuduj_kopie(csv_path, _arkusze_json(), mapa)
    imp = budzet.generuj_import(csv_path, _arkusze_json(), copy.deepcopy(mapa))
    model = rozwin(kopia)
    budzet.zastosuj_zdarzenia(model, [])
    raport = kontrola(rekordy_bs(imp, _wiersze_csv(csv_path)), rekordy_modelu(model, kopia['pochodzenie']), _nazwy_arkuszy(model['arkusze'], imp['arkusze']),
                      liczba_w_csv(csv_path))
    raport.update(eksport=csv_path.name, eksport_sha256=_sha(csv_path), weryfikacja=kopia['weryfikacja'], podsumowanie=_podsumowanie(kopia),
                  zdarzen_przed=_obowiazujace_zmiany(wpisy or []))
    return raport


def _obowiazujace_zmiany(wpisy: list[dict]) -> int:
    k = aktywna(wpisy)
    cofniete = {w['dane'].get('wpis') for w in wpisy if w.get('typ') == 'cofnij'}
    return sum(1 for w in (po_kopii(wpisy, k) if k else wpisy) if w.get('typ') == 'bud' and w.get('id') not in cofniete)


def _podsumowanie(kopia: dict) -> dict:
    return {'arkuszy': len(kopia['arkusze']), 'wierszy': len(kopia['wiersze']), 'komorek': len(kopia['komorki']),
            'podzialow': len(kopia['podzialy']), 'nadpisanych_planow': len(kopia['nadpisane']),
            'planow_z_komorek': len(kopia['pochodzenie']), 'platnosci_jawnych': len(kopia['platnosci']),
            'scalonych_umow': len(set(kopia['wiersze_bs'].values()))}


def przyjmij(eksport_sha256: str, wpisy: list[dict]) -> dict:
    """Zapis kopii: plik kopii_bs (600), raport kontroli w Excelu (600) i wpis „kopia_bs” w pliku zmian tego Maca.
    Tylko gdy eksport jest ten sam co w podglądzie i kontrola ma 0 różnic. Zwraca wpis."""
    from . import zmiany
    lista = eksporty()
    if not lista:
        raise ValueError('brak eksportu Budget Studio')
    csv_path = lista[0]
    if _sha(csv_path) != eksport_sha256:
        raise ValueError('eksport BS zmienił się od podglądu; odśwież podgląd i sprawdź raport jeszcze raz')
    mapa = copy.deepcopy(budzet.wczytaj_mape())
    kopia = zbuduj_kopie(csv_path, _arkusze_json(), mapa)
    imp = budzet.generuj_import(csv_path, _arkusze_json(), copy.deepcopy(mapa))
    model = rozwin(kopia)
    budzet.zastosuj_zdarzenia(model, [])
    raport = kontrola(rekordy_bs(imp, _wiersze_csv(csv_path)), rekordy_modelu(model, kopia['pochodzenie']), _nazwy_arkuszy(model['arkusze'], imp['arkusze']),
                      liczba_w_csv(csv_path))
    if raport['roznic']:
        raise ValueError(f"kontrola sum ma {raport['roznic']} różnic; kopii nie przyjęto")
    katalog = konfig.katalog_prywatny(konfig.katalog_finanse() / 'Aplikacja')
    nazwa, i = f'kopia_bs_{date.today().isoformat()}.json', 1
    while (katalog / nazwa).exists():
        i += 1
        nazwa = f'kopia_bs_{date.today().isoformat()}_{i}.json'
    tresc = json.dumps(kopia, ensure_ascii=False, indent=0, sort_keys=True)
    budzet._zapisz_jesli_inny(katalog / nazwa, tresc)
    raport.update(eksport=csv_path.name, weryfikacja=kopia['weryfikacja'], podsumowanie=_podsumowanie(kopia))
    _zapisz_xlsx(katalog / nazwa.replace('.json', '_kontrola.xlsx'), raport)
    return zmiany.zapisz('kopia_bs', {'plik': nazwa, 'sha256': hashlib.sha256(tresc.encode('utf-8')).hexdigest(),
                                      'eksport': csv_path.name, 'eksport_sha256': eksport_sha256, 'roznic': 0})


def zgodnosc(con, wpisy: list[dict]) -> dict:
    """Ekran „Zgodność z BS”: najnowszy eksport BS kontra plany apki (to, co jest w bazie), ten sam raport co przy kopii.
    Od 1.11 pokazuje, gdzie plan w BS różni się od planu w apce."""
    lista = eksporty()
    k = aktywna(wpisy)
    stan = {'kopia': None, 'import_bs': budzet.import_bs_wlaczony(wpisy)}
    pochodzenie = {}
    if k:
        stan['kopia'] = {'wpis': k['id'], 'czas': k['czas'], 'mac': k.get('mac', ''), **k['dane']}
        try:
            pochodzenie = wczytaj(k).get('pochodzenie', {})
        except ValueError as e:
            stan['kopia']['blad'] = str(e)
    if not lista or not stan['import_bs']:
        return {**stan, 'eksport': None}
    csv_path = lista[0]
    imp = budzet.generuj_import(csv_path, _arkusze_json(), copy.deepcopy(budzet.wczytaj_mape()))
    arkusze = [dict(r) for r in con.execute("SELECT id, folder, kategoria, nazwa FROM bud_arkusz WHERE zrodlo != 'system'")]
    raport = kontrola(rekordy_bs(imp, _wiersze_csv(csv_path), bez_podatkow=True), rekordy_bazy(con, pochodzenie),
                      _nazwy_arkuszy(arkusze, imp['arkusze']), liczba_w_csv(csv_path, bez_podatkow=True), z_aktualnoscia=not k)
    raport['bez_danin'] = True   # VAT i podatek dochodowy: arkusz systemowy, zgodność z BS na ekranie Podatki
    return {**stan, **raport, 'eksport': csv_path.name, 'eksport_sha256': _sha(csv_path)}


# ---------- Okres równoległy: nowy eksport BS kontra plany apki (grupy w Rozjazdach) ----------

def _klucz_wiersza(w: dict) -> tuple:
    return (w['arkusz'], w['rodzaj'], (w.get('kontrahent') or '').strip(), (w.get('komentarz') or '').strip(), (w.get('przedmiot') or '').strip())


def porownaj(imp: dict, model: dict, wiersze_bs: dict) -> list[tuple]:
    """Rozjazdy okresu równoległego (rodzaj, ref, kwota, opis): komórka (wiersz × miesiąc) z inną kwotą w BS niż w apce
    (suma netto wszystkich płatności komórki), wiersz tylko w BS, wiersz tylko w apce. Wiersze łączone po mapie stałych ID
    (pakiet 1), umowa po scaleniu składowych; wiersz dodany osobno w obu miejscach: po arkuszu, rodzaju, kontrahencie,
    komentarzu i przedmiocie."""
    # VAT i podatek dochodowy: w apce arkusz systemowy (pakiet 8), BS liczy swoje prognozy; porównujemy tylko zwykłe wiersze
    wb = {w['id']: w for w in imp['wiersze'] if w.get('podstawa') not in plan.PODATKI}
    wa = {w['id']: w for w in model['wiersze'] if w.get('podstawa') not in plan.PODATKI and w.get('zrodlo') != 'system'}
    pla_bs = [p for p in imp['platnosci'] if p['wiersz'] in wb]
    aliasy = {**aliasy_umow(wb, pla_bs), **wiersze_bs}   # także umowy dodane w BS po kopii
    mapuj = {wid: aliasy.get(wid, wid) for wid in wb}
    wolne_apka = defaultdict(list)
    uzyte = {v for v in mapuj.values() if v in wa}
    for wid, w in wa.items():
        if wid not in uzyte:
            wolne_apka[_klucz_wiersza(w)].append(wid)
    for wid, cel in mapuj.items():
        if cel not in wa:
            cel_bs = aliasy.get(wid, wid)
            kand = wolne_apka.get(_klucz_wiersza(wb.get(cel_bs, wb[wid])))
            if kand:
                nowy = kand.pop(0)
                for x in mapuj:   # składowe tej samej umowy idą za wierszem osoby
                    if aliasy.get(x, x) == cel_bs:
                        mapuj[x] = nowy
    kom_bs, kom_apka, rz_bs = defaultdict(int), defaultdict(int), defaultdict(set)
    ids_apka = {p['id'] for p in model['platnosci']} | {p.get('bs_id') for p in model['platnosci'] if p.get('bs_id')}
    for p in pla_bs:
        k = (mapuj[p['wiersz']], p['miesiac'])
        kom_bs[k] += p['netto']
        if p['typ'] != 'plan' and p['id'] not in ids_apka:
            rz_bs[k].add(p['id'])
    for p in model['platnosci']:
        if p['wiersz'] in wa:
            kom_apka[(p['wiersz'], p['miesiac'])] += p['netto']
    wynik = []
    w_bs = set(mapuj.values())
    for k in sorted(set(kom_bs) | set(kom_apka)):
        wid, m = k
        if wid not in wa or wid not in w_bs:
            continue
        b, a = kom_bs.get(k, 0), kom_apka.get(k, 0)
        if b != a:
            przyczyna = ('w BS jest płatność rzeczywista, której apka jeszcze nie ma' if rz_bs.get(k)
                         else 'komórka pusta w apce' if k not in kom_apka else 'komórka pusta w BS' if k not in kom_bs else 'inna kwota planu')
            wynik.append(('Plan w BS inny niż w apce', f'komorka:{wid}:{m}', b - a, f'BS {b} gr, apka {a} gr; {przyczyna}'))
    for wid in sorted(set(mapuj.values()) - set(wa)):
        zrodlowe = [x for x, c in mapuj.items() if c == wid]
        suma = sum(p['netto'] for p in pla_bs if p['wiersz'] in zrodlowe)
        w = wb[zrodlowe[0]]
        wynik.append(('Wiersz jest tylko w BS', f'wiersz-bs:{wid}', suma,
                      f"{w.get('kontrahent', '')} · {w.get('przedmiot', '') or w.get('podstawa', '')} · {w.get('komentarz', '')}; dodaj go w apce albo usuń w BS"))
    for wid in sorted(set(wa) - w_bs):
        suma = sum(p['netto'] for p in model['platnosci'] if p['wiersz'] == wid)
        wynik.append(('Wiersz jest tylko w apce', f'wiersz:{wid}', suma, 'dodaj go w BS albo usuń w apce'))
    return wynik


# ---------- Excel ----------

def _zapisz_xlsx(cel: Path, raport: dict):
    tresc = xlsx(raport)
    tmp = cel.with_suffix('.tmp')
    tmp.write_bytes(tresc)
    tmp.chmod(0o600)
    tmp.replace(cel)


def xlsx(raport: dict) -> bytes:
    """Raport kontroli sum w Excelu: podsumowanie, różnice komórek, arkusze, miesiące (cash flow i wynik), płatności."""
    from openpyxl import Workbook
    from .eksport import _arkusz_xls
    wb = Workbook()
    ws = wb.active
    ws.title = 'Podsumowanie'
    ws.append(['Eksport BS', raport.get('eksport') or ''])
    ws.append(['Różnic', raport['roznic']])
    ws.append(['Płatności w pliku BS (rodzaj 1 i -1)', raport.get('liczba_csv')])
    ws.append(['Płatności w imporcie BS', raport['liczba_bs']])
    ws.append(['Płatności w apce', raport['liczba_apka']])
    for s in ('bs', 'apka'):
        r = raport['razem'][s]
        ws.append([f"Saldo razem ({'BS' if s == 'bs' else 'apka'}), netto / VAT / brutto", r['netto'] / 100, r['vat'] / 100, r['brutto'] / 100])
    for k, v in (raport.get('weryfikacja') or {}).items():
        ws.append([f'Weryfikacja: {k}', json.dumps(v, ensure_ascii=False)])
    gr = lambda x: x / 100
    _arkusz_xls(wb.create_sheet('Różnice komórek'),
                ['Arkusz', 'Miesiąc', 'Rodzaj', 'Typ', 'Liczba BS', 'Liczba apka', 'Netto BS', 'Netto apka', 'VAT BS', 'VAT apka',
                 'Brutto BS', 'Brutto apka', 'Przyczyny'],
                [[x['arkusz'], x['miesiac'], 'wpływ' if x['rodzaj'] == 1 else 'obciążenie', x['typ'], x['bs']['liczba'], x['apka']['liczba'],
                  gr(x['bs']['netto']), gr(x['apka']['netto']), gr(x['bs']['vat']), gr(x['apka']['vat']), gr(x['bs']['brutto']),
                  gr(x['apka']['brutto']), '; '.join(x['przyczyny'])] for x in raport['komorki']], range(7, 13))
    _arkusz_xls(wb.create_sheet('Arkusze'),
                ['Folder', 'Kategoria', 'Arkusz', 'Liczba BS', 'Liczba apka', 'Netto BS', 'Netto apka', 'VAT BS', 'VAT apka',
                 'Brutto BS', 'Brutto apka', 'Zgodny'],
                [[x['folder'], x['kategoria'], x['nazwa'], x['bs']['liczba'], x['apka']['liczba'], gr(x['bs']['netto']), gr(x['apka']['netto']),
                  gr(x['bs']['vat']), gr(x['apka']['vat']), gr(x['bs']['brutto']), gr(x['apka']['brutto']), 'tak' if x['zgodny'] else 'nie']
                 for x in raport['arkusze']], range(6, 12))
    _arkusz_xls(wb.create_sheet('Miesiące'),
                ['Miesiąc', 'Cash flow BS', 'Cash flow apka', 'Wynik BS', 'Wynik apka', 'Zgodny'],
                [[x['miesiac'], gr(x['cf_bs']), gr(x['cf_apka']), gr(x['wynik_bs']), gr(x['wynik_apka']), 'tak' if x['zgodny'] else 'nie']
                 for x in raport['miesiace']], range(2, 6))
    _arkusz_xls(wb.create_sheet('Płatności'),
                ['ID w BS', 'ID w apce', 'Arkusz', 'Miesiąc', 'Typ', 'Przyczyny', 'Netto BS', 'Netto apka', 'Brutto BS', 'Brutto apka'],
                [[x['bs_pid'] or '', x['apka_id'] or '', x['arkusz'], x['miesiac'], x['typ'], '; '.join(x['przyczyny']),
                  gr(x['bs']['netto']) if x['bs'] else None, gr(x['apka']['netto']) if x['apka'] else None,
                  gr(x['bs']['brutto']) if x['bs'] else None, gr(x['apka']['brutto']) if x['apka'] else None] for x in raport['platnosci']],
                (7, 8, 9, 10))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
