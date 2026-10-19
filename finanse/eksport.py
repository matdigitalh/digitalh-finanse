"""Eksport do Excela (xlsx) i CSV (pakiet 10). Plik powstaje w pamięci i idzie prosto do przeglądarki jako pobranie
(Content-Disposition); apka niczego nie zapisuje na dysku i niczego nie wysyła w sieć.

Adres: /api/eksport/<ścieżka zapytania ekranu>?format=xlsx|csv&<te same parametry co zapytanie ekranu>.
  - Płatności budżetu (zapytanie „budzet/lista” z bieżącymi filtrami) w układzie eksportu Budget Studio: te same
    kolumny w tej samej kolejności co budget_studio_*.csv, kwoty ze znakiem; CSV w formacie pliku BS (tabulator,
    UTF-8 bez BOM, kwoty jak w BS), żeby dało się je porównać 1:1.
  - Każda lista apki z bieżącymi filtrami: eksport woła tę samą funkcję zapytań co ekran, z tymi samymi parametrami,
    bez stronicowania (wszystkie pozycje, nie tylko widoczna strona).
  - Zestawienia: arkusz budżetu (wiersze × miesiące), cash flow wg granulacji, wynik zarządczy, VAT, rentowność.
CSV list: UTF-8 z BOM, średnik, przecinek dziesiętny, bez separatora tysięcy (Excel w polskich ustawieniach otwiera
go wprost). W Excelu kwoty są liczbami (format „zł”), daty datami. Przy kilku tabelach CSV zawiera pierwszą (główną).
Tekst zaczynający się od =, +, @ (albo minus przed tekstem) dostaje apostrof, żeby arkusz nie wziął go za formułę.
"""
import csv
import inspect
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import cashflow, konfig, plan, rentownosc, wynik, zapytania, zapytania_budzet
from .budzet import stawka_proc

NAGLOWEK = Font(bold=True, color='FFFFFF')
TLO = PatternFill('solid', fgColor='171FB5')
ZL = '#,##0.00 "zł";[Red]-#,##0.00 "zł"'
KWOTA_BS = '#,##0.00;[Red]-#,##0.00'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
CSV = 'text/csv; charset=utf-8'
FORMATY = ('xlsx', 'csv')
POMIJANE = {'format', 'strona', 'ile'}   # eksport bierze wszystkie pozycje, nie stronę
MAKS_STRON = 400                          # bezpiecznik pętli stron (400 × 500 pozycji)

# Układ eksportu Budget Studio (budget_studio_*.csv), kolumna po kolumnie.
KOLUMNY_BS = ['ID', 'Rodzaj', 'Status aktualności', 'Etap płatności', 'Podstawa płatności', 'Numer wewnętrzny',
              'Numer dokumentu', 'Kontrahent', 'Tytuł', 'Netto (wartość)', 'Brutto (wartość)', 'VAT', 'Stawka VAT',
              'Miesiąc budżetowy', 'Dzień wystawienia', 'Dzień płatności', 'Dzień płatności - realny', 'Przesunięcie płatności',
              'Miesiąc księgowania', 'Metoda płatności', 'Księgowane', 'W cash flow', 'Komentarz', 'Komentarz w arkuszu',
              'Arkusz', 'Folder', 'Kategoria', 'Sekcja', 'Przedmiot transakcji', 'Data utworzenia', 'Data modyfikacji']
TYPY_BS = {'Rodzaj': 'liczba', 'Netto (wartość)': 'bs', 'Brutto (wartość)': 'bs', 'VAT': 'bs', 'Przesunięcie płatności': 'liczba',
           'Miesiąc budżetowy': 'data', 'Dzień wystawienia': 'data', 'Dzień płatności': 'data', 'Dzień płatności - realny': 'data',
           'Miesiąc księgowania': 'data'}
# Podstawa płatności tak, jak zapisuje ją BS w eksporcie (ze spacją na końcu przy fakturze i bez faktury).
PODSTAWA_BS = {'Faktura (jedna)': 'Faktura ', 'Bez faktury': 'Bez faktury '}
SKLADOWE_BS = {'osoba': 'dla wykonawcy', 'us': 'PIT', 'zus': 'ZUS'}


@dataclass
class Tabela:
    """Jedna tabela eksportu: arkusz w Excelu albo cały plik CSV. kolumny: [(nagłówek, typ)];
    typ: tekst, kwota (grosze), bs (grosze, w CSV zapis jak w BS), data (ISO), liczba, proc."""
    tytul: str
    kolumny: list
    wiersze: list = field(default_factory=list)


# ---------- zapis plików ----------

def _formula(s: str) -> bool:
    """Czy arkusz kalkulacyjny wziąłby ten tekst za formułę (=, +, @, minus przed literą)."""
    if not s:
        return False
    if s[0] in '=+@\t\r':
        return True
    return s[0] == '-' and len(s) > 1 and (s[1].isalpha() or s[1] in '=+@(')   # „-cmd|…”, ale nie „-12,50” ani „- zwrot”


def _kwota_xls(gr):
    return None if gr is None or gr == '' else Decimal(int(gr)).scaleb(-2)


def _data_xls(s):
    if isinstance(s, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', s):
        try:
            return date.fromisoformat(s)
        except ValueError:
            return s
    return s if s not in ('', None) else None


def _xlsx(tabele: list[Tabela]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for t in tabele:
        ws = wb.create_sheet(re.sub(r'[\[\]:*?/\\]', ' ', t.tytul)[:31] or 'Dane')
        ws.append([n for n, _ in t.kolumny])
        for c in ws[1]:
            c.font, c.fill = NAGLOWEK, TLO
        for w in t.wiersze:
            ws.append([_kwota_xls(v) if typ in ('kwota', 'bs') else _data_xls(v) if typ == 'data' else v
                       for v, (_, typ) in zip(w, t.kolumny)])
        for i, (n, typ) in enumerate(t.kolumny, start=1):
            for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                if typ == 'kwota':
                    c.number_format = ZL
                elif typ == 'bs':
                    c.number_format = KWOTA_BS
                elif typ == 'data' and isinstance(c.value, date):
                    c.number_format = 'yyyy-mm-dd'
                elif isinstance(c.value, str) and c.value.startswith('='):
                    c.data_type = 's'   # tekst (np. tytuł przelewu), nie formuła
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, len(str(n)) + 4))
        ws.freeze_panes = 'B2'
    if not tabele:
        wb.create_sheet('Dane')
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _arkusz_xls(ws, naglowki, wiersze, kolumny_kwot=()):
    """Arkusz z gotowych wierszy (kwoty już w złotych): raport kontroli sum kopii z BS (pakiet 6, kopia_bs.xlsx)."""
    ws.append(naglowki)
    for c in ws[1]:
        c.font, c.fill = NAGLOWEK, TLO
    for w in wiersze:
        ws.append(w)
    for i in kolumny_kwot:
        for c in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            c[0].number_format = ZL
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                c.data_type = 's'   # tekst, nie formuła
    for i, n in enumerate(naglowki, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(45, len(str(n)) + 4))
    ws.freeze_panes = 'B2'


def _kwota_csv(gr) -> str:
    if gr is None or gr == '':
        return ''
    gr = int(gr)
    return f"{'-' if gr < 0 else ''}{abs(gr) // 100},{abs(gr) % 100:02d}"


def kwota_bs(gr) -> str:
    """Kwota tak, jak zapisuje ją eksport BS: przecinek, bez zbędnych zer (-10,31; 0; 9999,9)."""
    if gr is None or gr == '':
        return ''
    t = format(Decimal(int(gr)).scaleb(-2).normalize(), 'f')
    if '.' in t:
        t = t.rstrip('0').rstrip('.')
    return t.replace('.', ',')


def _tekst_csv(v) -> str:
    s = '' if v is None else str(v)
    return "'" + s if _formula(s) else s


def _komorka_csv(v, typ):
    if typ == 'kwota':
        return _kwota_csv(v)
    if typ == 'bs':
        return kwota_bs(v)
    if typ == 'proc':
        return '' if v is None else str(v).replace('.', ',')
    if typ in ('liczba', 'data'):
        return '' if v is None else str(v)
    return _tekst_csv(v)


def _csv(t: Tabela, bs=False) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter='\t' if bs else ';', lineterminator='\n' if bs else '\r\n')
    w.writerow([n for n, _ in t.kolumny])
    for wiersz in t.wiersze:
        w.writerow([_komorka_csv(v, typ) for v, (_, typ) in zip(wiersz, t.kolumny)])
    return (buf.getvalue() if bs else '﻿' + buf.getvalue()).encode('utf-8')


def _nazwa(*czesci) -> str:
    """Nazwa pliku do pobrania: „Finanse <co> <dzień>”, bez znaków niedozwolonych w nazwach plików."""
    return re.sub(r'\s+', ' ', re.sub(r'[\\/:*?"<>|\n\r\t]+', ' ', ' '.join(str(c) for c in ('Finanse', *czesci) if c))).strip()


# ---------- wywołanie zapytania ekranu z tymi samymi parametrami ----------

def wolaj(fn, con, q: dict, aliasy: dict | None = None, konwersje: dict | None = None):
    """Woła funkcję zapytań ekranu z parametrami z adresu (po nazwie; aliasy jak w serwer.py), bez stronicowania:
    gdy funkcja ma strony, zbiera wszystkie po 500. Parametry, których funkcja nie zna, pomija."""
    sig = inspect.signature(fn)
    dowolne = any(p.kind == p.VAR_KEYWORD for p in sig.parameters.values())
    kw = {}
    for k, v in q.items():
        cel = (aliasy or {}).get(k, k)
        if k in POMIJANE or cel in ('con', 'strona', 'ile') or v in ('', None):
            continue
        if cel == 'rok' and not (len(str(v)) == 4 and str(v).isdigit()):   # jak serwer: zły rok = rok domyślny
            continue
        if cel in sig.parameters or dowolne:
            kw[cel] = (konwersje or {}).get(cel, lambda x: x)(v)
    if 'strona' not in sig.parameters:
        return fn(con, **kw)
    wynik_, wiersze = None, []
    for strona in range(1, MAKS_STRON + 1):
        wynik_ = fn(con, **kw, strona=strona, ile=500)
        wiersze += wynik_['wiersze']
        if not wynik_['wiersze'] or len(wiersze) >= wynik_['razem']:
            break
    wynik_['wiersze'] = wiersze
    return wynik_


def _tab(tytul, spec, dane) -> Tabela:
    """Tabela z listy słowników: spec = [(nagłówek, klucz albo funkcja, typ)]."""
    return Tabela(tytul, [(n, t) for n, _, t in spec],
                  [[(k(x) if callable(k) else x.get(k)) for _, k, _ in spec] for x in dane])


def _filtry(q: dict) -> str:
    return ', '.join(f'{k}: {v}' for k, v in sorted(q.items()) if k not in POMIJANE and v not in ('', None)) or 'bez filtrów'


def _podsumowanie(pozycje: list[tuple]) -> Tabela:
    """Arkusz „Podsumowanie”: (co, opis, kwota w groszach albo None)."""
    return Tabela('Podsumowanie', [('Co', 'tekst'), ('Wartość', 'tekst'), ('Kwota', 'kwota')], [list(p) for p in pozycje])


def _dzien_stanu(con) -> str:
    return zapytania_budzet._dzien_stanu(con)


def _tak(v) -> str:
    return 'Nie' if v is not None and not v else 'Tak'


# ---------- płatności budżetu w układzie Budget Studio ----------

def _czas_lokalny(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).astimezone().strftime('%Y-%m-%d %H:%M:%S')
    except (TypeError, ValueError):
        return ''


def _z_pliku_bs(con) -> tuple[dict, dict]:
    """Pola, których baza nie trzyma, a eksport BS ma: data utworzenia i modyfikacji oraz pełny „Miesiąc budżetowy”
    (BS zapisuje tam czasem dzień inny niż 1.). Płatność z importu BS bierze je z pliku eksportu (po numerze wiersza),
    zmieniona w apce: daty z historii zmian (pierwsza i ostatnia zmiana jej albo jej komórki)."""
    z_pliku = {}
    r = con.execute("SELECT wartosc FROM meta WHERE klucz='import_bs'").fetchone()
    nazwa = (json.loads(r[0]) or {}).get('plik', '') if r and r[0] else ''
    p = konfig.katalog_dane() / 'Budget Studio' / nazwa
    if nazwa and '/' not in nazwa and p.is_file():
        with open(p, encoding='utf-8-sig', newline='') as f:
            for nr, x in enumerate(csv.DictReader(f, delimiter='\t'), start=2):
                z_pliku[nr] = (x.get('Data utworzenia') or '', x.get('Data modyfikacji') or '', (x.get('Miesiąc budżetowy') or '')[:10])
    historia = {r[0]: (_czas_lokalny(r[1]), _czas_lokalny(r[2]))
                for r in con.execute('SELECT obiekt, MIN(czas), MAX(czas) FROM bud_historia GROUP BY obiekt')}
    return z_pliku, historia


def _stawka_bs(x: dict) -> str:
    """Stawka jak w BS: ułamek (0,23). Faktura bez stawki standardowej, ale z VAT (kilka stawek na fakturze): stawka
    efektywna VAT / netto, jak w BS; podstawa bez VAT: 0."""
    st = x.get('stawka')
    if st is not None:
        return kwota_bs(int(st))
    if not plan.ma_vat(x.get('podstawa') or ''):
        return '0'
    if x.get('netto') and x.get('vat'):
        return format((Decimal(int(x['vat'])) / Decimal(int(x['netto']))).quantize(Decimal('0.0001')).normalize(), 'f').replace('.', ',')
    return ''


def podstawa_bs(podstawa: str, skladowa: str | None) -> str:
    """Podstawa w zapisie BS: „Faktura ”, „Umowa o pracę - dla wykonawcy”, „Umowa zlecenie - PIT” itd."""
    if skladowa and plan.jest_umowa(podstawa):
        return f'{podstawa} - {SKLADOWE_BS.get(skladowa, skladowa)}'
    return PODSTAWA_BS.get(podstawa, podstawa or '')


def _dni(a, b) -> int:
    try:
        return (date.fromisoformat(b) - date.fromisoformat(a)).days
    except (TypeError, ValueError):
        return 0


def wiersz_bs(x: dict, d: dict, z_pliku: dict, historia: dict) -> list:
    """Płatność z listy (zapytania_budzet.lista) jako wiersz eksportu BS; d: cechy z wiersza i arkusza budżetu."""
    znak = int(x.get('rodzaj') or 1)
    realny = x.get('realny') or x.get('termin')
    ksieg = x.get('ksiegowanie') or ''
    if len(ksieg) == 7:
        ksieg += '-01'
    utw, mod, mies_bs = z_pliku.get(x.get('wiersz_pliku'), ('', '', '')) if x.get('bs_id') else ('', '', '')
    miesiac = mies_bs if x.get('miesiac') and mies_bs[:7] == x['miesiac'] else f"{x['miesiac']}-01" if x.get('miesiac') else ''
    h = historia.get(x['id']) or historia.get(f"{x.get('wiersz')}:{x.get('miesiac')}")
    if h:
        utw, mod = utw or h[0], h[1]
    aktualna = x.get('aktualna') in (1, True) and (x.get('aktualnosc') or 'Aktualna') == 'Aktualna'
    return [x.get('bs_id') or x['id'], znak, 'Aktualne' if aktualna else 'Nieaktualne', x.get('etap_wg_bs') or 'Planowana',
            podstawa_bs(x.get('podstawa') or '', x.get('skladowa')), x.get('nr_wewn') or '', x.get('nr_dok') or '',
            x.get('w_kontrahent') or '', x.get('tytul') or '', (x.get('netto') or 0) * znak, (x.get('brutto') or 0) * znak,
            (x.get('vat') or 0) * znak, _stawka_bs(x), miesiac, x.get('wystawienie') or '',
            x.get('termin') or '', realny or '', _dni(x.get('termin'), realny) if x.get('termin') and realny else 0, ksieg,
            x.get('metoda') or d.get('w_metoda') or '', _tak(x.get('ksiegowane')), _tak(x.get('w_cashflow')), x.get('komentarz') or '',
            d.get('w_komentarz') or '', x.get('arkusz') or '', x.get('folder') or '', d.get('kategoria') or '', d.get('sekcja') or '-',
            d.get('przedmiot') or 'Przedmiot transakcji', utw, mod]


def platnosci_bs(con, q: dict) -> tuple[str, list[Tabela], bool]:
    d = wolaj(zapytania_budzet.lista, con, q, aliasy={'lista': 'nazwa'})
    cechy = {r['id']: dict(r) for r in con.execute(
        'SELECT p.id, w.komentarz AS w_komentarz, w.sekcja, w.przedmiot, w.metoda AS w_metoda, a.kategoria FROM bud_platnosc p'
        ' JOIN bud_wiersz w ON w.id = p.wiersz JOIN bud_arkusz a ON a.id = w.arkusz')}
    z_pliku, historia = _z_pliku_bs(con)
    wiersze = [wiersz_bs(x, cechy.get(x['id'], {}), z_pliku, historia) for x in d['wiersze']]
    brutto = [w[10] for w in wiersze]
    tab = Tabela('Płatności', [(k, TYPY_BS.get(k, 'tekst')) for k in KOLUMNY_BS], wiersze)
    pods = _podsumowanie([
        ('Lista', d.get('nazwa', ''), None), ('Filtry', _filtry(q), None), ('Dzień stanu', d.get('dzien_stanu', ''), None),
        ('Płatności', str(len(wiersze)), None),
        ('Suma brutto (jak w stopce listy)', '', sum(brutto)), ('Suma netto', '', sum(w[9] for w in wiersze)),
        ('Wpływy brutto', '', sum(b for b in brutto if b > 0)), ('Obciążenia brutto', '', sum(b for b in brutto if b < 0)),
        ('Układ kolumn', 'jak w eksporcie Budget Studio (budget_studio_*.csv)', None)])
    return _nazwa('płatności budżetu', d.get('nazwa', ''), d.get('dzien_stanu', '')), [tab, pods], True


# ---------- listy apki ----------

STAN_WPLATY = {'dopasowana': 'połączona z fakturą', 'nieznana': 'do połączenia', 'sprzed_eksportu': 'faktura sprzed eksportu',
               'bez_faktury': 'nie dotyczy faktury', 'zwrot_zakupu': 'zwrot zakupu'}

BANK = [('Data księgowania', 'data_ksiegowania', 'data'), ('Data transakcji', 'data_transakcji', 'data'), ('Rachunek', 'rachunek', 'tekst'),
        ('Kwota', 'kwota', 'kwota'), ('Waluta', 'waluta', 'tekst'), ('Kwota oryginalna', 'kwota_oryg', 'kwota'),
        ('Waluta oryginalna', 'waluta_oryg', 'tekst'), ('Saldo po', 'saldo_po', 'kwota'), ('Kontrahent', 'kontrahent', 'tekst'),
        ('Rachunek kontrahenta', 'rachunek_kontrahenta', 'tekst'), ('Opis', 'opis', 'tekst'), ('Tytuł', 'tytul', 'tekst'),
        ('Numer referencyjny', 'nr_ref', 'tekst'), ('Ruch wewnętrzny', lambda x: 'tak' if x.get('wewnetrzna') else 'nie', 'tekst'),
        ('Faktury', 'faktury', 'tekst'), ('Stan wpłaty', lambda x: STAN_WPLATY.get(x.get('stan_wplaty'), x.get('stan_wplaty') or ''), 'tekst'),
        ('Plik', 'plik', 'tekst'), ('Wiersz w pliku', 'wiersz', 'liczba')]

FAKTURY = [('Numer', 'numer', 'tekst'), ('Typ', 'typ', 'tekst'), ('Status w Fakturowni', 'status', 'tekst'),
           ('Stan w banku', 'stan_bank', 'tekst'), ('Klient', 'klient', 'tekst'), ('Nabywca', 'nabywca', 'tekst'), ('NIP', 'nip', 'tekst'),
           ('Netto', 'netto', 'kwota'), ('VAT', 'vat', 'kwota'), ('Brutto', 'brutto', 'kwota'), ('Opłacono (Fakturownia)', 'oplacono', 'kwota'),
           ('Zapłacono (bank)', 'zaplacono', 'kwota'), ('Pozostało', 'pozostalo', 'kwota'), ('Dni po terminie', 'dni_po_terminie', 'liczba'),
           ('Data wystawienia', 'data_wystawienia', 'data'), ('Data sprzedaży', 'data_sprzedazy', 'data'), ('Termin', 'termin', 'data'),
           ('Ostatnia wpłata', 'ostatnia_wplata', 'data'), ('Numer KSeF', 'ksef', 'tekst'),
           ('Korekta', lambda x: 'tak' if x.get('korekta') else '', 'tekst'), ('Faktura pierwotna', 'oryginal', 'tekst'),
           ('Pozycje', 'pozycje', 'tekst'), ('Wpłaty', 'wplaty', 'tekst'), ('Uwagi', 'uwagi', 'tekst'), ('Plik', 'plik', 'tekst'),
           ('Wiersz w pliku', 'wiersz', 'liczba')]

NALEZNOSCI = [('Numer', 'numer', 'tekst'), ('Nabywca', 'nabywca', 'tekst'), ('NIP', 'nip', 'tekst'),
              ('Data wystawienia', 'data_wystawienia', 'data'), ('Termin', 'termin', 'data'), ('Brutto', 'brutto', 'kwota'),
              ('Korekty', 'korekty', 'kwota'), ('Należność', 'naleznosc', 'kwota'), ('Zapłacono', 'zaplacono', 'kwota'),
              ('Pozostało', 'pozostalo', 'kwota'), ('Dni po terminie', 'dni_po_terminie', 'liczba'), ('Wiekowanie', 'przedzial', 'tekst'),
              ('Stan w banku', 'stan_bank', 'tekst'), ('Ostatnia wpłata', 'ostatnia_wplata', 'data'), ('Plik', 'plik', 'tekst'),
              ('Wiersz w pliku', 'wiersz', 'liczba')]

ZAKUPY = [('Rejestr', 'rejestr', 'tekst'), ('Nr ewidencyjny', 'nr_ewid', 'tekst'), ('Numer dokumentu', 'nr_dok', 'tekst'),
          ('Data', 'data', 'data'), ('Data zaewidencjonowania', 'data_zaewid', 'data'), ('Kontrahent', 'podmiot', 'tekst'),
          ('Kod kontrahenta', 'podmiot_kod', 'tekst'), ('Netto', 'netto', 'kwota'), ('VAT', 'vat', 'kwota'), ('Brutto', 'brutto', 'kwota'),
          ('Płatność w banku', 'stan', 'tekst'), ('Data płatności', 'data_platnosci', 'data'),
          ('Kwota płatności', lambda x: -x['kwota_platnosci'] if x.get('kwota_platnosci') is not None else None, 'kwota'),
          ('Plik', 'plik', 'tekst'), ('Wiersz w pliku', 'wiersz', 'liczba')]

WPLATY = [('Dzień płatności', 'data_ksiegowania', 'data'), ('Płatnik', 'kontrahent', 'tekst'), ('Tytuł', 'tytul', 'tekst'),
          ('Kwota', 'kwota', 'kwota'), ('Stan', lambda x: STAN_WPLATY.get(x.get('stan'), x.get('stan') or ''), 'tekst'),
          ('Pewność', 'pewnosc', 'tekst'), ('Nadpłata', 'nadplata', 'kwota'), ('Faktury', 'faktury', 'tekst'),
          ('Podpowiedź', 'podpowiedz', 'tekst'), ('Uwaga', 'uwaga', 'tekst'), ('Plik', 'plik', 'tekst'), ('Wiersz w pliku', 'wiersz', 'liczba')]

WERYFIKACJA = [('Rodzaj', 'tytul', 'tekst'), ('Data', 'data', 'data'), ('Kwota', 'kwota', 'kwota'), ('Kontrahent', 'kontrahent', 'tekst'),
               ('Faktury', 'faktury', 'tekst'), ('Opis', 'opis', 'tekst')]

WYDATKI = [('Data', 'data_ksiegowania', 'data'), ('Kontrahent', 'kontrahent', 'tekst'), ('Tytuł', 'tytul', 'tekst'),
           ('Kwota', 'kwota', 'kwota'), ('Rodzaj wydatku', 'rodzaj', 'tekst'), ('Klient', 'klient', 'tekst'),
           ('Skąd przypisanie', 'zrodlo', 'tekst'), ('Dokument', 'dokument', 'tekst'), ('Netto z dokumentu', 'netto_dok', 'kwota'),
           ('Dokumenty z rejestru', 'dokumenty', 'tekst'), ('Rachunek', 'rachunek', 'tekst'), ('Plik', 'plik', 'tekst'),
           ('Wiersz w pliku', 'wiersz', 'liczba')]

ROZJAZDY = [('Rodzaj', 'rodzaj', 'tekst'), ('Arkusz', 'arkusz', 'tekst'), ('Kontrahent', 'kontrahent', 'tekst'), ('Miesiąc', 'miesiac', 'tekst'),
            ('Kwota', 'kwota', 'kwota'), ('Opis', 'opis', 'tekst'), ('Odnośnik', 'ref', 'tekst')]
NIEPRZYPISANE = [('Kontrahent (klucz)', 'klucz', 'tekst'), ('Przykład nazwy', 'przyklad', 'tekst'), ('Płatności', 'ile', 'liczba'),
                 ('Razem', 'kwota', 'kwota'), ('Ostatnia', 'ostatnia', 'data'), ('Konto reklamowe', 'konto', 'tekst')]
KONTA = [('Konto', 'konto', 'tekst'), ('Płatności', 'platnosci', 'liczba'), ('Razem', 'suma', 'kwota'), ('Ostatnia', 'ostatnia', 'data'),
         ('Klient', 'klient', 'tekst'), ('Podpowiedź', 'podpowiedz', 'tekst')]
WYNIKI_BIURA = [('Okres', 'okres', 'tekst'), ('Kod', 'kod', 'tekst'), ('Pozycja', 'opis', 'tekst'), ('Bieżący okres', 'wartosc', 'kwota'),
                ('Poprzedni rok', 'wartosc_poprzednia', 'kwota'), ('Plik', 'plik', 'tekst'), ('Wiersz w pliku', 'wiersz', 'liczba')]


def _lista(tytul, spec, fn, aliasy=None, konwersje=None, nazwa_pliku=None, suma=None):
    """Eksport listy ekranu: ta sama funkcja zapytań, wszystkie pozycje, arkusz Podsumowanie z liczbą i sumami listy."""
    def eksport(con, q):
        d = wolaj(fn, con, q, aliasy, konwersje)
        dane = d['wiersze'] if isinstance(d, dict) else d
        pods = [('Lista', tytul, None), ('Filtry', _filtry(q), None), ('Pozycji', str(len(dane)), None)]
        if suma and isinstance(d, dict):
            pods += [(n, '', d.get(k)) for n, k in suma if d.get(k) is not None]
        return _nazwa(nazwa_pliku or tytul.lower(), _dzien_stanu(con)), [_tab(tytul, spec, dane), _podsumowanie(pods)], False
    return eksport


def _naleznosci(con, q):
    d = wolaj(zapytania.naleznosci, con, q)
    n = zapytania.podsumowanie_naleznosci(con)
    wiek = Tabela('Wiekowanie', [('Przedział', 'tekst'), ('Faktur', 'liczba'), ('Pozostało', 'kwota')],
                  [[k, v[0], v[1]] for k, v in n['przedzialy'].items()] + [['Razem', None, n['razem']]])
    pods = _podsumowanie([('Lista', 'Należności z wiekowaniem', None), ('Filtry', _filtry(q), None),
                          ('Stan na dzień', d.get('dzien_stanu', ''), None), ('Pozycji', str(len(d['wiersze'])), None),
                          ('Pozostało do zapłaty', '', d.get('pozostalo'))])
    return _nazwa('należności', d.get('dzien_stanu', '')), [_tab('Należności', NALEZNOSCI, d['wiersze']), wiek, pods], False


def _macierz(tytul, miesiace, wiersze, nazwa_kol='Pozycja', razem=True) -> Tabela:
    """Tabela pozycje × miesiące (kwoty w groszach): wiersze = [(nazwa, {miesiąc: kwota}, razem)]."""
    kol = [(nazwa_kol, 'tekst')] + [(m, 'kwota') for m in miesiace] + ([('Razem', 'kwota')] if razem else [])
    return Tabela(tytul, kol, [[n, *[w.get(m, 0) for m in miesiace], *([r] if razem else [])] for n, w, r in wiersze])


def _wynik(con, q):
    d = wolaj(wynik.wynik_zarzadczy, con, q)
    t = _macierz('Wynik zarządczy', d['miesiace'], [(w['nazwa'], w['miesiace'], w['razem']) for w in d['wiersze']])
    pods = _podsumowanie([('Rok', d['rok'], None), ('Miesiące niepełne', ', '.join(d.get('niepelne') or []) or 'brak', None),
                          ('Obrót netto', '', d['obrot']), ('Wynik zarządczy', '', d['wynik']),
                          ('Marża', f"{d['marza_proc']} %" if d.get('marza_proc') is not None else '', None)])
    return _nazwa('wynik zarządczy', d['rok']), [t, pods], False


def _obroty(con, q):
    d = wolaj(wynik.obroty_klientow, con, q)
    t = _macierz('Obroty wg klientów', d['miesiace'], [(k['klient'], k['miesiace'], k['obrot']) for k in d['klienci']], 'Klient')
    t.kolumny += [('Koszty klienta', 'kwota'), ('Po kosztach klienta', 'kwota'), ('Udział w obrocie %', 'proc'), ('Marża %', 'proc')]
    for w, k in zip(t.wiersze, d['klienci']):
        w += [k['koszty'], k['pna'], k['udzial'], k['marza_proc']]
    return _nazwa('obroty wg klientów', d['rok']), [t], False


def _koszty(con, q):
    d = wolaj(wynik.koszty_kategorii, con, q)
    t = _macierz('Koszty wg kategorii', d['miesiace'], [(x['nazwa'], x['miesiace'], x['razem']) for x in d['pozycje']],
                 'Kontrahent' if d.get('rodzaj') else 'Rodzaj kosztu')
    return _nazwa('koszty', d.get('rodzaj') or 'wg kategorii', d['rok']), [t], False


VAT = [('Miesiąc', 'miesiac', 'tekst'), ('VAT należny (faktury)', 'nalezny', 'kwota'), ('VAT naliczony (rejestr biura)', 'naliczony', 'kwota'),
       ('Do zapłaty', 'do_zaplaty', 'kwota'), ('Zapłacony (bank)', 'zaplacony', 'kwota'),
       ('Rejestr biura obejmuje', lambda x: 'tak' if x.get('rejestr') else 'nie', 'tekst')]


def _vat(con, q):
    d = wolaj(wynik.vat, con, q)
    return _nazwa('VAT', d['rok']), [_tab('VAT', VAT, d['miesiace'])], False


RENT_KLIENCI = [('Klient', 'klient', 'tekst'), ('Obrót netto', 'obrot', 'kwota'), ('Koszty klienta', 'koszty', 'kwota'),
                ('Po kosztach klienta', 'pna', 'kwota'), ('Marża %', 'marza_proc', 'proc'), ('Udział czasu %', 'udzial_czasu', 'proc'),
                ('Koszt pracy', 'koszt_pracy', 'kwota'), ('Wynik klienta', 'wynik_klienta', 'kwota')]
RENT_ARKUSZE = [('Folder', 'folder', 'tekst'), ('Kategoria', 'kategoria', 'tekst'), ('Arkusz', 'nazwa', 'tekst'),
                ('Plan przychód', 'plan_przychod', 'kwota'), ('Plan koszt', 'plan_koszt', 'kwota'), ('Saldo planu', 'plan_saldo', 'kwota'),
                ('Wykonanie przychód', 'wyk_przychod', 'kwota'), ('Wykonanie koszt', 'wyk_koszt', 'kwota'),
                ('Saldo wykonania', 'wyk_saldo', 'kwota'), ('Realizacja %', 'realizacja_proc', 'proc'), ('Po terminie', 'po_terminie', 'liczba')]


def _rent_klienci(con, q):
    d = wolaj(rentownosc.klienci, con, q)
    return _nazwa('rentowność klientów', d['rok']), [_tab('Klienci', RENT_KLIENCI, d['klienci'])], False


def _rent_arkusze(con, q):
    return _nazwa('arkusze plan i wykonanie', _dzien_stanu(con)), [_tab('Arkusze', RENT_ARKUSZE, rentownosc.arkusze(con))], False


def _rentownosc_oba(con, q):
    """Eksport rentowności sprzed pakietu 10: klienci i arkusze w jednym pliku."""
    d = wolaj(rentownosc.klienci, con, q)
    return (_nazwa('rentowność', d['rok']),
            [_tab('Klienci', RENT_KLIENCI, d['klienci']), _tab('Arkusze', RENT_ARKUSZE, rentownosc.arkusze(con))], False)


KLIENCI = [('Klient', 'klient', 'tekst'), ('Faktur', 'faktur', 'liczba'), ('Przychód netto', 'przychod', 'kwota'),
           ('Wpłaty', 'wplaty', 'kwota'), ('Należność', 'naleznosc', 'kwota'), ('Po terminie', 'po_terminie', 'kwota'),
           ('Koszty klienta', 'koszty', 'kwota'), ('Koszty bez dokumentu', 'koszty_bez_dok', 'kwota'), ('Marża', 'marza', 'kwota'),
           ('Marża %', 'marza_proc', 'proc')]


def _klienci(con, q):
    d = wolaj(zapytania.klienci_lista, con, q)
    return _nazwa('klienci', d.get('rok') or 'wszystkie lata'), [_tab('Klienci', KLIENCI, d['klienci'])], False


def _klient(con, q):
    d = wolaj(zapytania.klient, con, q)
    fak = [('Numer', 'numer', 'tekst'), ('Nabywca', 'nabywca', 'tekst'), ('Netto', 'netto', 'kwota'), ('Brutto', 'brutto', 'kwota'),
           ('Data wystawienia', 'data_wystawienia', 'data'), ('Data sprzedaży', 'data_sprzedazy', 'data'), ('Stan w banku', 'stan', 'tekst'),
           ('Pozostało', 'pozostalo', 'kwota')]
    kos = [('Data', 'data_ksiegowania', 'data'), ('Kontrahent', 'kontrahent', 'tekst'), ('Tytuł', 'tytul', 'tekst'),
           ('Kwota z banku', 'kwota', 'kwota'), ('Koszt klienta', 'koszt', 'kwota'), ('Dokument', 'dokument', 'tekst')]
    mies = [('Miesiąc', 'miesiac', 'tekst'), ('Przychód netto', 'przychod', 'kwota'), ('Koszty', 'koszty', 'kwota')]
    return (_nazwa('klient', d['klient'], d.get('rok') or ''),
            [_tab('Miesiące', mies, d['miesiace']), _tab('Faktury', fak, d['faktury']), _tab('Koszty', kos, d['koszty'])], False)


def _stawka_wiersza(w: dict) -> str:
    if not plan.ma_vat(w.get('podstawa') or ''):
        return ''
    v = stawka_proc(w.get('stawka_vat'))
    return f'{v} %' if v is not None else str(w.get('stawka_vat') or '')


def _arkusz(con, q):
    """Arkusz budżetu jak na ekranie: wiersze z ustawieniami × miesiące (netto ze znakiem), suma wiersza, pod spodem
    saldo okresu, wpływy, obciążenia i saldo narastająco."""
    d = zapytania_budzet.arkusz(con, q.get('id') or '')
    a, m = d['arkusz'], d['miesiace']
    kol = [('Przedmiot transakcji', 'tekst'), ('Kontrahent', 'tekst'), ('Sekcja', 'tekst'), ('Komentarz', 'tekst'), ('Rodzaj', 'tekst'),
           ('Podstawa', 'tekst'), ('Stawka VAT', 'tekst'), ('Dzień wystawienia', 'tekst'), ('Dni na płatność', 'tekst'), ('Metoda', 'tekst'),
           ('Księgowane', 'tekst'), ('W cash flow', 'tekst'), *[(x, 'kwota') for x in m], ('Suma', 'kwota')]
    wiersze = []
    for w in d['wiersze']:
        kwoty = [w['komorki'][x]['netto'] if x in w['komorki'] else None for x in m]
        wiersze.append([w.get('przedmiot') or '', w.get('kontrahent') or '', w.get('sekcja') or '', w.get('komentarz') or '',
                        'wpływy' if w['rodzaj'] == 1 else 'obciążenia', w.get('podstawa') or '', _stawka_wiersza(w),
                        w.get('dzien_opis') or w.get('dzien_wystawienia') or '', str(w.get('dni_na_platnosc') or ''), w.get('metoda') or '',
                        _tak(w.get('ksiegowane')), _tak(w.get('w_cashflow')), *kwoty, sum(k or 0 for k in kwoty)])
    p, pusto = d['podsumowanie'], [''] * 11
    for nazwa, klucz in (('Saldo okresu', 'saldo'), ('Wpływy', 'wplywy'), ('Obciążenia', 'obciazenia'), ('Saldo narastająco', 'narastajaco')):
        wiersze.append([nazwa, *pusto, *[x[klucz] for x in p], sum(x[klucz] for x in p) if klucz != 'narastajaco' else None])
    pods = _podsumowanie([('Arkusz', a['nazwa'], None), ('Folder', a['folder'], None), ('Kategoria', a['kategoria'], None),
                          ('Zakres', f"{a['od']} do {a['do']}", None),
                          ('Kwoty', 'netto; wpływy ze znakiem plus, obciążenia ze znakiem minus', None)])
    return _nazwa('arkusz', a['nazwa']), [Tabela('Arkusz', kol, wiersze), pods], False


POZYCJE_CF = [('Data', 'data', 'data'), ('Okres', 'okres_etykieta', 'tekst'), ('Kwota', 'kwota', 'kwota'), ('Opis', 'opis', 'tekst'),
              ('Rodzaj', 'rodzaj', 'tekst'), ('Typ', 'typ', 'tekst'), ('Scenariusz', 'scenariusz', 'tekst'), ('Źródło', 'zrodlo', 'tekst'),
              ('Dzień płatności (termin)', 'termin', 'data'), ('Przesunięcie (dni)', 'przesuniecie', 'liczba')]
GRANULACJE = {'dzien': 'dzień', 'tydzien': 'tydzień', 'miesiac': 'miesiąc'}


def _cashflow(con, q):
    """Cash flow wg granulacji (te same okresy, saldo startowe i filtry co ekran): okres w wierszu, pozycje osobno."""
    d = wolaj(cashflow.okresy, con, q)
    if d.get('brak_danych'):
        raise ValueError('brak danych z banku: nie ma od czego liczyć cash flow')
    ok = d['okresy']
    t = Tabela('Cash flow', [('Okres', 'tekst'), ('Od', 'data'), ('Do', 'data'), ('Wpływy', 'kwota'), ('Obciążenia', 'kwota'),
                             ('Saldo okresu', 'kwota'), ('Saldo narastająco', 'kwota'), ('Pozycji', 'liczba'), ('Stan okresu', 'tekst')],
               [[o['etykieta'], o['od'], o['do'], o['wplywy'], o['obciazenia'], o['saldo'], o['narastajaco'], o['pozycji'],
                 'bieżący' if o.get('biezacy') else 'miniony' if o.get('przeszly') else 'prognoza'] for o in ok])
    pozycje = [{**z, 'okres_etykieta': ok[z['okres']]['etykieta'] if 0 <= z.get('okres', -1) < len(ok) else ''} for z in d['pozycje']]
    pods = _podsumowanie([
        ('Zakres', f"{d['od']} do {d['do']}", None), ('Granulacja', GRANULACJE.get(d['granulacja'], d['granulacja']), None),
        ('Scenariusz', d['scenariusz'], None), ('Filtry', _filtry(q), None), ('Saldo startowe', d['zrodlo_salda'], d['saldo_start']),
        ('SUM (saldo na koniec)', '', d['sum']), ('MAX', ok[d['max']['okres']]['etykieta'], d['max']['saldo']),
        ('MIN', ok[d['min']['okres']]['etykieta'], d['min']['saldo']), ('Dzień stanu', d['dzien_stanu'], None)])
    return (_nazwa('cash flow', GRANULACJE.get(d['granulacja'], ''), d['od'], 'do', d['do']),
            [t, _tab('Pozycje', POZYCJE_CF, pozycje), pods], False)


def _cashflow_13(con, q):
    """Eksport sprzed pakietu 10: prognoza 13 tygodni (scenariusz ostrożny i pełny), pozycje, saldo dzienne."""
    d = cashflow.prognoza(con)
    if d.get('brak_danych'):
        raise ValueError('brak danych z banku: nie ma od czego liczyć cash flow')
    t = Tabela('13 tygodni', [('Od', 'data'), ('Do', 'data'), ('Wpływy', 'kwota'), ('Wydatki', 'kwota'), ('Tylko w scenariuszu pełnym', 'kwota'),
                              ('Saldo ostrożne', 'kwota'), ('Saldo pełne', 'kwota')],
               [[str(x['od']), str(x['do']), x['wplywy'], x['wydatki'], x['tylko_pelny'], x['saldo_ostrozny'], x['saldo_pelny']]
                for x in d['tygodnie']])
    poz = Tabela('Pozycje', [('Data', 'data'), ('Kwota', 'kwota'), ('Opis', 'tekst'), ('Rodzaj', 'tekst'), ('Scenariusz', 'tekst')],
                 [[str(p['data']), p['kwota'], p['opis'], p['rodzaj'], p['scenariusz']] for x in d['tygodnie'] for p in x['pozycje']])
    dni = Tabela('Dziennie', [('Data', 'data'), ('Saldo ostrożne', 'kwota'), ('Saldo pełne', 'kwota')],
                 [[str(x['data']), x['ostrozny'], x['pelny']] for x in d['dni']])
    return _nazwa('cash flow 13 tygodni', d['dzien_stanu']), [t, poz, dni], False


# Ścieżka zapytania ekranu → eksport. Funkcja zwraca (nazwa pliku bez rozszerzenia, tabele, czy CSV w formacie BS).
EKSPORTY = {
    'budzet/lista': platnosci_bs,
    'operacje': _lista('Bank ING', BANK, zapytania.operacje, aliasy={'kolejnosc': 'malejaco'},
                       konwersje={'malejaco': lambda v: v == 'malejaco'}, nazwa_pliku='bank ING',
                       suma=[('Wpływy', 'wplywy'), ('Obciążenia', 'obciazenia')]),
    'faktury': _lista('Faktury sprzedażowe', FAKTURY, zapytania.faktury, suma=[('Netto', 'netto'), ('Brutto', 'brutto')]),
    'naleznosci': _naleznosci,
    'zakupy': _lista('Rejestr zakupów', ZAKUPY, zapytania.zakupy, suma=[('Netto', 'netto')]),
    'wplaty': _lista('Wpłaty do połączenia', WPLATY, zapytania.wplaty, nazwa_pliku='wpłaty', suma=[('Suma', 'suma')]),
    'do-weryfikacji': _lista('Do weryfikacji', WERYFIKACJA, zapytania.do_weryfikacji),
    'wyniki-biura': _lista('Wyniki z biura', WYNIKI_BIURA, zapytania.wyniki_biura),
    'wydatki': _lista('Wydatki z banku', WYDATKI, zapytania.wydatki, suma=[('Suma', 'suma')]),
    'nieprzypisane-grupy': _lista('Nieprzypisane wg kontrahenta', NIEPRZYPISANE, zapytania.nieprzypisane_grupy),
    'konta-reklamowe': _lista('Konta reklamowe', KONTA, zapytania.konta_reklamowe),
    'budzet/rozjazdy': _lista('Rozjazdy z Budget Studio', ROZJAZDY, zapytania_budzet.rozjazdy),
    'klienci': _klienci,
    'klient': _klient,
    'wynik': _wynik,
    'obroty-klientow': _obroty,
    'koszty-kategorii': _koszty,
    'vat': _vat,
    'rentownosc/klienci': _rent_klienci,
    'rentownosc/arkusze': _rent_arkusze,
    'cashflow/okresy': _cashflow,
    'budzet/arkusz': _arkusz,
    # adresy sprzed pakietu 10 (przyciski w arkuszu, cash flow i rentowności)
    'arkusz': _arkusz,
    'cashflow': _cashflow_13,
    'rentownosc': _rentownosc_oba,
}


def plik(con, rodzaj: str, q: dict) -> tuple[str, bytes, str]:
    """Plik do pobrania: (nazwa z rozszerzeniem, treść, typ MIME). ValueError przy nieznanym eksporcie albo formacie."""
    fmt = q.get('format') or 'xlsx'
    if fmt not in FORMATY:
        raise ValueError('format: xlsx albo csv')
    fn = EKSPORTY.get(rodzaj)
    if not fn:
        raise ValueError('nieznany eksport')
    nazwa, tabele, bs = fn(con, q)
    if fmt == 'csv':
        return f'{nazwa}.csv', _csv(tabele[0], bs=bs), CSV
    return f'{nazwa}.xlsx', _xlsx(tabele), XLSX


# ---------- funkcje sprzed pakietu 10 (testy i wywołania z kodu) ----------

def arkusz(con, aid: str) -> tuple[str, bytes]:
    nazwa, tabele, _ = _arkusz(con, {'id': aid})
    return f'{nazwa}.xlsx', _xlsx(tabele)


def cash_flow(con) -> tuple[str, bytes]:
    nazwa, tabele, _ = _cashflow_13(con, {})
    return f'{nazwa}.xlsx', _xlsx(tabele)


def rentownosc_xls(con, rok=None) -> tuple[str, bytes]:
    nazwa, tabele, _ = _rentownosc_oba(con, {'rok': rok} if rok else {})
    return f'{nazwa}.xlsx', _xlsx(tabele)
