"""Podsumowanie importu jak w Budget Studio (B-IMP-07, B-LIS-14): każdy plik w Dane/ to jeden import, a jego wiersze
dzielą się na Dodane, Zaktualizowane (ten sam rekord, inne dane), Bez zmian i Pominięte (z powodem).

Liczone przy otwarciu Historii importów, nie przy odbudowie bazy (odbudowa się nie wydłuża), z samych plików, więc
każdy Mac dostaje ten sam wynik. Pliki źródła idą w kolejności, w jakiej trafiły do folderu (data pliku, potem nazwa),
a każdy jest porównany ze stanem sprzed niego:
  - ING, rejestr zakupów, wyniki biura (apka czyta wszystkie pliki): ze wszystkimi wcześniejszymi plikami razem,
  - Fakturownia, Budget Studio (apka czyta tylko najnowszy, bo każdy plik to pełny stan): z poprzednim plikiem.
Klucz rekordu jak w bazie: ING numer referencyjny + rachunek (bez numeru: dzień, kwota, tytuł, rachunek), faktura numer,
rejestr numer ewidencyjny, wyniki okres + wiersz, Budget Studio ID (kolejne wystąpienie tego samego ID: #2, #3).
Pominięte: wiersz, którego baza nie przyjmie (drugi taki sam klucz w pliku, rejestr bez numeru ewidencyjnego, BS bez
rodzaju 1 albo -1) albo cały plik, gdy nie da się go wczytać.
"""
import csv
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from .. import konfig
from ..format import zl
from . import dziennik, fakturownia, ing, rejestr_biura, wyniki_biura
from .budget_studio import WYMAGANE as WYMAGANE_BS
from .wspolne import data_z_nazwy, pliki, sha256

KATEGORIE = ('pominiete', 'dodane', 'zaktualizowane', 'bez_zmian')
KWOTY = {'kwota', 'kwota_oryg', 'saldo_po', 'netto', 'vat', 'brutto', 'oplacono', 'wartosc', 'wartosc_poprzednia'}
ETYKIETY = {
    'data_ksiegowania': 'data księgowania', 'data_transakcji': 'data transakcji', 'kwota_oryg': 'kwota oryginalna',
    'waluta_oryg': 'waluta oryginalna', 'saldo_po': 'saldo po', 'rachunek_kontrahenta': 'rachunek kontrahenta', 'tytul': 'tytuł',
    'nr_ref': 'numer referencyjny', 'data_wystawienia': 'data wystawienia', 'data_sprzedazy': 'data sprzedaży', 'termin': 'termin płatności',
    'nip': 'NIP', 'vat': 'VAT', 'oplacono': 'kwota opłacona', 'ksef': 'numer KSeF', 'oryginal': 'faktura pierwotna',
    'przyczyna_korekty': 'przyczyna korekty', 'nr_dok': 'numer dokumentu', 'data_zaewid': 'data zaewidencjonowania',
    'podmiot': 'kontrahent', 'podmiot_kod': 'kod kontrahenta', 'wartosc': 'wartość', 'wartosc_poprzednia': 'poprzedni rok',
}


@dataclass
class Rekord:
    klucz: str
    wiersz: int | None
    opis: str
    pola: dict


@dataclass
class Wczytanie:
    rekordy: dict = field(default_factory=dict)   # klucz → Rekord
    pominiete: list = field(default_factory=list)  # {'wiersz', 'opis', 'powod'}
    blad: str | None = None

    @property
    def wierszy(self) -> int:
        return len(self.rekordy) + len(self.pominiete)


def _krotko(s, n=60) -> str:
    s = ' '.join(str(s or '').split())
    return s if len(s) <= n else s[:n - 1] + '…'


def _opis(*czesci) -> str:
    return ' · '.join(str(c) for c in czesci if c not in (None, ''))


# ---------- odczyt rekordów z pliku, źródło po źródle ----------

POLA_ING = ('data_ksiegowania', 'data_transakcji', 'rachunek', 'kwota', 'waluta', 'kwota_oryg', 'waluta_oryg', 'saldo_po', 'kontrahent',
            'rachunek_kontrahenta', 'opis', 'tytul', 'nr_ref')


def _ing(p: Path, okres=None) -> Wczytanie:
    w = Wczytanie()
    for o in ing.wczytaj(p):
        r = Rekord(o['klucz'], o['wiersz'], _opis(o['data_ksiegowania'], zl(o['kwota']), _krotko(o['kontrahent'], 40), _krotko(o['tytul'])),
                   {k: o[k] for k in POLA_ING})
        if r.klucz in w.rekordy:   # baza przyjmie tylko pierwszą (ten sam klucz)
            w.pominiete.append({'wiersz': r.wiersz, 'opis': r.opis, 'powod': (
                f'taka sama operacja jak w wierszu {w.rekordy[r.klucz].wiersz} (ten sam numer referencyjny i rachunek, a gdy numeru brak: '
                'dzień, kwota, tytuł i rachunek); apka liczy ją raz')})
            continue
        w.rekordy[r.klucz] = r
    return w


POLA_FAKTURY = ('typ', 'status', 'data_wystawienia', 'data_sprzedazy', 'termin', 'nabywca', 'nip', 'netto', 'vat', 'brutto', 'oplacono',
                'waluta', 'ksef', 'oryginal', 'przyczyna_korekty', 'uwagi')


def _fakturownia(p: Path, okres=None) -> Wczytanie:
    faktury, pozycje = fakturownia.wczytaj(p)
    poz = defaultdict(list)
    for x in pozycje:
        poz[x['numer']].append(f"{x['nazwa']} {zl(x['netto'])}")
    w = Wczytanie()
    for f in faktury:
        pola = {k: f.get(k, '') for k in POLA_FAKTURY}
        pola['pozycje'] = '; '.join(poz.get(f['numer'], []))
        w.rekordy[f['numer']] = Rekord(f['numer'], f['wiersz'], _opis(f['numer'], _krotko(f['nabywca'], 40), zl(f['brutto']), f['status']), pola)
    # ten sam numer w kilku wierszach: apka bierze ostatni, wcześniejsze są pominięte
    with open(p, encoding='utf-8-sig', newline='') as plik:
        wiersze = list(csv.reader(plik))
    h = wiersze[0] if wiersze else []
    i_num, i_typ = h.index('numer'), h.index('Typ')
    ostatni = {n: f['wiersz'] for n, f in ((f['numer'], f) for f in faktury)}
    for nr, x in enumerate(wiersze[1:], start=2):
        if len(x) > max(i_num, i_typ) and x[i_typ].strip() and ostatni.get(x[i_num].strip()) not in (None, nr):
            w.pominiete.append({'wiersz': nr, 'opis': x[i_num].strip(),
                                'powod': f'ten sam numer faktury jest też w wierszu {ostatni[x[i_num].strip()]}; apka bierze tamten (ostatni)'})
    return w


POLA_REJESTRU = ('rejestr', 'nr_dok', 'data', 'data_zaewid', 'podmiot', 'podmiot_kod', 'netto', 'vat', 'brutto')


def _rejestr(p: Path, okres=None) -> Wczytanie:
    w = Wczytanie()
    for z in rejestr_biura.wczytaj(p):
        r = Rekord(z['nr_ewid'], z['wiersz'], _opis(z['nr_ewid'], z['rejestr'], _krotko(z['podmiot'], 40), z['nr_dok'], f"netto {zl(z['netto'])}"),
                   {k: z.get(k) for k in POLA_REJESTRU})
        if r.klucz in w.rekordy:
            w.pominiete.append({'wiersz': r.wiersz, 'opis': r.opis,
                                'powod': f'numer ewidencyjny powtórzony (pierwszy w arkuszu {w.rekordy[r.klucz].pola["rejestr"]}, wiersz '
                                         f'{w.rekordy[r.klucz].wiersz}); apka bierze pierwszy'})
            continue
        w.rekordy[r.klucz] = r
    import openpyxl   # wiersze bez numeru ewidencyjnego, w których coś jest (sumy na końcu arkusza pomijamy bez słowa)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        for arkusz in ('krajowe', 'zagraniczne'):
            if arkusz not in wb.sheetnames:
                continue
            wiersze = wb[arkusz].iter_rows(values_only=True)
            naglowek = [str(c).strip() if c is not None else '' for c in next(wiersze, ())]
            if 'Nr ewid.' not in naglowek:
                continue
            i = naglowek.index('Nr ewid.')
            for nr, x in enumerate(wiersze, start=2):
                if (x[i] if i < len(x) else None) not in (None, '') and str(x[i]).strip():
                    continue
                teksty = [str(c).strip().lower() for c in x if isinstance(c, str) and str(c).strip()]
                if not teksty or all(t.startswith(('suma', 'razem', 'ogółem')) for t in teksty):
                    continue   # pusty wiersz albo suma arkusza
                w.pominiete.append({'wiersz': nr, 'opis': _opis(arkusz, f'wiersz {nr}', _krotko(' '.join(teksty), 50)),
                                    'powod': 'brak numeru ewidencyjnego, a w wierszu są dane; apka bierze tylko wiersze z numerem'})
    finally:
        wb.close()
    return w


def _wyniki(p: Path, okres=None) -> Wczytanie:
    okres = okres or data_z_nazwy(p)
    w = Wczytanie()
    for x in wyniki_biura.wczytaj(p):
        k = f"{okres}|{x['zestawienie']}|{x['wiersz']}"
        w.rekordy[k] = Rekord(k, x['wiersz'], _opis(okres, x['kod'], _krotko(x['opis']), zl(x['wartosc'])),
                              {'kod': x['kod'], 'opis': x['opis'], 'wartosc': x['wartosc'], 'wartosc_poprzednia': x['wartosc_poprzednia']})
    return w


def _budget_studio(p: Path, okres=None) -> Wczytanie:
    w = Wczytanie()
    ile = Counter()
    with open(p, encoding='utf-8-sig', newline='') as f:
        czytnik = csv.DictReader(f, delimiter='\t')
        brak = [k for k in WYMAGANE_BS if k not in (czytnik.fieldnames or [])]
        if brak:
            raise ValueError(f'brak kolumn {brak}')
        for nr, x in enumerate(czytnik, start=2):
            pola = {k: (v or '').strip() for k, v in x.items() if k}
            opis = _opis(f"ID {pola.get('ID', '')}", _krotko(pola.get('Kontrahent'), 40), pola.get('Netto (wartość)'),
                         (pola.get('Miesiąc budżetowy') or '')[:7], pola.get('Etap płatności'))
            if pola.get('Rodzaj') not in ('1', '-1'):
                w.pominiete.append({'wiersz': nr, 'opis': opis, 'powod': 'rodzaj inny niż 1 (wpływ) albo -1 (obciążenie); apka go nie czyta'})
                continue
            ile[pola['ID']] += 1
            k = pola['ID'] if ile[pola['ID']] == 1 else f"{pola['ID']}#{ile[pola['ID']]}"
            w.rekordy[k] = Rekord(k, nr, opis, pola)
    return w


CZYTNIKI = {'ING': _ing, 'Fakturownia': _fakturownia, 'Rejestr zakupów': _rejestr, 'Wyniki biura': _wyniki, 'Budget Studio': _budget_studio}
OPISY = {'ING': 'historia operacji ING', 'Fakturownia': 'faktury sprzedaży z Fakturowni', 'Rejestr zakupów': 'rejestr zakupów z biura',
         'Wyniki biura': 'wyniki (RZiS) z biura', 'Budget Studio': 'eksport płatności z Budget Studio'}


def wczytaj(rodzaj: str, p: Path, okres: str | None = None) -> Wczytanie:
    """Rekordy pliku do porównania. Plik, którego nie da się wczytać: jeden wpis w Pominiętych z powodem, bez wyjątku."""
    try:
        return CZYTNIKI[rodzaj](p, okres)
    except Exception as e:   # zły plik nie psuje historii: pokazujemy powód
        return Wczytanie(pominiete=[{'wiersz': None, 'opis': p.name, 'powod': f'nie da się wczytać pliku: {e}'}], blad=str(e))


# ---------- porównanie ze stanem sprzed importu ----------

def _wartosc(k, v) -> str:
    if k in KWOTY and isinstance(v, int):
        return zl(v)
    return '' if v is None else str(v)


def roznice(stare: dict, nowe: dict) -> list[dict]:
    return [{'pole': ETYKIETY.get(k, k.replace('_', ' ')), 'bylo': _wartosc(k, stare.get(k)), 'jest': _wartosc(k, nowe.get(k))}
            for k in nowe if stare.get(k) != nowe.get(k)]


def porownaj(w: Wczytanie, stan: dict) -> dict:
    """Rekordy pliku kontra stan sprzed niego (klucz → Rekord): dodane, zaktualizowane (z listą zmian), bez zmian."""
    wynik = {'dodane': [], 'zaktualizowane': [], 'bez_zmian': []}
    for k, r in w.rekordy.items():
        s = stan.get(k)
        if s is None:
            wynik['dodane'].append(r)
        elif s.pola == r.pola:
            wynik['bez_zmian'].append(r)
        else:
            wynik['zaktualizowane'].append((r, roznice(s.pola, r.pola)))
    return wynik


_PLIKI = {}      # (rodzaj, sha, nazwa) → Wczytanie; pliki się nie zmieniają, więc wystarczy pamięć procesu
_PRZEBIEGI = {}  # rodzaj → (odcisk plików, przebieg)


def _wczytaj_z_pamieci(rodzaj, p: Path, sha: str) -> Wczytanie:
    k = (rodzaj, sha, p.name)
    if k not in _PLIKI:
        if len(_PLIKI) > 60:
            _PLIKI.clear()
        _PLIKI[k] = wczytaj(rodzaj, p)
    return _PLIKI[k]


def _zrodla():
    from ..odbudowa import ZRODLA   # (rodzaj, wzorzec względem Dane/, moduł, wszystkie pliki czy tylko najnowszy)
    return [(r, wz, wsz) for r, wz, _, wsz in ZRODLA]


def przebieg(rodzaj: str, wzorzec: str, wszystkie: bool) -> dict:
    """Pliki źródła w kolejności trafiania do folderu, każdy z podsumowaniem; plus stan po ostatnim (do podglądu nowego
    pliku): przy „wszystkie” suma wszystkich plików, przy „tylko najnowszy” plik, którego apka używa."""
    dane = konfig.katalog_dane()
    lista = sorted((p for p in dane.glob(wzorzec) if p.is_file()), key=lambda p: (p.stat().st_mtime, p.name))
    info = [(p, sha256(p), p.stat().st_mtime) for p in lista]
    odcisk = (str(dane), *((p.name, h, m) for p, h, m in info))
    if rodzaj in _PRZEBIEGI and _PRZEBIEGI[rodzaj][0] == odcisk:
        return _PRZEBIEGI[rodzaj][1]
    uzywany = (pliki(dane, wzorzec) or [None])[0]
    stan, widziane, wynik, stan_uzywanego = {}, {}, [], {}
    for p, h, m in info:
        w = _wczytaj_z_pamieci(rodzaj, p, h)
        e = {'rodzaj': rodzaj, 'plik': p.name, 'rel': f'{p.parent.name}/{p.name}', 'sha': h, 'mtime': m, 'wczytanie': w}
        if w.blad:
            e.update(stan='błąd odczytu', dodane=[], zaktualizowane=[], bez_zmian=[])
        elif h in widziane:
            e.update(stan=f'kopia pliku {widziane[h]}, apka czyta go raz', dodane=[], zaktualizowane=[], bez_zmian=list(w.rekordy.values()))
        else:
            widziane[h] = p.name
            e.update(porownaj(w, stan))
            e['stan'] = 'wczytany' if wszystkie or p == uzywany else 'zastąpiony nowszym plikiem (apka czyta tylko najnowszy)'
            stan = {**stan, **w.rekordy} if wszystkie else dict(w.rekordy)
            if p == uzywany:
                stan_uzywanego = dict(w.rekordy)
        wynik.append(e)
    wynik_ = {'pliki': wynik, 'stan': stan if wszystkie else stan_uzywanego}
    _PRZEBIEGI[rodzaj] = (odcisk, wynik_)
    return wynik_


def stan_przed(rodzaj: str) -> dict:
    """Stan źródła przed nowym importem (do podglądu pliku przeciągniętego do okna)."""
    for r, wz, wsz in _zrodla():
        if r == rodzaj:
            return przebieg(r, wz, wsz)['stan']
    return {}


# ---------- Historia importów i podsumowanie pojedynczego importu ----------

def _czas_pliku(m: float) -> str:
    return datetime.fromtimestamp(m, timezone.utc).isoformat(timespec='seconds')


def _wiersz_historii(e: dict, wpis: dict | None, uwagi: str) -> dict:
    w = e['wczytanie']
    return {'id': f"plik:{e['rel']}", 'zrodlo': e['rodzaj'], 'opis_zrodla': OPISY.get(e['rodzaj'], ''), 'plik': e['plik'],
            'katalog': e['rel'].split('/')[0], 'sposob': 'przeciągnięty do okna' if wpis else 'folder Dane',
            'plik_zrodlowy': (wpis or {}).get('plik_zrodlowy', ''), 'czas': (wpis or {}).get('czas') or _czas_pliku(e['mtime']),
            'stan': e['stan'], 'dodane': len(e['dodane']), 'zaktualizowane': len(e['zaktualizowane']), 'bez_zmian': len(e['bez_zmian']),
            'pominiete': len(w.pominiete), 'wierszy': w.wierszy, 'uwagi': uwagi or '', 'blad': w.blad or ''}


def _wiersz_wpisu(wpis: dict) -> dict:
    """Import pliku, który już był w Dane/ (pliku nie zapisano drugi raz): liczniki z chwili importu."""
    l = wpis.get('liczniki') or {}
    return {'id': f"wpis:{wpis['id']}", 'zrodlo': wpis.get('zrodlo', ''), 'opis_zrodla': OPISY.get(wpis.get('zrodlo', ''), ''),
            'plik': wpis.get('plik_zrodlowy', ''), 'katalog': '', 'sposob': 'przeciągnięty do okna', 'plik_zrodlowy': wpis.get('plik_zrodlowy', ''),
            'czas': wpis.get('czas', ''), 'stan': f"ten sam plik co {wpis.get('duplikat_z', '')}, nie zapisano drugi raz",
            **{k: int(l.get(k, 0)) for k in KATEGORIE}, 'wierszy': int(wpis.get('wierszy') or 0), 'uwagi': '', 'blad': ''}


def _przebiegi() -> list[dict]:
    return [e for r, wz, wsz in _zrodla() for e in przebieg(r, wz, wsz)['pliki']]


def historia(con) -> dict:
    """Wszystkie importy (pliki w Dane/ i ponowne importy z dziennika), najnowsze na górze, z licznikami jak w BS."""
    wpisy = dziennik.wczytaj()
    po_pliku = {w['zapisany_jako']: w for w in wpisy if w.get('zapisany_jako')}
    uwagi, inne = {}, []
    for r in con.execute('SELECT rodzaj, plik, wierszy, uwagi FROM import_log ORDER BY id'):
        uwagi[(r['rodzaj'], r['plik'])] = r['uwagi'] or ''
        if r['rodzaj'] not in CZYTNIKI:
            inne.append(dict(r))
    importy = [_wiersz_historii(e, po_pliku.get(e['rel']), uwagi.get((e['rodzaj'], e['plik']), '')) for e in _przebiegi()]
    importy += [_wiersz_wpisu(w) for w in wpisy if not w.get('zapisany_jako')]
    importy.sort(key=lambda x: x['czas'], reverse=True)
    return {'importy': importy, 'inne': inne, 'katalog': 'Finanse/Raporty/Dane'}


def _pozycje(e: dict, kategoria: str) -> list[dict]:
    if kategoria == 'pominiete':
        return [dict(x) for x in e['wczytanie'].pominiete]
    if kategoria == 'zaktualizowane':
        return [{'wiersz': r.wiersz, 'opis': r.opis, 'zmiany': '; '.join(f"{z['pole']}: {z['bylo'] or '(puste)'} → {z['jest'] or '(puste)'}"
                                                                          for z in zm)} for r, zm in e['zaktualizowane']]
    return [{'wiersz': r.wiersz, 'opis': r.opis} for r in e[kategoria]]


def szczegoly(con, id_: str, kategoria: str | None = None, q: str | None = None, strona=1, ile=100) -> dict:
    """Podsumowanie jednego importu: nagłówek z licznikami i pozycje jednej kategorii (domyślnie Pominięte, gdy są)."""
    wpisy = dziennik.wczytaj()
    if id_.startswith('plik:'):
        e = next((x for x in _przebiegi() if x['rel'] == id_[5:]), None)
        if not e:
            raise ValueError('nie ma takiego importu (plik zniknął z folderu Dane?)')
        uw = con.execute('SELECT uwagi FROM import_log WHERE rodzaj = ? AND plik = ?', (e['rodzaj'], e['plik'])).fetchone()
        naglowek = _wiersz_historii(e, next((w for w in wpisy if w.get('zapisany_jako') == e['rel']), None), uw[0] if uw else '')
    elif id_.startswith('wpis:'):
        wpis = next((w for w in wpisy if w['id'] == id_[5:] and not w.get('zapisany_jako')), None)
        if not wpis:
            raise ValueError('nie ma takiego importu')
        naglowek = _wiersz_wpisu(wpis)
        oryg = next((x for x in _przebiegi() if x['rel'] == wpis.get('duplikat_z')), None)
        e = {'wczytanie': Wczytanie(), 'dodane': [], 'zaktualizowane': [],
             'bez_zmian': list(oryg['wczytanie'].rekordy.values()) if oryg else []}
    else:
        raise ValueError('nie ma takiego importu')
    if kategoria not in KATEGORIE:
        kategoria = next((k for k in KATEGORIE if naglowek[k]), 'dodane')
    pozycje = _pozycje(e, kategoria)
    if q:
        f = q.lower()
        pozycje = [x for x in pozycje if f in ' '.join(str(v) for v in x.values()).lower()]
    ile = max(1, min(int(ile or 100), 500))
    strona = max(1, int(strona or 1))
    return {'import': naglowek, 'kategoria': kategoria, 'razem': len(pozycje), 'strona': strona, 'ile': ile,
            'wiersze': pozycje[(strona - 1) * ile:strona * ile]}


def api(con, sciezka: str, q: dict):
    """GET /api/importy/historia i /api/importy/podsumowanie?id=&kategoria=&q=&strona=&ile=."""
    if sciezka == 'historia':
        return historia(con)
    if sciezka == 'podsumowanie':
        return szczegoly(con, q.get('id', ''), q.get('kategoria'), q.get('q'), q.get('strona'), q.get('ile'))
    raise ValueError('nieznany adres')
