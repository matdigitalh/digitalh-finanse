"""Rejestr VAT zakupów z biura (xlsx, arkusze „krajowe” i „zagraniczne”). Pełna lista kosztów fakturowych.

Podmiot ma sufiks kontrahenta „(00024)”, który odcinamy i trzymamy osobno. Wiersze bez numeru
ewidencyjnego (puste i sumy na końcu) pomijamy. Kolejne rejestry mogą się nakładać: importujemy
wszystkie od najnowszego, klucz to numer ewidencyjny.
"""
import re
from pathlib import Path

from ..format import data_iso, grosze
from .wspolne import zapisz_log

SUFIKS = re.compile(r'\s*\((\d+)\)\s*$')


def wczytaj(p: Path) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    wynik = []
    try:
        for rejestr in ('krajowe', 'zagraniczne'):
            if rejestr not in wb.sheetnames:
                continue
            wiersze = wb[rejestr].iter_rows(values_only=True)
            naglowek = [str(c).strip() if c is not None else '' for c in next(wiersze)]
            i = {k: n for n, k in enumerate(naglowek)}
            for nr, w in enumerate(wiersze, start=2):
                g = lambda k: w[i[k]] if k in i and i[k] < len(w) else None
                nr_ewid = g('Nr ewid.')
                if not nr_ewid or not str(nr_ewid).strip():
                    continue
                podmiot = str(g('Podmiot') or '').strip()
                m = SUFIKS.search(podmiot)
                if rejestr == 'krajowe':
                    netto = grosze(g('8% Netto') or 0) + grosze(g('23% Netto') or 0)
                    vat = grosze(g('8% VAT') or 0) + grosze(g('23% VAT') or 0)
                    brutto = grosze(g('Brutto') or 0)
                else:
                    netto, vat, brutto = grosze(g('Netto') or 0), 0, None
                wynik.append({'nr_ewid': str(nr_ewid).strip(), 'nr_dok': str(g('Nr dok.') or '').strip(),
                              'data': data_iso(g('Data')), 'data_zaewid': data_iso(g('Data zaewid.')),
                              'podmiot': SUFIKS.sub('', podmiot), 'podmiot_kod': m.group(1) if m else '',
                              'rejestr': rejestr, 'netto': netto, 'vat': vat, 'brutto': brutto, 'wiersz': nr})
    finally:
        wb.close()
    return wynik


def importuj(con, p: Path) -> int:
    zakupy = wczytaj(p)
    zrodlo = f'{p.parent.name}/{p.name}'
    nowych = 0
    for z in zakupy:
        cur = con.execute('INSERT OR IGNORE INTO zakupy (nr_ewid, nr_dok, data, data_zaewid, podmiot, podmiot_kod, rejestr,'
                          ' netto, vat, brutto, plik, wiersz) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
                          (z['nr_ewid'], z['nr_dok'], z['data'], z['data_zaewid'], z['podmiot'], z['podmiot_kod'],
                           z['rejestr'], z['netto'], z['vat'], z['brutto'], zrodlo, z['wiersz']))
        nowych += cur.rowcount
    zapisz_log(con, 'Rejestr zakupów', p, len(zakupy), nowych)
    return nowych
