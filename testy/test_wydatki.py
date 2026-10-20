"""Analiza wydatków na zmyślonych danych: kategorie wbudowane, pierwszeństwo decyzji Mata, cel i cofnij, koszty stałe,
kafle bez podwójnego liczenia, ostrzeżenie o kategorii niezgodnej z rodzajem, kategoria nie zmienia rodzaju ani wyniku."""
import io
import unittest

import openpyxl

from finanse import baza, eksport, konfig, odbudowa, wydatki, wynik, zmiany
from finanse.wydatki import kategoria_wbudowana, klucz_grupy
from testy.pomoc import ONLINE, Srodowisko, operacja, zapisz_ing

KARTA = 'Płatność kartą'


def karta(data, kwota, kontrahent, ref):
    return operacja(data, kwota, '0', kontrahent, KARTA, rachunek=ONLINE, ref=ref, opis=KARTA)


def dane_testowe(dane):
    zapisz_ing(dane, 'ing_historia_2025-09-01_2026-10-15.csv', [
        operacja('2026-10-15', '-100,00', '0', 'URZĄD SKARBOWY', '/OKR/26K03/CIT-8', ref='T4'),
        operacja('2026-09-25', '-1500,00', '0', 'URZĄD SKARBOWY', '/OKR/26M08/VAT-7', ref='T1'),
        operacja('2026-09-20', '-300,00', '0', 'URZĄD SKARBOWY', '/OKR/26M08/PIT-4', ref='T2'),
        operacja('2026-09-15', '-600,00', '0', 'ZUS', 'DRA 08/2026', ref='Z3'),
        operacja('2026-07-15', '-800,00', '0', 'ZUS', 'DRA 06/2026', ref='Z2'),
        operacja('2026-06-15', '-800,00', '0', 'ZUS', 'DRA 05/2026', ref='Z1'),
        operacja('2026-09-15', '-60,00', '0', 'PKO EMERYTURA SFIO', 'SKŁADKI PPK', ref='E1'),
        operacja('2026-09-10', '-5000,00', '0', 'KOWALSKA ANNA UL. TESTOWA', 'WYNAGRODZENIE ZA 08.2026', ref='P1'),
        operacja('2026-09-08', '-2000,00', '0', 'NOWAK PIOTR', 'DYWIDENDA ZA 2025', ref='D1'),
        operacja('2026-09-05', '-200,00', '0', 'KLIENT ZWROT SP. Z O.O.', 'Zwrot nadpłaty do faktury FV1', ref='ZW1'),
        operacja('2026-09-04', '-1230,00', '0', 'Drukomat.pl Pila 64-920', 'Faktura 1/09', ref='DR1'),
        karta('2026-09-29', '-50,00', 'FIGMA SAN FRANCISCO 94102', 'F4'),
        karta('2026-08-29', '-41,00', 'FIGMA SAN FRANCISCO 94102', 'F3'),
        karta('2026-07-29', '-41,00', 'FIGMA SAN FRANCISCO 94102', 'F2'),
        karta('2026-06-29', '-41,00', 'FIGMA SAN FRANCISCO 94102', 'F1'),
        karta('2026-09-12', '-310,00', 'nazwa.pl Krakow 30-727', 'N2'),
        karta('2025-09-10', '-300,00', 'nazwa.pl Krakow 30-727', 'N1'),
        karta('2026-08-15', '-60,00', 'GOOGLE *ADS1112223333 DUBLIN', 'G1'),
        karta('2026-08-10', '-25,00', 'BOLT.EU/S/2610160930 Warsaw', 'B1'),
        karta('2026-08-11', '-30,00', 'BOLT.EU/S/9999999999 Warsaw', 'B2'),
        karta('2026-08-02', '-99,00', 'NIEZNANY SKLEP', 'X1'),
        # poprawki po kontroli liczb: opłata roczna, wygasły, flagi, kilka kart w miesiącu, zwroty kosztów obok pensji
        karta('2026-09-28', '-4678,90', 'Adobe DUBLIN D24', 'AD1'),
        karta('2026-08-20', '-12,00', 'TRELLO.COM* ATLASSIAN', 'TR3'),
        karta('2026-07-20', '-12,00', 'TRELLO.COM* ATLASSIAN', 'TR2'),
        karta('2026-06-20', '-12,00', 'TRELLO.COM* ATLASSIAN', 'TR1'),
        karta('2026-10-15', '-282,78', 'WISPR SAN FRANCISCO', 'WI5'),
        karta('2026-09-20', '-30,00', 'WISPR SAN FRANCISCO', 'WI4'),
        karta('2026-08-20', '-30,00', 'WISPR SAN FRANCISCO', 'WI3'),
        karta('2026-07-20', '-30,00', 'WISPR SAN FRANCISCO', 'WI2'),
        karta('2026-06-20', '-30,00', 'WISPR SAN FRANCISCO', 'WI1'),
        karta('2026-10-02', '-92,99', 'Google Workspace_digit Warsaw', 'GW5'),
        karta('2026-10-02', '-464,94', 'Google Workspace_digit Warsaw', 'GW6'),
        karta('2026-10-02', '-46,49', 'Google Workspace_digit Warsaw', 'GW7'),
        karta('2026-09-02', '-92,99', 'Google Workspace_digit Warsaw', 'GW3'),
        karta('2026-09-02', '-464,94', 'Google Workspace_digit Warsaw', 'GW4'),
        karta('2026-08-02', '-92,99', 'Google Workspace_digit Warsaw', 'GW2'),
        karta('2026-07-02', '-92,99', 'Google Workspace_digit Warsaw', 'GW1'),
        operacja('2026-09-03', '-9367,93', '0', 'HEROLD MATEUSZ UL. TESTOWA 1', 'WYNAGRODZENIE ZA 08.2026 ZARZĄD', ref='H4'),
        operacja('2026-08-06', '-9367,93', '0', 'HEROLD MATEUSZ UL. TESTOWA 1', 'WYNAGRODZENIE ZA 07.2026 ZARZĄD', ref='H3'),
        operacja('2026-07-07', '-9367,93', '0', 'HEROLD MATEUSZ UL. TESTOWA 1', 'WYNAGRODZENIE ZA 06.2026 ZARZĄD', ref='H2'),
        operacja('2026-06-05', '-9367,93', '0', 'HEROLD MATEUSZ UL. TESTOWA 1', 'WYNAGRODZENIE ZA 05.2026 ZARZĄD', ref='H1'),
        operacja('2026-09-01', '-2388,55', '0', 'Mateusz Herold', 'Zwrot za wyjazdy służbowe w sierpniu', ref='HZ3'),
        operacja('2026-08-01', '-581,90', '0', 'Mateusz Herold', 'Zwrot za wyjazdy służbowe w lipcu', ref='HZ2'),
        operacja('2026-07-13', '-410,55', '0', 'Mateusz Herold', 'Zwrot za wyjazdy służbowe w czerwcu', ref='HZ1'),
    ])


def wyd(con, ref):
    return con.execute('SELECT w.* FROM wydatki w JOIN operacje o ON o.id = w.operacja_id WHERE o.nr_ref = ?', (ref,)).fetchone()


def klucz_op(con, ref):
    return con.execute('SELECT klucz FROM operacje WHERE nr_ref = ?', (ref,)).fetchone()[0]


class TestRegulyWbudowane(unittest.TestCase):
    def k(self, **o):
        return kategoria_wbudowana(o)[0]

    def test_po_rodzaju_i_nazwie(self):
        self.assertEqual(self.k(rodzaj='vat', kontrahent='Urząd Skarbowy'), 'vat')
        self.assertEqual(self.k(rodzaj='dywidenda', kontrahent='X'), 'podzial_zysku')
        self.assertEqual(self.k(rodzaj='zwrot', kontrahent='X'), 'zwroty')
        self.assertEqual(self.k(rodzaj='zwrocony', kontrahent='Allegro'), 'zwroty')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='ZUS'), 'zus')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='Zakład Ubezpieczeń Społecznych'), 'zus')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='Urząd Skarbowy', tytul='/OKR/26M09/PIT-4'), 'pit4')
        self.assertEqual(self.k(rodzaj='podatki', kontrahent='Urząd Skarbowy', tytul='CIT-8'), 'cit')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='PKO EMERYTURA- SFIO'), 'ppk')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='KOWALSKA ANNA', tytul='WYNAGRODZENIE ZA 08.2026'), 'pensje')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='Jan Nowak', tytul='Faktura 3/2026'), 'wspolpracownicy')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='KRÓL ZUZANNA', tytul='Rach. umowa o dzieło'), 'wspolpracownicy')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='KRÓL ZUZANNA', tytul='Rach. umowa o dzieło', kk='krol zuzanna'),
                         'wspolpracownicy')
        self.assertEqual(kategoria_wbudowana({'rodzaj': 'praca', 'kontrahent': 'Jan Nowak', 'tytul': 'przelew', 'kk': 'jan nowak'},
                                             {'jan nowak'})[0], 'pensje')   # osoba z listy płac
        self.assertEqual(wydatki.NAZWY['pensje'], 'Pensje (etat, powołanie)')
        self.assertIn('wspolpracownicy', wydatki.KOSZT_PRACY)
        self.assertEqual(self.k(rodzaj='praca', kontrahent='Mateusz Herold', tytul='Zwrot za wyjazdy służbowe w maju'), 'podroze')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='MedicCentre Sp.z.o.o'), 'benefity')
        for n in ('Google Workspace_digit Warsaw', 'FIGMA SAN FRANCISCO', 'BUDGET STUDIO WARSZAWA', 'ANTHROPIC* CLAUDE SUB', 'OPENAI *CHATGPT',
                  'CLICKUP 888', 'IPR* VXTGK TORONTO', 'Microsoft-G186021732 msbill.info', 'nazwa.pl Krakow', 'fakturownia.pl Warszawa'):
            self.assertEqual(self.k(rodzaj='staly', kontrahent=n, karta=True), 'narzedzia', n)
        self.assertEqual(self.k(rodzaj='staly', kontrahent='BRIXEN-FK Magdalena'), 'ksiegowosc')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='Kancelaria Radcy Prawnego'), 'ksiegowosc')
        self.assertEqual(self.k(rodzaj='nieprzypisany', kontrahent='BOLT.EU/S/261 Warsaw', karta=True), 'podroze')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='Booking.com Hotel Amsterdam', karta=True), 'podroze')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='MEDIAEXPERT PL ZLOTOW', karta=True), 'sprzet')
        self.assertEqual(self.k(rodzaj='staly', kontrahent='Allegro Poznan', karta=True), 'sprzet')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='RISTORANTE SEMOLINO', karta=True), 'reprezentacja')
        self.assertEqual(self.k(rodzaj='staly', kontrahent='JUST JOIN IT Sp. z o. o.'), 'rekrutacja')

    def test_koszty_klientow_i_reklama(self):
        self.assertEqual(self.k(rodzaj='przenoszony', kontrahent='GOOGLE *ADS1112223333 DUBLIN', karta=True), 'klienci_media')
        self.assertEqual(self.k(rodzaj='przenoszony', kontrahent='FACEBK *W9NGQGHVA2'), 'klienci_media')
        self.assertEqual(self.k(rodzaj='przenoszony', kontrahent='Drukomat.pl Pila'), 'klienci_druk')
        self.assertEqual(self.k(rodzaj='przenoszony', kontrahent='Studio Grafiki'), 'klienci_podwykonawcy')
        self.assertEqual(self.k(rodzaj='staly', kontrahent='FACEBK *ABC'), 'reklama')
        self.assertEqual(self.k(rodzaj='nieprzypisany', kontrahent='FACEBK *ABC', konto='Meta DigitalH'), 'reklama')
        self.assertEqual(self.k(rodzaj='nieprzypisany', kontrahent='FACEBK *ABC', konto='Meta ROPA'), 'klienci_media')

    def test_domyslne_dla_rodzaju(self):
        self.assertEqual(self.k(rodzaj='nieprzypisany', kontrahent='NIEZNANY SKLEP', karta=True), 'nieprzypisane')
        self.assertEqual(self.k(rodzaj='staly', kontrahent='NIEZNANY SKLEP', karta=True), 'inne')
        self.assertEqual(self.k(rodzaj='praca', kontrahent='NIEZNANY SKLEP SP. Z O.O.', karta=True), 'inne')
        self.assertEqual(self.k(rodzaj='podatki', kontrahent='Gmina Testowa'), 'cit')

    def test_grupa_kontrahenta_kart(self):
        self.assertEqual(klucz_grupy('FACEBK *W9NGQGHVA2 fb.me/ads', 'facebk wngqghva'), 'facebk')
        self.assertEqual(klucz_grupy('FACEBK *ACY5FFMVA2 fb.me/ads', 'acyffmva facebk'), 'facebk')
        self.assertEqual(klucz_grupy('IPR* VXTGK TORONTO', 'ipr'), 'iproyal')
        self.assertEqual(klucz_grupy('Zakład Ubezpieczeń Społecznych (ZUS)', 'ubezpieczen zaklad'), 'zus')
        self.assertEqual(klucz_grupy('FIGMA SAN FRANCISCO', 'figma san'), 'figma san')

    def test_zgodnosc(self):
        self.assertFalse(wydatki.zgodna('narzedzia', 'praca'))
        self.assertTrue(wydatki.zgodna('narzedzia', 'staly'))
        self.assertTrue(wydatki.zgodna('inne', 'praca'))
        self.assertFalse(wydatki.zgodna('pensje', 'nieprzypisany'))


class TestWydatki(unittest.TestCase):
    def setUp(self):
        self.env = Srodowisko().__enter__()
        dane_testowe(self.env.dane)
        self.przebuduj()

    def tearDown(self):
        self.con.close()
        self.env.__exit__()

    def przebuduj(self):
        if getattr(self, 'con', None):
            self.con.close()
        self.raport = odbudowa.odbuduj()
        self.con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)

    def wpis(self, typ, dane):
        return zmiany.zapisz(typ, zmiany.sprawdz(typ, dane))['id']

    def test_kategorie_po_odbudowie(self):
        self.assertEqual([b for b in self.raport['bledy'] if not b.startswith('Brak plików')], [])
        oczekiwane = {'T1': 'vat', 'T2': 'pit4', 'T4': 'cit', 'Z1': 'zus', 'E1': 'ppk', 'P1': 'pensje', 'D1': 'podzial_zysku', 'ZW1': 'zwroty',
                      'F1': 'narzedzia', 'N1': 'narzedzia', 'B1': 'podroze', 'X1': 'nieprzypisane', 'G1': 'klienci_media'}
        for ref, kat in oczekiwane.items():
            w = wyd(self.con, ref)
            self.assertEqual(w['kategoria'], kat, ref)
            self.assertTrue(w['kategoria_zrodlo'].startswith('reguła wbudowana'), ref)
        self.assertEqual(wyd(self.con, 'B1')['kontrahent_grupa'], wyd(self.con, 'B2')['kontrahent_grupa'])   # Bolt: jeden kontrahent

    def test_pierwszenstwo_decyzji_mata_i_rodzaj_bez_zmian(self):
        rodzaje = {r['operacja_id']: r['rodzaj'] for r in self.con.execute('SELECT operacja_id, rodzaj FROM wydatki')}
        wynik_przed = wynik.wynik_zarzadczy(self.con, '2026')
        grupa = wyd(self.con, 'F1')['kontrahent_grupa']
        self.wpis('kategoria_regula', {'kontrahent': grupa, 'kategoria': 'inne'})
        self.wpis('kategoria_operacja', {'operacja': klucz_op(self.con, 'F2'), 'kategoria': 'sprzet'})
        self.przebuduj()
        self.assertEqual(wyd(self.con, 'F1')['kategoria'], 'inne')
        self.assertEqual(wyd(self.con, 'F1')['kategoria_zrodlo'], 'reguła Mata (kontrahent)')
        self.assertEqual(wyd(self.con, 'F2')['kategoria'], 'sprzet')   # płatność wygrywa z regułą kontrahenta
        self.assertEqual(wyd(self.con, 'F2')['kategoria_zrodlo'], 'decyzja Mata (ta płatność)')
        self.assertEqual({r['operacja_id']: r['rodzaj'] for r in self.con.execute('SELECT operacja_id, rodzaj FROM wydatki')}, rodzaje)
        self.assertEqual(wynik.wynik_zarzadczy(self.con, '2026'), wynik_przed)

    def test_cel_czesciowy_i_cofnij(self):
        g = wyd(self.con, 'F1')['kontrahent_grupa']
        self.wpis('cel_wydatku', {'kontrahent': g, 'obszar': 'Zespół', 'kto': 'Anna', 'uzasadnienie': 'projekty graficzne'})
        w2 = self.wpis('cel_wydatku', {'kontrahent': g, 'ocena': 'do_wypowiedzenia'})
        self.przebuduj()
        c = dict(self.con.execute('SELECT * FROM wydatki_cel WHERE kontrahent = ?', (g,)).fetchone())
        self.assertEqual((c['obszar'], c['kto'], c['uzasadnienie'], c['ocena']), ('Zespół', 'Anna', 'projekty graficzne', 'do_wypowiedzenia'))
        self.assertRegex(c['data_oceny'], r'^\d{4}-\d{2}-\d{2}$')
        self.wpis('cofnij', {'wpis': w2})
        self.przebuduj()
        c = dict(self.con.execute('SELECT * FROM wydatki_cel WHERE kontrahent = ?', (g,)).fetchone())
        self.assertEqual((c['obszar'], c['ocena'], c['data_oceny']), ('Zespół', '', ''))

    def test_walidacja(self):
        with self.assertRaises(ValueError):
            zmiany.sprawdz('kategoria_regula', {'kontrahent': 'figma san', 'kategoria': 'nieznana'})
        with self.assertRaises(ValueError):
            zmiany.sprawdz('cel_wydatku', {'kontrahent': 'figma san', 'ocena': 'super'})
        with self.assertRaises(ValueError):
            zmiany.sprawdz('cel_wydatku', {'kontrahent': 'figma san'})   # brak zmian
        with self.assertRaises(ValueError):
            zmiany.sprawdz('cel_wydatku', {'kontrahent': 'figma san', 'staly': 'tak'})
        with self.assertRaises(ValueError):
            wydatki.sprawdz_odwolania(self.con, 'cel_wydatku', {'kontrahent': 'nie ma takiego', 'kto': 'x'})
        with self.assertRaises(ValueError):
            wydatki.sprawdz_odwolania(self.con, 'cel_wydatku', {'kontrahent': 'figma san', 'obszar': 'Szkolenia'})
        self.wpis('ustawienie', {'klucz': 'obszary_wydatkow', 'wartosc': ['Szkolenia']})
        self.przebuduj()
        wydatki.sprawdz_odwolania(self.con, 'cel_wydatku', {'kontrahent': 'figma san', 'obszar': 'Szkolenia'})
        self.assertIn('Szkolenia', wydatki.obszary(self.con))

    def stale(self):
        s = wydatki.koszty_stale(self.con)
        return s, {x['kontrahent'] + '|' + x['kategoria']: x for x in s['pozycje']}

    def grupa(self, ref):
        return wyd(self.con, ref)['kontrahent_grupa']

    def test_koszty_stale_kwota_z_ostatniego_pelnego_miesiaca(self):
        s, po = self.stale()
        self.assertEqual(s['okno'], ['2026-06', '2026-07', '2026-08', '2026-09'])
        self.assertEqual(po['zus|zus']['miesiecznie'], 60000)   # wrzesień, nie średnia 800, 800, 600
        figma = po[self.grupa('F1') + '|narzedzia']
        self.assertEqual((figma['miesiecznie'], figma['rocznie']), (5000, 60000))
        self.assertEqual(figma['flagi'], [])   # +22%: zmiana ceny pokazana, bez flagi (próg 50%)
        self.assertEqual(figma['zmiana_ceny']['proc'], 22.0)
        self.assertNotIn(self.grupa('X1') + '|nieprzypisane', po)   # jednorazowy
        self.assertFalse(any(x['kategoria'] in ('vat', 'klienci_media') for x in s['pozycje']))
        self.assertEqual(s['razem']['miesiecznie'], sum(x['miesiecznie'] for x in s['pozycje']))

    def test_wygasle(self):
        s, po = self.stale()
        trello = self.grupa('TR1')
        self.assertNotIn(trello + '|narzedzia', po)   # 3 z 4 miesięcy, ale nic we wrześniu ani w październiku
        w = next(x for x in s['wygasle'] if x['kontrahent'] == trello)
        self.assertEqual((w['ostatnia_data'], w['ostatnia_kwota']), ('2026-08-20', 1200))
        self.assertNotIn(trello, [x['kontrahent'] for x in s['pozycje']])
        self.assertEqual(wydatki.kontrahent(self.con, trello)['wygasly']['ostatnia_data'], '2026-08-20')

    def test_oplata_roczna(self):
        s, po = self.stale()
        adobe = po[self.grupa('AD1') + '|narzedzia']
        self.assertEqual((adobe['roczna'], adobe['rocznie'], adobe['miesiecznie']), (True, 467890, 38991))
        nazwa = po[self.grupa('N1') + '|narzedzia']   # jedna płatność od 300 zł w 12 miesiącach
        self.assertEqual((nazwa['roczna'], nazwa['miesiecznie']), (True, 2583))
        # ręcznie: sklep jako opłata roczna, Adobe jako miesięczna (wtedy 1 z 4 miesięcy, więc nie jest stały)
        self.wpis('cel_wydatku', {'kontrahent': self.grupa('X1'), 'roczna': True})
        self.wpis('cel_wydatku', {'kontrahent': self.grupa('AD1'), 'roczna': False})
        self.przebuduj()
        s, po = self.stale()
        self.assertEqual(po[self.grupa('X1') + '|nieprzypisane']['miesiecznie'], 825)
        self.assertIn('decyzja Mata', po[self.grupa('X1') + '|nieprzypisane']['sposob'])
        self.assertNotIn(self.grupa('AD1') + '|narzedzia', po)

    def test_flagi_zmiany_kwoty_i_poza_cyklem(self):
        s, po = self.stale()
        wispr = po[self.grupa('WI1') + '|narzedzia']
        self.assertEqual(wispr['miesiecznie'], 3000)
        self.assertEqual(set(wispr['flagi']), {'zmiana kwoty', 'poza cyklem'})
        self.assertEqual((wispr['ostatnia_data'], wispr['ostatnia_kwota']), ('2026-10-15', 28278))

    def test_ostatnia_platnosc_suma_miesiaca(self):
        s, po = self.stale()
        gw = po[self.grupa('GW1') + '|narzedzia']
        self.assertEqual((gw['ostatnia_data'], gw['ostatnia_kwota'], gw['ostatnia_ile']), ('2026-10-02', 9299 + 46494 + 4649, 3))
        self.assertEqual(gw['miesiecznie'], 9299 + 46494)   # wrzesień, wszystkie karty
        self.assertNotIn('poza cyklem', gw['flagi'])

    def test_zwroty_kosztow_poza_pensja(self):
        g = self.grupa('H1')
        self.assertEqual(g, self.grupa('HZ1'))   # ten sam kontrahent w banku
        self.assertEqual((wyd(self.con, 'H1')['kategoria'], wyd(self.con, 'HZ1')['kategoria']), ('pensje', 'podroze'))
        s, po = self.stale()
        self.assertEqual(po[g + '|pensje']['miesiecznie'], 936793)
        self.assertNotIn(g + '|podroze', po)   # zwroty za wyjazdy nie są kosztem stałym
        self.assertEqual([x['kategoria'] for x in s['pozycje'] if x['kontrahent'] == g], ['pensje'])
        a = wydatki.analiza(self.con, '2026-01-01', '2026-10-15')
        pensje = next(x for x in a['kategorie'] if x['klucz'] == 'pensje')
        self.assertEqual(pensje['suma'], 500000 + 4 * 936793)   # zwroty też poza pensjami w kategoriach

    def test_decyzje_mata_staly(self):
        s, po = self.stale()
        figma = self.grupa('F1')
        self.wpis('cel_wydatku', {'kontrahent': figma, 'staly': False})
        self.wpis('cel_wydatku', {'kontrahent': self.grupa('X1'), 'staly': True, 'obszar': 'Rozwój', 'kto': 'Anna'})
        self.wpis('cel_wydatku', {'kontrahent': self.grupa('WI1'), 'obszar': 'Zespół', 'kto': 'Anna'})
        self.przebuduj()
        s, po = self.stale()
        self.assertNotIn(figma + '|narzedzia', po)
        self.assertNotIn(self.grupa('X1') + '|nieprzypisane', po)   # sierpień: ani we wrześniu, ani w październiku, więc wygasły
        self.assertIn(self.grupa('X1'), [x['kontrahent'] for x in s['wygasle']])
        self.assertIn({'nazwa': 'Zespół', 'miesiecznie': 3000, 'rocznie': 36000, 'ile': 1}, s['wg_obszaru'])
        self.assertEqual(next(x for x in s['wg_osoby'] if x['nazwa'] == 'Anna')['miesiecznie'], 3000)

    def test_kafle_bez_podwojnego_liczenia(self):
        a = wydatki.analiza(self.con, '2026-01-01', '2026-10-15')
        k = a['kafle']
        self.assertEqual(k['koszt_pracy']['suma'], 500000 + 4 * 936793 + 2 * 80000 + 60000 + 30000 + 6000)   # pensje, ZUS, PIT-4, PPK
        self.assertEqual(k['podatki']['suma'], 150000 + 10000)   # VAT i CIT, bez PIT-4 i ZUS
        self.assertEqual(a['razem'], sum(x['suma'] for x in a['kategorie']))
        self.assertEqual(a['razem'], -self.con.execute("SELECT SUM(o.kwota) FROM wydatki w JOIN operacje o ON o.id = w.operacja_id"
                                                        " WHERE o.data_ksiegowania >= '2026-01-01'").fetchone()[0])
        self.assertEqual(a['niepelne'], ['2026-10'])
        self.assertAlmostEqual(sum(x['udzial'] for x in a['kategorie']), 100, delta=0.5)

    def test_kafel_narzedzi_pelne_miesiace_i_roczne(self):
        n = wydatki.analiza(self.con, '2026-01-01', '2026-10-15')['kafle']['narzedzia']
        self.assertEqual((n['miesiecy'], n['od_miesiaca'], n['do_miesiaca']), (9, '2026-01', '2026-09'))
        biezace = 3 * 4100 + 5000 + 3 * 1200 + 4 * 3000 + 9299 * 3 + 46494   # styczeń do września, bez opłat rocznych
        self.assertEqual(n['biezace_mies'], (biezace + 4) // 9)
        self.assertEqual((n['roczne'], n['roczne_ile'], n['roczne_mies']), (467890 + 31000, 2, (467890 + 31000 + 6) // 12))
        self.assertEqual(n['srednia'], n['biezace_mies'] + n['roczne_mies'])

    def test_ostrzezenie_kategoria_niezgodna_z_rodzajem(self):
        g = wyd(self.con, 'P1')['kontrahent_grupa']
        self.assertEqual(wyd(self.con, 'P1')['rodzaj'], 'praca')
        self.wpis('kategoria_regula', {'kontrahent': g, 'kategoria': 'narzedzia'})
        self.przebuduj()
        self.assertEqual(wyd(self.con, 'P1')['rodzaj'], 'praca')   # rodzaj bez zmian, tylko ostrzeżenie
        d = wydatki.do_decyzji(self.con)
        x = next(y for y in d['niezgodne'] if y['klucz'] == g)
        self.assertTrue(x['od_mata'])
        self.assertEqual((x['kategoria'], x['rodzaj'], x['typ']), ('narzedzia', 'praca', 'niezgodna'))
        self.assertEqual(d['niezgodne'][0]['klucz'], g)   # decyzje Mata na górze
        self.assertIn(wyd(self.con, 'X1')['kontrahent_grupa'], [y['klucz'] for y in d['nieprzypisane']])
        bolt = next(y for y in d['niezgodne'] if y['klucz'] == 'bolt')
        self.assertEqual((bolt['typ'], bolt['rodzaj_podpowiedz'], bolt['ile']), ('bez_rodzaju', 'praca', 2))

    def test_karta_kontrahenta_i_eksport(self):
        g = wyd(self.con, 'F1')['kontrahent_grupa']
        k = wydatki.kontrahent(self.con, g, '2026-01-01', '2026-10-15')
        self.assertEqual((k['ile'], k['suma'], k['kategoria']), (4, 17300, 'narzedzia'))
        self.assertEqual(k['platnosci'][0]['data'], '2026-09-29')
        self.assertIn('reguła wbudowana', k['platnosci'][0]['kategoria_zrodlo'])
        with self.assertRaises(ValueError):
            wydatki.kontrahent(self.con, 'nie ma', None, None)
        nazwa, tresc, typ = eksport.plik(self.con, 'wydatki/analiza', {'format': 'xlsx', 'od': '2026-01-01', 'do': '2026-10-15'})
        wb = openpyxl.load_workbook(io.BytesIO(tresc))
        self.assertEqual(wb.sheetnames[:3], ['Kategorie', 'Kontrahenci', 'Koszty stałe'])
        suma = sum(r[1] for r in wb['Kategorie'].iter_rows(min_row=2, values_only=True) if r[1] is not None)
        self.assertAlmostEqual(suma, wydatki.analiza(self.con, '2026-01-01', '2026-10-15')['razem'] / 100, places=2)
        nazwa, tresc, typ = eksport.plik(self.con, 'wydatki/stale', {'format': 'csv'})
        self.assertTrue(nazwa.endswith('.csv'))
        self.assertIn('Kontrahent', tresc.decode('utf-8-sig').splitlines()[0])

    def test_schemat_i_odcisk(self):
        kolumny = {r[1] for r in self.con.execute('PRAGMA table_info(wydatki)')}
        self.assertTrue({'kategoria', 'kategoria_zrodlo', 'kategoria_wpis', 'kontrahent_grupa'} <= kolumny)
        self.assertIn('wydatki_cel', baza.SCHEMAT)
        import inspect
        self.assertIn('wydatki.__file__', inspect.getsource(odbudowa.odcisk_kodu))


if __name__ == '__main__':
    unittest.main()
