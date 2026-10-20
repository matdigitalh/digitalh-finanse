"""Pakiet 11 na zmyślonych danych: rentowność budżetu (tabela przestawna, wynik folderu = suma arkuszy, tylko Księgowane),
ostrzeżenia ustawień wierszy (worek przy stałej opłacie, faktura kosztowa poza cash flow, 23% u firmy z zagranicy, dzień
wystawienia inny niż na fakturach), karta klienta z umową (stan, ostatni dzień na wypowiedzenie z obu stron, termin odejścia,
plany po końcu umowy, punkty podpięcia pakietu 9), alerty umów na pulpit, przypięcia i widoki arkusza (wpisy z cofnięciem),
licznik planów nieaktualnych, filtry roku, kolumny VAT, nowe adresy serwera i odrzucanie odwołań do nieistniejących arkuszy."""
import csv
import http.client
import io
import json
import threading
import unittest
from datetime import date
from http.server import ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook

from finanse import (baza, ekrany, eksport, karta_klienta, konfig, odbudowa, ostrzezenia_wierszy, przestawna, rentownosc, serwer,
                     widoki, wynik, zmiany)
from testy.pomoc import Srodowisko
from testy.test_budzet import KOLUMNY, bs, dane

DZIS = '2026-10-13'   # dzień stanu w danych testowych (ostatnia operacja ING)
KAT = ('Koszty Beta', 'Beta', 'Koszty')


def wiersz(id_, rodzaj, podstawa, kontrahent, netto, miesiac, wyst, arkusz=KAT, nr='', ksiegowane='Tak', cf='Tak', status='Aktualne',
           etap='Planowana', stawka='0.23'):
    w = bs(id_, rodzaj, etap, podstawa, kontrahent, netto, miesiac, wyst, *arkusz, nr=nr, status=status)
    w.update({'Księgowane': ksiegowane, 'W cash flow': cf, 'Stawka VAT': stawka})
    return w


def dane11(srod):
    """Dane z test_budzet (folder Alfa: plany i rzeczywiste, wiersze VAT) plus folder Beta z wierszami do ostrzeżeń,
    pozycja z Księgowane = nie i folder DigitalH z planem wpływu (folder firmy, bez alertu „brak umowy”)."""
    dane(srod)
    p = srod.dane / 'Budget Studio' / 'budget_studio_2026-10-13.csv'
    wiersze = [
        # worek, a co miesiąc dokładnie jedna faktura (stała opłata)
        *[wiersz(f'10{i}', '-1', 'Faktura (w worku)', 'Drukarnia Worek', '-300,00', m, f'{m}-10', nr=f'FK/{i}', etap='Opłacona')
          for i, m in enumerate(('2026-06', '2026-07', '2026-08'), start=1)],
        wiersz('104', '-1', 'Faktura ', 'Koszt prywatny', '-200,00', '2026-11', '2026-11-30', cf='Nie'),
        wiersz('105', '-1', 'Faktura ', 'Tool Inc.', '-100,00', '2026-11', '2026-11-30'),
        # faktury wystawiane 15. dnia kolejnego miesiąca, a plany i reguła wiersza: ostatni dzień miesiąca budżetu
        *[wiersz(f'11{i}', '-1', 'Faktura ', 'Biuro Dzien', '-400,00', m, w_, nr=f'FB/{i}', etap='Opłacona')
          for i, (m, w_) in enumerate((('2026-06', '2026-07-15'), ('2026-07', '2026-08-15'), ('2026-08', '2026-09-15')), start=1)],
        *[wiersz(f'12{i}', '-1', 'Faktura ', 'Biuro Dzien', '-400,00', m, d)
          for i, (m, d) in enumerate((('2026-09', '2026-09-30'), ('2026-10', '2026-10-31'), ('2026-11', '2026-11-30'), ('2026-12', '2026-12-31')), start=1)],
        # wpłata właściciela: Księgowane = nie, poza rentownością
        wiersz('130', '1', 'Bez faktury', 'Wlasciciel', '5000,00', '2026-09', '2026-09-30', ksiegowane='Nie'),
        # plan wpływu w folderze firmy: bez alertu „brak umowy”
        wiersz('140', '1', 'Faktura ', 'Klient X', '700,00', '2026-11', '2026-11-30', arkusz=('New Business', 'DigitalH', 'Sprzedaż')),
    ]
    with open(p, 'a', encoding='utf-8', newline='') as f:
        csv.DictWriter(f, fieldnames=KOLUMNY, delimiter='\t').writerows(wiersze)


class Baza(unittest.TestCase):
    def setUp(self):
        self.srod = Srodowisko().__enter__()
        dane11(self.srod)
        self.raport = odbudowa.odbuduj()
        self.con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)

    def tearDown(self):
        self.con.close()
        self.srod.__exit__()

    def przebuduj(self):
        self.con.close()
        odbudowa.odbuduj()
        self.con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)

    def arkusz(self, nazwa):
        return self.con.execute('SELECT id FROM bud_arkusz WHERE nazwa = ?', (nazwa,)).fetchone()[0]


class TestPrzestawna(Baza):
    def test_wynik_folderu_to_suma_arkuszy_w_kazdym_miesiacu(self):
        d = przestawna.policz(self.con, '2026', ['folder', 'arkusz'])
        self.assertTrue(d['wiersze'])
        for f in d['wiersze']:
            self.assertEqual(f['wynik'], sum(a['wynik'] for a in f['dzieci']), f['nazwa'])
            for m in d['miesiace']:
                ark = sum((a['miesiace'].get(m) or {}).get('wynik', 0) for a in f['dzieci'])
                self.assertEqual((f['miesiace'].get(m) or {}).get('wynik', 0), ark, (f['nazwa'], m))
        self.assertEqual(d['razem']['wynik'], sum(f['wynik'] for f in d['wiersze']))

    def test_tylko_ksiegowane_i_kwoty_netto(self):
        d = przestawna.policz(self.con, '2026', ['folder'])
        alfa = next(f for f in d['wiersze'] if f['nazwa'] == 'Alfa')
        # wpływy Alfa: 1000 (sie) + 1000 (wrz) + 1000 (paź) + 500 (lip), obciążenie 500 (wrz)
        self.assertEqual((alfa['wplywy'], alfa['obciazenia'], alfa['wynik'], alfa['marza']), (350000, 50000, 300000, 85.7))
        self.assertEqual(alfa['miesiace']['2026-09'], {'wplywy': 100000, 'obciazenia': 50000, 'wynik': 50000, 'marza': 50.0})
        beta = next(f for f in d['wiersze'] if f['nazwa'] == 'Beta')
        self.assertEqual(beta['wplywy'], 0)   # wpłata właściciela ma Księgowane = nie
        self.assertNotIn('Zobowiązania fiskalne', [f['nazwa'] for f in d['wiersze']])   # daniny: Księgowane = nie
        self.assertIsNotNone(d['zarzadczy'])

    def test_poziomy_folder_i_os(self):
        d = przestawna.policz(self.con, '2026', 'kontrahent,folder,nieznany', folder='Alfa')
        self.assertEqual(d['poziomy'], ['kontrahent', 'folder'])
        self.assertEqual(sorted(w['nazwa'] for w in d['wiersze']), ['Alfa', 'Drukarnia'])
        self.assertIsNone(d['zarzadczy'])   # wynik zarządczy jest dla całej firmy, nie dla folderu
        self.assertEqual(przestawna.policz(self.con, '2026', os='ksiegowanie')['razem']['wynik'],
                         przestawna.policz(self.con, '2026')['razem']['wynik'])   # bez miesięcy księgowania: miesiąc budżetu
        self.assertEqual(przestawna.marza(-100, 0), None)
        self.assertEqual(przestawna.marza(1, 3), 33.3)


class TestOstrzezeniaWierszy(Baza):
    def test_regulay_bez_bazy(self):
        rz = lambda m, w='', s=None: {'typ': 'rzeczywista', 'miesiac': m, 'wystawienie': w, 'stawka': s}
        worek = {'podstawa': 'Faktura (w worku)', 'rodzaj': -1, 'stawka_vat': 0, 'w_cashflow': True}
        self.assertEqual([o['rodzaj'] for o in ostrzezenia_wierszy.dla_wiersza(worek, [rz('2026-06'), rz('2026-07'), rz('2026-08')])], ['worek_stala'])
        self.assertEqual(ostrzezenia_wierszy.dla_wiersza(worek, [rz('2026-06'), rz('2026-07')]), [])   # dwa miesiące to za mało
        self.assertEqual(ostrzezenia_wierszy.dla_wiersza(worek, [rz('2026-06'), rz('2026-07'), rz('2026-08'), rz('2026-08')]), [])
        koszt = {'podstawa': 'Faktura (jedna)', 'rodzaj': -1, 'stawka_vat': 23, 'w_cashflow': False, 'kontrahent': 'Drukarnia'}
        self.assertEqual([o['rodzaj'] for o in ostrzezenia_wierszy.dla_wiersza(koszt, [])], ['cf_off_koszt'])
        self.assertEqual(ostrzezenia_wierszy.dla_wiersza({**koszt, 'rodzaj': 1}, []), [])   # wpływ poza cash flow: inna sprawa
        zagr = {**koszt, 'w_cashflow': True, 'kontrahent': 'Canva Pty. Ltd.'}
        self.assertEqual([o['rodzaj'] for o in ostrzezenia_wierszy.dla_wiersza(zagr, [])], ['vat_zagranica'])
        self.assertEqual(ostrzezenia_wierszy.dla_wiersza({**zagr, 'kontrahent': 'Drukarnia'}, []), [])
        self.assertTrue(ostrzezenia_wierszy.dla_wiersza({**zagr, 'kontrahent': 'Drukarnia'}, [], zagraniczny_rejestr=True))
        zero = [rz('2026-06', s='0'), rz('2026-07', s='0')]
        self.assertIn('0%', ostrzezenia_wierszy.dla_wiersza({**zagr, 'kontrahent': 'Drukarnia'}, zero)[0]['opis'])
        self.assertEqual(ostrzezenia_wierszy.dla_wiersza({**zagr, 'stawka_vat': 0}, []), [])
        dz = {'podstawa': 'Faktura (jedna)', 'rodzaj': -1, 'stawka_vat': 23, 'w_cashflow': True, 'kontrahent': 'Biuro', 'dzien_wystawienia': 'Ostatni'}
        fakt = [rz('2026-06', '2026-07-15'), rz('2026-07', '2026-08-15'), rz('2026-08', '2026-09-15')]
        o = ostrzezenia_wierszy.dla_wiersza(dz, fakt)
        self.assertEqual([x['rodzaj'] for x in o], ['dzien_historia'])
        self.assertIn('„15. +1”', o[0]['rada'])
        self.assertEqual(ostrzezenia_wierszy.dla_wiersza({**dz, 'dzien_wystawienia': '15. +1'}, fakt), [])
        self.assertEqual(ostrzezenia_wierszy.dla_wiersza({**dz, 'dzien_wystawienia': 'Ostatni'}, fakt[:2]), [])   # za mało faktur

    def test_ostrzezenia_w_budzecie(self):
        o = ostrzezenia_wierszy.policz(self.con)
        po = {(x['kontrahent'], x['rodzaj']) for x in o['ostrzezenia']}
        # odbiór pakietu 11: worek z jedną płatnością miesięcznie dostaje ostrzeżenie
        self.assertIn(('Drukarnia Worek', 'worek_stala'), po)
        self.assertIn(('Koszt prywatny', 'cf_off_koszt'), po)
        self.assertIn(('Tool Inc.', 'vat_zagranica'), po)
        self.assertIn(('Biuro Dzien', 'dzien_historia'), po)
        self.assertFalse([x for x in o['ostrzezenia'] if x['folder'] == 'Alfa'])
        jeden = ostrzezenia_wierszy.policz(self.con, self.arkusz('Koszty Beta'))
        self.assertEqual(jeden['ile'], o['ile'])
        self.assertEqual(ostrzezenia_wierszy.policz(self.con, self.arkusz('Obsługa stała Alfa'))['ostrzezenia'], [])


class TestKartaKlienta(Baza):
    def test_opis_umowy_terminy(self):
        """Po scaleniu z pakietem 9: terminy i plany po końcu z umowy.termin i umowy.plany_po_koncu, karta zmienia tylko kształt."""
        from finanse import umowy
        dz = date.fromisoformat(DZIS)
        u = {'id': 'u1', 'klient': 'Alfa', 'rodzaj': 'okreslony', 'do': '2026-12-31', 'przedluzenie': True, 'wypowiedzenie_mies': 1,
             'koniec_miesiaca': True, 'alert_dni': 60, 'zawarta': '2026-01-01', 'wartosc_mies': 500000}
        o = karta_klienta.opis_umowy({**u, **umowy.termin(u, dz)}, DZIS)
        self.assertEqual((o['stan'], o['ostatni_dzien_wypowiedzenia'], o['termin_odejscia']),
                         ('alert', {'nasza': '2026-11-30', 'ich': '2026-11-30'}, '2026-12-31'))
        # termin na wypowiedzenie minął: odejście wg pakietu 9 (umowy._strona: koniec albo dziś + okres, co później)
        u2 = {**u, 'do': '2026-10-31', 'zawarta': '2025-11-01'}
        o = karta_klienta.opis_umowy({**u2, **umowy.termin(u2, dz)}, DZIS)
        self.assertEqual((o['stan'], o['termin_odejscia']), ('minął', '2026-11-30'))
        u3 = {'id': 'u3', 'klient': 'Alfa', 'rodzaj': 'nieokreslony', 'wypowiedzenie_mies': 3, 'koniec_miesiaca': True}
        o = karta_klienta.opis_umowy({**u3, **umowy.termin(u3, dz)}, DZIS)
        self.assertEqual((o['stan'], o['ostatni_dzien_wypowiedzenia'], o['termin_odejscia'], o['typ']), ('ok', None, '2027-01-31', 'nieokreślony'))
        u4 = {**u, 'przedluzenie': False, 'do': '2026-10-31'}
        p9 = {**u4, **umowy.termin(u4, dz), 'plany_po_koncu': {'ile': 2, 'netto': 500, 'od_miesiaca': '2026-11', 'arkusze': {}},
              'folder': 'Alfa', 'arkusze': [{'id': 'x', 'nazwa': 'X'}], 'arkusze_ids': ['x'], 'folder_zrodlo': 'wpis'}
        o = karta_klienta.opis_umowy(p9, DZIS, surowe={**u4, 'folder': 'Alfa', 'arkusze': ['x'], 'wpis': 'w1', 'czas': 't'})
        self.assertEqual(o['plany_po_koncu'], {'ile': 2, 'kwota': 500, 'od': '2026-11'})
        self.assertIsNone(o['ostatni_dzien_wypowiedzenia'])
        self.assertEqual(o['termin_odejscia'], '2026-10-31')
        # formularz karty odsyła surowy wpis: folder i arkusze tak, jak zapisane, bez pól wyliczonych
        self.assertEqual((o['dane']['folder'], o['dane']['arkusze']), ('Alfa', ['x']))
        self.assertFalse(set(o['dane']) & karta_klienta.POCHODNE)
        self.assertEqual(zmiany.sprawdz('umowa', o['dane'])['arkusze'], ['x'])
        self.assertIsNone(karta_klienta.opis_umowy({**u4, **umowy.termin(u4, dz), 'plany_po_koncu': {'ile': 0}}, DZIS)['plany_po_koncu'])

    def test_lista_karta_i_alerty_bez_umowy(self):
        l = karta_klienta.lista(self.con)
        alfa = next(x for x in l['klienci'] if x['klient'] == 'Alfa')
        self.assertEqual((alfa['stan'], alfa['folder'], [a['nazwa'] for a in alfa['arkusze']]), ('brak umowy', 'Alfa', ['Obsługa stała Alfa']))
        self.assertEqual(alfa['plany_12m'], 100000)   # plan wpływu w październiku (od miesiąca dnia stanu)
        self.assertEqual([o['rodzaj'] for o in alfa['ostrzezenia']], ['brak_umowy'])
        a = karta_klienta.alerty(self.con)
        self.assertEqual([(x['klient'], x['rodzaj']) for x in a], [('Alfa', 'brak umowy')])   # DigitalH to folder firmy
        k = karta_klienta.karta(self.con, 'Alfa')
        self.assertEqual((k['nip'], k['folder'], k['umowa'], k['stan']), (['111'], 'Alfa', None, 'brak umowy'))
        self.assertEqual(k['budzet']['wynik'], 300000)   # wynik budżetu folderu z tabeli przestawnej
        self.assertEqual([x['miesiac'] for x in k['plany']][:3], ['2026-10', '2026-11', '2026-12'])
        self.assertEqual(k['plany'][0]['plan_wplywy'], 100000)
        self.assertEqual(k['obroty']['przychod'], 200000)
        with self.assertRaises(ValueError):
            karta_klienta.karta(self.con, ' ')

    def test_umowa_z_karty_alerty_i_plany_po_koncu(self):
        zmiany.zapisz('umowa', zmiany.sprawdz('umowa', {'id': 'u-alfa', 'klient': 'alfa', 'rodzaj': 'okreslony', 'do': '2026-09-30',
                                                        'wypowiedzenie_mies': '1', 'wartosc_mies': '100000', 'alert_dni': '30'}))
        self.przebuduj()
        x = next(x for x in karta_klienta.lista(self.con)['klienci'] if x['klient'] == 'Alfa')   # nazwa klienta bez wielkości liter
        self.assertEqual((x['stan'], x['umowa']['wartosc_mies']), ('minął', 100000))
        self.assertEqual([o['rodzaj'] for o in x['ostrzezenia']], ['plany_po_koncu'])
        a = ekrany.alerty_pulpitu(self.con)
        self.assertEqual([(y['klient'], y['rodzaj']) for y in a['umowy']], [('Alfa', 'minął'), ('Alfa', 'plany po końcu')])
        self.assertEqual(a['liczniki_umow']['minął'], 1)
        # zmiana na umowę na czas nieokreślony usuwa alert
        zmiany.zapisz('umowa', zmiany.sprawdz('umowa', {**x['umowa']['dane'], 'rodzaj': 'nieokreslony', 'do': ''}))
        self.przebuduj()
        x = next(x for x in karta_klienta.lista(self.con)['klienci'] if x['klient'] == 'Alfa')
        self.assertEqual((x['stan'], x['ostrzezenia']), ('ok', []))


class TestPrzypieteIWidokiArkusza(Baza):
    def test_przypiecie_z_cofnieciem(self):
        """Po scaleniu z pakietem 4: przypięcie to pole „przypiety” arkusza w zdarzeniu „bud” (bez decyzji przypięty),
        osobnego wpisu „przypiecie” nie ma; GET budzet/przypiete czyta to samo pole."""
        aid, drugi = self.arkusz('Obsługa stała Alfa'), self.arkusz('Koszty Beta')
        with self.assertRaises(ValueError):
            zmiany.sprawdz('przypiecie', {'arkusz': aid})
        self.assertLessEqual({aid, drugi}, {a['id'] for a in widoki.przypiete(self.con)})
        pin = lambda a, tak: zmiany.zapisz('bud', zmiany.sprawdz('bud', {'op': 'arkusz', 'obiekt': {'id': a, 'przypiety': tak}}))
        pin(drugi, False)
        w = pin(aid, False)
        self.przebuduj()
        ids = {a['id'] for a in widoki.przypiete(self.con)}
        self.assertFalse({aid, drugi} & ids)
        zmiany.zapisz('cofnij', {'wpis': w['id']})
        self.przebuduj()
        ids = {a['id'] for a in widoki.przypiete(self.con)}
        self.assertIn(aid, ids)
        self.assertNotIn(drugi, ids)

    def test_widoki_arkusza(self):
        aid, inny = self.arkusz('Obsługa stała Alfa'), self.arkusz('Koszty Beta')
        d = zmiany.sprawdz('widok', {'ekran': 'arkusz', 'id': 'a1', 'nazwa': 'Wg kontrahenta',
                                     'parametry': {'arkusz': aid, 'grupuj': 'kontrahent', 'serie': 'wplywy,obciazenia'}, 'kolumny': []})
        ekrany.sprawdz_odwolania(self.con, 'widok', d)
        with self.assertRaises(ValueError):
            ekrany.sprawdz_odwolania(self.con, 'widok', {**d, 'parametry': {'grupuj': 'kontrahent'}})   # widok bez arkusza
        with self.assertRaises(ValueError):
            ekrany.sprawdz_odwolania(self.con, 'widok', {**d, 'parametry': {'arkusz': 'nie-ma'}})
        zmiany.zapisz('widok', d)
        zmiany.zapisz('widok', zmiany.sprawdz('widok', {'ekran': 'arkusz', 'id': 'b1', 'nazwa': 'Inny', 'parametry': {'arkusz': inny}}))
        self.przebuduj()
        self.assertEqual(widoki.arkusza(self.con, aid),
                         [{'id': 'a1', 'nazwa': 'Wg kontrahenta', 'wpis': widoki.arkusza(self.con, aid)[0]['wpis'],
                           'parametry': {'grupuj': 'kontrahent', 'serie': 'wplywy,obciazenia'}}])
        zmiany.zapisz('widok', zmiany.sprawdz('widok', {'ekran': 'arkusz', 'id': 'a1', 'usun': True}))
        self.przebuduj()
        self.assertEqual(widoki.arkusza(self.con, aid), [])
        self.assertEqual([v['id'] for v in widoki.arkusza(self.con, inny)], ['b1'])


class TestPulpitIRentownosc(Baza):
    def test_nieaktualne_plany_jak_lista(self):
        from finanse import zapytania_listy
        n = ekrany.nieaktualne_plany(self.con)
        lista = zapytania_listy.lista(self.con, 'platnosci', {'typ': 'plan', 'aktualnosc': 'Nieaktualna', 'ile': '1'})
        self.assertEqual(n['ile'], lista['razem'])
        self.assertGreater(n['ile'], 0)   # plany z minioną datą wystawienia (np. lipiec, wrzesień)

    def test_lata_i_vat(self):
        self.assertIn('2026', wynik.obroty_klientow(self.con, '2026')['lata'])
        self.assertIn('2026', rentownosc.klienci(self.con, '2026')['lata'])
        v = wynik.vat(self.con, '2026')
        self.assertIn('2026', v['lata'])
        meta = json.loads(self.con.execute("SELECT wartosc FROM meta WHERE klucz = 'podatki'").fetchone()[0])
        okresy = {x['okres'] for x in meta['vat'] if x['okres'].startswith('2026')}
        self.assertTrue(okresy <= {m['miesiac'] for m in v['miesiace']})   # miesiące z samą prognozą też na liście
        self.assertEqual([m['miesiac'] for m in v['miesiace']], sorted(m['miesiac'] for m in v['miesiace']))
        self.assertTrue(all(m['prognoza'] is not None for m in v['miesiace'] if m.get('tylko_prognoza')))
        nazwa, tresc, _ = eksport.plik(self.con, 'vat', {'rok': '2026', 'format': 'xlsx'})
        naglowki = [c.value for c in next(load_workbook(io.BytesIO(tresc))['VAT'].iter_rows(max_row=1))]
        self.assertTrue({'Prognoza z budżetu', 'Kwota od biura', 'Różnica prognozy i biura'} <= set(naglowki))


class TestSerwer(Baza):
    def test_nowe_adresy_i_odrzucenie_odwolan(self):
        stary = konfig.PORT
        h = ThreadingHTTPServer((konfig.HOST, 0), serwer.Obsluga)
        konfig.PORT = h.server_address[1]
        threading.Thread(target=h.serve_forever, daemon=True).start()
        try:
            def zapytaj(metoda, sciezka, cialo=None):
                c = http.client.HTTPConnection('127.0.0.1', konfig.PORT, timeout=20)
                c.request(metoda, sciezka, body=json.dumps(cialo) if cialo else None,
                          headers={'Host': f'127.0.0.1:{konfig.PORT}', 'X-Finanse': '1', 'Content-Type': 'application/json'})
                r = c.getresponse()
                return r.status, json.loads(r.read() or b'{}')
            aid = self.arkusz('Obsługa stała Alfa')
            for s in ('klienci/umowy', 'klienci/karta?nazwa=Alfa', 'pulpit/alerty', 'rentownosc/przestawna?rok=2026&poziomy=folder,kontrahent',
                      'budzet/ostrzezenia', f'budzet/ostrzezenia?arkusz={aid}', 'budzet/przypiete', f'budzet/widoki-arkusza?arkusz={aid}'):
                kod, w = zapytaj('GET', f'/api/{s}')
                self.assertEqual(kod, 200, (s, w))
            self.assertEqual(set(ekrany.ODCZYTY) - {'klienci/umowy', 'klienci/karta', 'pulpit/alerty', 'rentownosc/przestawna',
                                                    'budzet/ostrzezenia', 'budzet/przypiete', 'budzet/widoki-arkusza'}, set())
            kod, w = zapytaj('GET', '/api/klienci/karta')
            self.assertEqual(kod, 400)
            kod, w = zapytaj('POST', '/api/zmiana', {'typ': 'widok', 'dane': {'ekran': 'arkusz', 'id': 'x', 'nazwa': 'X',
                                                                            'parametry': {'arkusz': 'nie-ma'}}})
            self.assertEqual((kod, w['blad']), (400, 'nie ma takiego arkusza'))
            kod, w = zapytaj('POST', '/api/zmiana', {'typ': 'przypiecie', 'dane': {'arkusz': aid}})   # typ z pakietu 11 usunięty
            self.assertEqual(kod, 400, w)
            # przypięcie to zmiana budżetu (pakiet 4): przy wyłączonym planowaniu odrzucona, arkusz zostaje przypięty
            kod, w = zapytaj('POST', '/api/zmiana', {'typ': 'bud', 'dane': {'op': 'arkusz', 'obiekt': {'id': aid, 'przypiety': False}}})
            self.assertEqual(kod, 400, w)
            self.assertIn('Planowanie', w['blad'])
            kod, w = zapytaj('GET', '/api/budzet/przypiete')
            self.assertIn(aid, [a['id'] for a in w['arkusze']])
            kod, w = zapytaj('GET', '/api/ostrzezenia')   # pakiet 5: kafel „Lista płac do wpisania” na pulpicie
            self.assertEqual(kod, 200)
            self.assertTrue(all({'rodzaj', 'tytul', 'osoba', 'miesiac', 'link'} <= set(x) for x in w['ostrzezenia']))
        finally:
            h.shutdown()
            h.server_close()
            konfig.PORT = stary


class TestInterfejs(unittest.TestCase):
    WEB = Path(__file__).resolve().parent.parent / 'web'

    def test_pliki_podpiete_w_index(self):
        html = (self.WEB / 'index.html').read_text(encoding='utf-8')
        for p in ('klienci.js', 'pulpit.js', 'przestawna.js', 'arkusz_widoki.js', 'skroty.js', 'dodatki.css'):
            self.assertIn(p, html)
            self.assertTrue((self.WEB / p).is_file())

    def test_bez_martwego_kodu_i_z_haczykami(self):
        app = (self.WEB / 'app.js').read_text(encoding='utf-8')
        self.assertNotIn('wBudowie', app)   # APK-BUG-09
        self.assertIn('pulpitDodatki', app)
        self.assertNotIn('wysuwPrzypiete', app)   # po scaleniu z pakietem 4: drzewo Budżetu pokazuje przypięte, bez drugiej listy
        self.assertIn('roznica_biuro', app)
        self.assertIn('poArkuszu', (self.WEB / 'arkusz.js').read_text(encoding='utf-8'))
        self.assertIn('przypnijArkusz', (self.WEB / 'budzet.js').read_text(encoding='utf-8'))   # pinezka z pakietu 4
        pulpit = (self.WEB / 'pulpit.js').read_text(encoding='utf-8')
        self.assertIn('bez umowy`', pulpit)   # kafel umów: bez umów do decyzji pokazuje liczbę klientów bez umowy, nie „0 do decyzji”
        self.assertIn("odczyt('ostrzezenia')", pulpit)   # ostrzeżenia listy płac z pakietu 5 (decyzja 12), przy 404 kafla nie ma

    def test_tokeny_danin_w_obu_trybach(self):
        css = (self.WEB / 'dodatki.css').read_text(encoding='utf-8')
        self.assertEqual(css.count('--danina:'), 3)   # jasny, ciemny systemowy, ciemny wybrany przełącznikiem
        self.assertIn('prefers-color-scheme: dark', css)
        self.assertIn('[data-theme="dark"]', css)


if __name__ == '__main__':
    unittest.main()
