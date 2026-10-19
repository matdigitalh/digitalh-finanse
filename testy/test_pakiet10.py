"""Pakiet 10 (eksport i import) na zmyślonych danych: płatności budżetu w układzie eksportu BS (kolumny, zapis kwot,
sumy jak w stopce listy przy tych samych filtrach), eksport list w Excelu i CSV, rozpoznanie źródła po zawartości,
podsumowanie importu (Dodane, Zaktualizowane, Bez zmian, Pominięte z powodem), import przez przeciągnięcie do okna
(nazwa wg konwencji, plik 600, ten sam plik drugi raz = Bez zmian), adresy serwera."""
import csv
import http.client
import io
import json
import os
import threading
import unittest
from datetime import date
from decimal import Decimal
from http.server import ThreadingHTTPServer

from openpyxl import load_workbook

from finanse import baza, cashflow, eksport, konfig, odbudowa, serwer, zapytania, zapytania_budzet
from finanse.importy import dziennik, podsumowanie, reczny, rozpoznanie
from testy.pomoc import Srodowisko, faktura, operacja, zapisz_fakturownie, zapisz_ing, zapisz_rejestr
from testy.test_budzet import KOLUMNY, bs, dane
from testy.test_importy import wrzesien_pazdziernik


def _xlsx_tabela(tresc: bytes, arkusz: str):
    ws = load_workbook(io.BytesIO(tresc))[arkusz]
    wiersze = list(ws.iter_rows(values_only=True))
    return list(wiersze[0]), [list(w) for w in wiersze[1:]], ws


def _bs_na_grosze(s: str) -> int:
    return int(Decimal(s.replace(',', '.')) * 100) if s else 0


def _ing_bajty(katalog, nazwa, wiersze) -> bytes:
    return zapisz_ing(katalog, nazwa, wiersze).read_bytes()


class Baza(unittest.TestCase):
    def setUp(self):
        self.srod = Srodowisko().__enter__()
        dane(self.srod)
        odbudowa.odbuduj()
        self.con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)

    def tearDown(self):
        self.con.close()
        self.srod.__exit__()


class TestEksportPlatnosciJakBS(Baza):
    def test_uklad_kolumn_jak_w_pliku_bs(self):
        nazwa, tresc, typ = eksport.plik(self.con, 'budzet/lista', {'format': 'csv'})
        self.assertTrue(nazwa.endswith('.csv'))
        self.assertEqual(typ, eksport.CSV)
        self.assertFalse(tresc.startswith(b'\xef\xbb\xbf'), 'plik BS nie ma BOM, eksport też nie')
        naglowek = tresc.decode('utf-8').split('\n')[0].split('\t')
        with open(self.srod.dane / 'Budget Studio' / 'budget_studio_2026-10-13.csv', encoding='utf-8') as f:
            naglowek_bs = next(csv.reader(f, delimiter='\t'))
        self.assertEqual(naglowek, naglowek_bs)
        self.assertEqual(eksport.KOLUMNY_BS, KOLUMNY)
        _, tresc_x, _ = eksport.plik(self.con, 'budzet/lista', {'format': 'xlsx'})
        self.assertEqual(_xlsx_tabela(tresc_x, 'Płatności')[0], KOLUMNY)

    def test_wiersz_jak_w_bs(self):
        _, tresc, _ = eksport.plik(self.con, 'budzet/lista', {'format': 'csv'})
        wiersze = list(csv.DictReader(io.StringIO(tresc.decode('utf-8')), delimiter='\t'))
        druk = next(w for w in wiersze if w['ID'] == '5')
        self.assertEqual((druk['Rodzaj'], druk['Netto (wartość)'], druk['Brutto (wartość)']), ('-1', '-500', '-615'))
        self.assertEqual((druk['Podstawa płatności'], druk['Sekcja'], druk['Przedmiot transakcji']), ('Faktura ', '-', 'Przedmiot transakcji'))
        self.assertEqual((druk['Księgowane'], druk['W cash flow'], druk['Stawka VAT']), ('Tak', 'Tak', '0,23'))
        self.assertEqual((druk['Miesiąc budżetowy'], druk['Dzień płatności - realny'], druk['Przesunięcie płatności']), ('2026-09-01', '2026-09-10', '0'))
        self.assertEqual((druk['Arkusz'], druk['Folder'], druk['Kategoria']), ('Obsługa stała Alfa', 'Alfa', 'Obsługa stała'))
        self.assertEqual(druk['Data utworzenia'], '2026-03-01 10:00:00')   # z pliku eksportu BS
        self.assertEqual(sum(1 for w in wiersze if w['ID'] == '6'), 2)      # dokument na dwie komórki: dwa wiersze, jak w BS

    def test_zapis_kwot_jak_w_bs(self):
        self.assertEqual([eksport.kwota_bs(x) for x in (-1031, 0, 999990, 12300, -100, 5)], ['-10,31', '0', '9999,9', '123', '-1', '0,05'])
        self.assertEqual(eksport.podstawa_bs('Umowa o pracę', 'osoba'), 'Umowa o pracę - dla wykonawcy')
        self.assertEqual(eksport.podstawa_bs('Umowa zlecenie', 'us'), 'Umowa zlecenie - PIT')
        self.assertEqual(eksport.podstawa_bs('Bez faktury', None), 'Bez faktury ')
        self.assertEqual(eksport.podstawa_bs('Faktura (w worku)', None), 'Faktura (w worku)')

    def test_sumy_jak_w_stopce_listy_przy_tych_samych_filtrach(self):
        for filtry in ({}, {'lista': 'po_terminie'}, {'rodzaj': '1'}, {'q': 'Drukarnia'}, {'lista': 'zaplacone', 'rodzaj': '-1'}):
            lista = zapytania_budzet.lista(self.con, filtry.get('lista', 'wszystkie'), filtry.get('rodzaj'), filtry.get('q'))
            _, tresc, _ = eksport.plik(self.con, 'budzet/lista', {**filtry, 'format': 'xlsx', 'strona': '2', 'ile': '50'})
            naglowek, wiersze, _ = _xlsx_tabela(tresc, 'Płatności')
            i = naglowek.index('Brutto (wartość)')
            self.assertEqual(len(wiersze), lista['razem'], filtry)
            self.assertEqual(int(sum(Decimal(str(w[i])) for w in wiersze) * 100), lista['suma'], filtry)
            _, tresc_csv, _ = eksport.plik(self.con, 'budzet/lista', {**filtry, 'format': 'csv'})
            csv_w = list(csv.DictReader(io.StringIO(tresc_csv.decode('utf-8')), delimiter='\t'))
            self.assertEqual(sum(_bs_na_grosze(w['Brutto (wartość)']) for w in csv_w), lista['suma'], filtry)

    def test_podsumowanie_w_excelu(self):
        _, tresc, _ = eksport.plik(self.con, 'budzet/lista', {'format': 'xlsx', 'lista': 'wszystkie'})
        _, wiersze, _ = _xlsx_tabela(tresc, 'Podsumowanie')
        suma = next(w for w in wiersze if w[0].startswith('Suma brutto'))
        self.assertEqual(int(suma[2] * 100), zapytania_budzet.lista(self.con)['suma'])


class TestWolanieZapytan(unittest.TestCase):
    def test_wszystkie_strony_i_tylko_znane_parametry(self):
        dane_ = list(range(1234))
        wolania = []

        def lista(con, nazwa='wszystkie', q=None, strona=1, ile=100):
            wolania.append((nazwa, q, int(strona), int(ile)))
            s, i = int(strona), int(ile)
            return {'razem': len(dane_), 'wiersze': dane_[(s - 1) * i:s * i]}

        w = eksport.wolaj(lista, None, {'lista': 'po_terminie', 'q': 'x', 'obcy': '1', 'strona': '3', 'ile': '50', 'format': 'csv'},
                          aliasy={'lista': 'nazwa'})
        self.assertEqual(w['wiersze'], dane_)
        self.assertEqual(wolania[0], ('po_terminie', 'x', 1, 500))
        self.assertEqual(len(wolania), 3)

    def test_zly_rok_pomijany(self):
        widziane = []
        eksport.wolaj(lambda con, rok=None: widziane.append(rok) or {}, None, {'rok': '2026; DROP'})
        eksport.wolaj(lambda con, rok=None: widziane.append(rok) or {}, None, {'rok': '2025'})
        self.assertEqual(widziane, [None, '2025'])


class TestEksportList(Baza):
    def test_csv_bom_srednik_przecinek(self):
        nazwa, tresc, typ = eksport.plik(self.con, 'operacje', {'format': 'csv'})
        self.assertTrue(tresc.startswith(b'\xef\xbb\xbf'))
        wiersze = list(csv.reader(io.StringIO(tresc.decode('utf-8-sig')), delimiter=';'))
        self.assertEqual(wiersze[0][:4], ['Data księgowania', 'Data transakcji', 'Rachunek', 'Kwota'])
        self.assertEqual(len(wiersze) - 1, zapytania.operacje(self.con)['razem'])
        self.assertIn('-10,00', [w[3] for w in wiersze])

    def test_filtry_z_adresu_listy(self):
        _, tresc, _ = eksport.plik(self.con, 'operacje', {'format': 'xlsx', 'kierunek': 'wplywy', 'sortuj': 'kwota', 'kolejnosc': 'rosnaco'})
        naglowek, wiersze, _ = _xlsx_tabela(tresc, 'Bank ING')
        kwoty = [int(w[naglowek.index('Kwota')] * 100) for w in wiersze]
        self.assertTrue(all(k > 0 for k in kwoty))
        self.assertEqual(sum(kwoty), zapytania.operacje(self.con, kierunek='wplywy')['wplywy'])

    def test_tekst_jak_formula_zabezpieczony(self):
        zapisz_ing(self.srod.dane, 'ing_historia_2026-10-14_2026-10-14.csv',
                   [operacja('2026-10-14', '-1,00', '49999,00', '=HYPERLINK("http://x")', '+48 tytuł', ref='Z1')])
        odbudowa.odbuduj()
        con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)
        try:
            _, tresc, _ = eksport.plik(con, 'operacje', {'format': 'csv', 'q': 'HYPERLINK'})
            w = list(csv.reader(io.StringIO(tresc.decode('utf-8-sig')), delimiter=';'))[1]
            self.assertEqual(w[8], '\'=HYPERLINK("http://x")')
            self.assertEqual(w[11], "'+48 tytuł")
            _, tresc_x, _ = eksport.plik(con, 'operacje', {'format': 'xlsx', 'q': 'HYPERLINK'})
            _, _, ws = _xlsx_tabela(tresc_x, 'Bank ING')
            self.assertEqual(ws.cell(2, 9).data_type, 's')
        finally:
            con.close()

    def test_naleznosci_z_wiekowaniem(self):
        _, tresc, _ = eksport.plik(self.con, 'naleznosci', {'format': 'xlsx'})
        wb = load_workbook(io.BytesIO(tresc))
        self.assertEqual(wb.sheetnames, ['Należności', 'Wiekowanie', 'Podsumowanie'])
        nal = zapytania.naleznosci(self.con)
        _, wiersze, _ = _xlsx_tabela(tresc, 'Należności')
        self.assertEqual(len(wiersze), nal['razem'])
        _, wiek, _ = _xlsx_tabela(tresc, 'Wiekowanie')
        self.assertEqual(int(wiek[-1][2] * 100), zapytania.podsumowanie_naleznosci(self.con)['razem'])

    def test_cash_flow_wg_granulacji(self):
        for gran in ('dzien', 'tydzien', 'miesiac'):
            q = {'granulacja': gran, 'od': '2026-09-01', 'do': '2026-12-31'}
            d = cashflow.okresy(self.con, **q)
            _, tresc, _ = eksport.plik(self.con, 'cashflow/okresy', {**q, 'format': 'xlsx'})
            naglowek, wiersze, _ = _xlsx_tabela(tresc, 'Cash flow')
            self.assertEqual(len(wiersze), len(d['okresy']), gran)
            i = naglowek.index('Saldo narastająco')
            self.assertEqual([int(w[i] * 100) for w in wiersze], [o['narastajaco'] for o in d['okresy']], gran)

    def test_arkusz_wiersze_razy_miesiace(self):
        aid = self.con.execute("SELECT id FROM bud_arkusz WHERE nazwa = 'Obsługa stała Alfa'").fetchone()[0]
        _, tresc, _ = eksport.plik(self.con, 'budzet/arkusz', {'id': aid, 'format': 'csv'})
        wiersze = list(csv.reader(io.StringIO(tresc.decode('utf-8-sig')), delimiter=';'))
        d = zapytania_budzet.arkusz(self.con, aid)
        self.assertEqual(wiersze[0][12:-1], d['miesiace'])
        saldo = next(w for w in wiersze if w[0] == 'Saldo okresu')
        self.assertEqual(saldo[12:-1], [eksport._kwota_csv(p['saldo']) for p in d['podsumowanie']])

    def test_wszystkie_eksporty_dzialaja(self):
        aid = self.con.execute('SELECT id FROM bud_arkusz LIMIT 1').fetchone()[0]
        for rodzaj in eksport.EKSPORTY:
            q = {'id': aid} if rodzaj in ('budzet/arkusz', 'arkusz') else {'nazwa': 'Alfa'} if rodzaj == 'klient' else {}
            for fmt in eksport.FORMATY:
                nazwa, tresc, typ = eksport.plik(self.con, rodzaj, {**q, 'format': fmt})
                self.assertTrue(nazwa.startswith('Finanse ') and nazwa.endswith('.' + fmt), nazwa)
                self.assertNotIn('/', nazwa)
                if fmt == 'xlsx':
                    self.assertTrue(load_workbook(io.BytesIO(tresc)).sheetnames)

    def test_nieznany_eksport_i_format(self):
        with self.assertRaises(ValueError):
            eksport.plik(self.con, 'nie-ma-takiego', {})
        with self.assertRaises(ValueError):
            eksport.plik(self.con, 'operacje', {'format': 'pdf'})


class TestRozpoznanie(unittest.TestCase):
    def setUp(self):
        self.srod = Srodowisko().__enter__()

    def tearDown(self):
        self.srod.__exit__()

    def test_zrodla_po_naglowkach(self):
        ing_ = _ing_bajty(self.srod.dane, 'x.csv', [operacja('2026-10-01', '1,00', '1,00', 'A', 'B', ref='R')])
        self.assertEqual(rozpoznanie.rozpoznaj(ing_, 'cokolwiek.csv')['rodzaj'], 'ING')
        fak = zapisz_fakturownie(self.srod.dane, 'f.csv', [faktura(1, 'FV1', 'Wystawiona', '2026-10-01', '2026-10-08', 'K', '1', [('X', 10)])])
        self.assertEqual(rozpoznanie.rozpoznaj(fak.read_bytes(), 'export.csv')['rodzaj'], 'Fakturownia')
        dane(self.srod)
        bs_ = (self.srod.dane / 'Budget Studio' / 'budget_studio_2026-10-13.csv').read_bytes()
        self.assertEqual(rozpoznanie.rozpoznaj(bs_, 'payments.csv')['rodzaj'], 'Budget Studio')
        rej = zapisz_rejestr(self.srod.dane, 'r.xlsx', [['ZR/1', 'A/1', date(2026, 9, 1), date(2026, 9, 1), 'D (00001)', 0, 0, 100.0, 23.0, 123.0]], [])
        self.assertEqual(rozpoznanie.rozpoznaj(rej.read_bytes(), 'rejestr.xlsx')['rodzaj'], 'Rejestr zakupów')

    def test_nieznane_z_ostrzezeniem(self):
        for tresc, fragment in ((b'%PDF-1.7 ...', 'PDF'), (b'', 'pusty'), (b'a;b;c\n1;2;3\n', 'Nieznany układ'),
                                ('<?xml version="1.0"?><Faktura xmlns="http://crd.gov.pl/wzor/2025/06/25/13775/">'.encode(), 'KSeF'),
                                (b'PK\x03\x04zepsuty', 'uszkodzony'), ('zażółć'.encode('cp1250'), 'UTF-8')):
            r = rozpoznanie.rozpoznaj(tresc, 'plik')
            self.assertIsNone(r['rodzaj'], fragment)
            self.assertIn(fragment, r['powod'])

    def test_zly_separator(self):
        buf = io.StringIO()
        csv.writer(buf).writerow(['data transakcji', 'data księgowania', 'kwota', 'saldo po transakcji', 'nazwa i adres kontrahenta',
                                  'rachunek ING "NRB" (26 znaków)', 'numer referencyjny'])
        r = rozpoznanie.rozpoznaj(buf.getvalue().encode(), 'ing.csv')
        self.assertIsNone(r['rodzaj'])
        self.assertIn('separatorem jest przecinek zamiast średnik', r['powod'])

    def test_wlasny_eksport_nie_jest_plikiem_bs(self):
        dane(self.srod)
        odbudowa.odbuduj()
        con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)
        try:
            nazwa, tresc, _ = eksport.plik(con, 'budzet/lista', {'format': 'csv'})
        finally:
            con.close()
        self.assertEqual(rozpoznanie.rozpoznaj(tresc, nazwa)['powod'], rozpoznanie.WLASNY)
        przemianowany = tresc.decode('utf-8').replace('\n5\t', '\nplan:w_1:2026-11:0\t', 1).encode()
        self.assertEqual(rozpoznanie.rozpoznaj(przemianowany, 'budget_studio_2026-10-20.csv')['powod'], rozpoznanie.WLASNY)

    def test_nazwa_wg_konwencji(self):
        p = zapisz_ing(self.srod.dane, 'x.csv', [operacja('2026-10-13', '1,00', '2,00', 'A', 'B', ref='R1'),
                                                  operacja('2026-09-02', '1,00', '1,00', 'A', 'B', ref='R2')])
        w = podsumowanie.wczytaj('ING', p)
        self.assertEqual(rozpoznanie.nazwa_docelowa('ING', 'ING/ing_historia_*.csv', w.rekordy), 'ing_historia_2026-09-02_2026-10-13.csv')
        self.assertEqual(rozpoznanie.nazwa_docelowa('Fakturownia', 'Fakturownia/fakturownia_sprzedaz_*.csv', {}, dzis=date(2026, 10, 19)),
                         'fakturownia_sprzedaz_2026-10-19.csv')
        self.assertEqual(rozpoznanie.nazwa_docelowa('Wyniki biura', 'Biuro/biuro_wyniki_*.xls', {}, okres='2026-06'), 'biuro_wyniki_2026-06.xls')


def _mtime(p, t):
    os.utime(p, (t, t))


class TestPodsumowanieImportu(unittest.TestCase):
    def setUp(self):
        self.srod = Srodowisko().__enter__()

    def tearDown(self):
        self.srod.__exit__()

    def historia(self):
        """Historia po odbudowie: {nazwa pliku: wiersz historii} i otwarte połączenie (zamyka test)."""
        odbudowa.odbuduj()
        con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)
        return {x['plik']: x for x in podsumowanie.historia(con)['importy']}, con

    def test_dodane_i_bez_zmian_przy_nakladajacych_sie_eksportach(self):
        wrzesien_pazdziernik(self.srod.dane)
        _mtime(self.srod.dane / 'ING' / 'ing_historia_2026-09-01_2026-09-30.csv', 1_700_000_000)
        _mtime(self.srod.dane / 'ING' / 'ing_historia_2026-09-15_2026-10-13.csv', 1_700_100_000)
        h, con = self.historia()
        con.close()
        a, b = h['ing_historia_2026-09-01_2026-09-30.csv'], h['ing_historia_2026-09-15_2026-10-13.csv']
        self.assertEqual((a['dodane'], a['zaktualizowane'], a['bez_zmian'], a['pominiete']), (2, 0, 0, 0))
        self.assertEqual((b['dodane'], b['zaktualizowane'], b['bez_zmian'], b['pominiete']), (5, 0, 2, 0))
        self.assertEqual(b['sposob'], 'folder Dane')

    def test_zaktualizowane_ze_zmianami_i_najnowszy_plik(self):
        f1 = zapisz_fakturownie(self.srod.dane, 'fakturownia_sprzedaz_2026-10-01.csv', [
            faktura(1, 'FV2026/10/1', 'Wystawiona', '2026-10-01', '2026-10-08', 'Klient', '1', [('X', 100)]),
            faktura(2, 'FV2026/10/2', 'Wystawiona', '2026-10-02', '2026-10-09', 'Klient', '1', [('Y', 50)])])
        f2 = zapisz_fakturownie(self.srod.dane, 'fakturownia_sprzedaz_2026-10-13.csv', [
            faktura(1, 'FV2026/10/1', 'Opłacona', '2026-10-01', '2026-10-08', 'Klient', '1', [('X', 100)], oplacono='123.0'),
            faktura(2, 'FV2026/10/2', 'Wystawiona', '2026-10-02', '2026-10-09', 'Klient', '1', [('Y', 50)]),
            faktura(3, 'FV2026/10/3', 'Wystawiona', '2026-10-03', '2026-10-10', 'Klient', '1', [('Z', 10)])])
        _mtime(f1, 1_700_000_000)
        _mtime(f2, 1_700_100_000)
        h, con = self.historia()
        try:
            stary, nowy = h[f1.name], h[f2.name]
            self.assertIn('zastąpiony nowszym', stary['stan'])
            self.assertEqual(nowy['stan'], 'wczytany')
            self.assertEqual((nowy['dodane'], nowy['zaktualizowane'], nowy['bez_zmian']), (1, 1, 1))
            sz = podsumowanie.szczegoly(con, nowy['id'], 'zaktualizowane')
            self.assertEqual(sz['razem'], 1)
            self.assertIn('status: Wystawiona → Opłacona', sz['wiersze'][0]['zmiany'])
            self.assertIn('kwota opłacona: 0,00 zł → 123,00 zł', sz['wiersze'][0]['zmiany'])
        finally:
            con.close()

    def test_pominiete_z_powodem_najpierw(self):
        dane(self.srod)
        p = self.srod.dane / 'Budget Studio' / 'budget_studio_2026-10-13.csv'
        with open(p, 'a', encoding='utf-8', newline='') as f:
            w = bs('99', '', 'Planowana', 'Faktura ', 'Bez rodzaju', '10,00', '2026-10', '2026-10-20', 'Obsługa stała Alfa', 'Alfa', 'Obsługa stała')
            csv.DictWriter(f, fieldnames=KOLUMNY, delimiter='\t').writerow(w)
        zapisz_ing(self.srod.dane, 'ing_historia_2026-10-14_2026-10-14.csv', [
            operacja('2026-10-14', '-5,00', '10,00', 'KARTA', 'SKLEP', ref=''),
            operacja('2026-10-14', '-5,00', '15,00', 'KARTA', 'SKLEP', ref='')])
        zapisz_rejestr(self.srod.dane, 'biuro_rejestr_zakupow_2026-09_2026-09.xlsx',
                       [['ZR/1', 'A/1', date(2026, 9, 1), date(2026, 9, 1), 'DOSTAWCA (00001)', 0, 0, 100.0, 23.0, 123.0],
                        [None, 'A/2', date(2026, 9, 2), None, 'BEZ NUMERU (00002)', 0, 0, 10.0, 2.3, 12.3]], [])
        h, con = self.historia()
        try:
            b = h['budget_studio_2026-10-13.csv']
            self.assertEqual(b['pominiete'], 1)
            sz = podsumowanie.szczegoly(con, b['id'])
            self.assertEqual(sz['kategoria'], 'pominiete')   # jak w BS: najpierw Pominięte
            self.assertIn('rodzaj inny niż 1', sz['wiersze'][0]['powod'])
            i = h['ing_historia_2026-10-14_2026-10-14.csv']
            self.assertEqual((i['dodane'], i['pominiete']), (1, 1))
            self.assertIn('taka sama operacja', podsumowanie.szczegoly(con, i['id'], 'pominiete')['wiersze'][0]['powod'])
            r = h['biuro_rejestr_zakupow_2026-09_2026-09.xlsx']
            self.assertEqual((r['dodane'], r['pominiete']), (1, 1))   # wiersz bez numeru z danymi; suma arkusza nie
        finally:
            con.close()

    def test_zly_plik_caly_w_pominietych(self):
        wrzesien_pazdziernik(self.srod.dane)
        (self.srod.dane / 'ING' / 'ing_historia_2026-10-14_2026-10-14.csv').write_text('to nie jest eksport', encoding='utf-8')
        h, con = self.historia()
        try:
            zly = h['ing_historia_2026-10-14_2026-10-14.csv']
            self.assertEqual((zly['stan'], zly['pominiete'], zly['dodane']), ('błąd odczytu', 1, 0))
            self.assertIn('brak kolumn', podsumowanie.szczegoly(con, zly['id'])['wiersze'][0]['powod'])
        finally:
            con.close()

    def test_szczegoly_szukaj_i_strony(self):
        wrzesien_pazdziernik(self.srod.dane)
        h, con = self.historia()
        try:
            x = h['ing_historia_2026-09-15_2026-10-13.csv']
            self.assertEqual(podsumowanie.szczegoly(con, x['id'], 'dodane', q='dostawca')['razem'], 2)
            s = podsumowanie.szczegoly(con, x['id'], 'dodane', strona=2, ile=2)
            self.assertEqual((s['razem'], len(s['wiersze'])), (5, 2))
            with self.assertRaises(ValueError):
                podsumowanie.szczegoly(con, 'plik:../../etc/passwd')
        finally:
            con.close()


class TestImportPrzeciagniecie(unittest.TestCase):
    def setUp(self):
        self.srod = Srodowisko().__enter__()
        wrzesien_pazdziernik(self.srod.dane)
        odbudowa.odbuduj()
        self.inny = self.srod.katalog / 'pobrane'
        self.inny.mkdir()

    def tearDown(self):
        self.srod.__exit__()

    def pliki_ing(self):
        return sorted(p.name for p in (self.srod.dane / 'ING').iterdir())

    def test_podglad_nic_nie_zapisuje(self):
        przed = self.pliki_ing()
        tresc = _ing_bajty(self.inny, 'Historia.csv', [operacja('2026-10-20', '7,00', '7,00', 'NOWY', 'FV9', ref='N1')])
        r = reczny.przyjmij(tresc, 'Historia.csv', zapisz=False)
        self.assertTrue(r['rozpoznany'])
        self.assertEqual(r['liczniki'], {'dodane': 1, 'zaktualizowane': 0, 'bez_zmian': 0, 'pominiete': 0})
        self.assertEqual(r['cel'], 'ING/ing_historia_2026-10-20_2026-10-20.csv')
        self.assertFalse(r['zapisano'])
        self.assertEqual(self.pliki_ing(), przed)
        self.assertEqual(dziennik.wczytaj(), [])

    def test_zapis_pliku_wg_konwencji_z_uprawnieniami_600(self):
        tresc = _ing_bajty(self.inny, 'Lista_transakcji.csv', [
            operacja('2026-10-20', '7,00', '7,00', 'NOWY', 'FV9', ref='N1'),
            operacja('2026-10-13', '-50,00', '8450,00', '', 'opłata', ref='')])
        r = reczny.przyjmij(tresc, '/Users/kto/Downloads/Lista_transakcji.csv', zapisz=True)
        self.assertTrue(r['zapisano'])
        cel = self.srod.dane / 'ING' / 'ing_historia_2026-10-13_2026-10-20.csv'
        self.assertEqual(cel.read_bytes(), tresc)              # surowy plik bez zmian
        self.assertEqual(cel.stat().st_mode & 0o777, 0o600)
        self.assertEqual(r['liczniki'], {'dodane': 1, 'zaktualizowane': 0, 'bez_zmian': 1, 'pominiete': 0})
        wpis = dziennik.wczytaj()[0]
        self.assertEqual((wpis['plik_zrodlowy'], wpis['zapisany_jako']), ('Lista_transakcji.csv', 'ING/ing_historia_2026-10-13_2026-10-20.csv'))
        self.assertEqual(next(dziennik.katalog().glob('*.jsonl')).stat().st_mode & 0o777, 0o600)
        self.assertFalse(list((self.srod.dane / 'ING').glob('.*')), 'plik tymczasowy został sprzątnięty')
        odbudowa.odbuduj()
        con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)
        try:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM operacje WHERE nr_ref = 'N1'").fetchone()[0], 1)
            h = {x['plik']: x for x in podsumowanie.historia(con)['importy']}
            x = h['ing_historia_2026-10-13_2026-10-20.csv']
            self.assertEqual((x['sposob'], x['plik_zrodlowy'], x['dodane'], x['bez_zmian']),
                             ('przeciągnięty do okna', 'Lista_transakcji.csv', 1, 1))
        finally:
            con.close()

    def test_ten_sam_plik_drugi_raz_bez_zmian(self):
        istniejacy = self.srod.dane / 'ING' / 'ing_historia_2026-09-15_2026-10-13.csv'
        tresc = istniejacy.read_bytes()
        przed = self.pliki_ing()
        r = reczny.przyjmij(tresc, 'ten sam.csv', zapisz=True)
        self.assertEqual(r['duplikat'], 'ING/ing_historia_2026-09-15_2026-10-13.csv')
        self.assertIsNone(r['cel'])
        self.assertEqual(r['liczniki']['bez_zmian'], r['wierszy'])
        self.assertEqual(r['wierszy'], 7)
        self.assertEqual((r['liczniki']['dodane'], r['liczniki']['zaktualizowane']), (0, 0))
        self.assertEqual(self.pliki_ing(), przed)
        con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)
        try:
            wpisy = [x for x in podsumowanie.historia(con)['importy'] if x['id'].startswith('wpis:')]
            self.assertEqual((len(wpisy), wpisy[0]['bez_zmian']), (1, 7))
            self.assertEqual(podsumowanie.szczegoly(con, wpisy[0]['id'])['razem'], 7)
        finally:
            con.close()

    def test_inna_tresc_pod_ta_sama_nazwa_dostaje_numer(self):
        tresc = _ing_bajty(self.inny, 'a.csv', [operacja('2026-09-01', '1,00', '1,00', 'X', 'Y', ref='Q1'),
                                                 operacja('2026-09-30', '1,00', '2,00', 'X', 'Y', ref='Q2')])
        r = reczny.przyjmij(tresc, 'a.csv', zapisz=True)
        self.assertEqual(r['cel'], 'ING/ing_historia_2026-09-01_2026-09-30_2.csv')

    def test_fakturownia_i_nieznany_format(self):
        p = zapisz_fakturownie(self.inny, 'eksport.csv', [faktura(1, 'FV2026/10/9', 'Wystawiona', '2026-10-01', '2026-10-08', 'K', '1', [('X', 10)])])
        r = reczny.przyjmij(p.read_bytes(), 'eksport.csv', zapisz=True)
        self.assertEqual(r['cel'], f'Fakturownia/fakturownia_sprzedaz_{date.today().isoformat()}.csv')
        self.assertTrue(r['uwagi'])
        n = reczny.przyjmij(b'kolumna;inna\n1;2\n', 'cos.csv', zapisz=True)
        self.assertFalse(n['rozpoznany'])
        self.assertIn('Nieznany układ', n['ostrzezenie'])
        self.assertEqual(len(dziennik.wczytaj()), 1)

    def test_za_duzy_plik(self):
        r = reczny.przyjmij(b'x' * (reczny.MAKS_ROZMIAR + 1), 'duzy.csv', zapisz=True)
        self.assertFalse(r['rozpoznany'])
        self.assertIn('20 MB', r['ostrzezenie'])

    def test_nazwa_od_przegladarki_tylko_do_opisu(self):
        self.assertEqual(reczny.czysta_nazwa('../../etc/passwd'), 'passwd')
        self.assertEqual(reczny.czysta_nazwa('C:\\Users\\x\\plik\x00.csv'), 'plik.csv')
        self.assertEqual(reczny.czysta_nazwa(''), 'plik bez nazwy')


class TestSerwerEksportImport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.srod = Srodowisko().__enter__()
        dane(cls.srod)
        odbudowa.odbuduj()
        cls.stary_port = konfig.PORT
        cls.http = ThreadingHTTPServer((konfig.HOST, 0), serwer.Obsluga)
        konfig.PORT = cls.http.server_address[1]
        threading.Thread(target=cls.http.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.http.shutdown()
        cls.http.server_close()
        konfig.PORT = cls.stary_port
        cls.srod.__exit__()

    def zadanie(self, metoda, sciezka, cialo=None, naglowki=None):
        c = http.client.HTTPConnection('127.0.0.1', konfig.PORT, timeout=10)
        c.request(metoda, sciezka, body=cialo, headers={'Host': f'127.0.0.1:{konfig.PORT}', **(naglowki or {})})
        r = c.getresponse()
        return r, r.read()

    def test_eksport_jako_pobranie(self):
        r, tresc = self.zadanie('GET', '/api/eksport/budzet/lista?lista=wszystkie&format=csv')
        self.assertEqual(r.status, 200)
        self.assertTrue(r.getheader('Content-Disposition').startswith("attachment; filename*=UTF-8''Finanse"))
        self.assertEqual(r.getheader('Content-Type'), eksport.CSV)
        self.assertTrue(tresc.decode('utf-8').startswith('ID\tRodzaj'))
        r, tresc = self.zadanie('GET', '/api/eksport/operacje?format=xlsx&kierunek=wplywy')
        self.assertEqual((r.status, r.getheader('Content-Type')), (200, eksport.XLSX))
        r, _ = self.zadanie('GET', '/api/eksport/nieznany?format=xlsx')
        self.assertEqual(r.status, 400)

    def test_import_tylko_z_naglowkiem_i_limitem(self):
        tresc = (self.srod.dane / 'Budget Studio' / 'budget_studio_2026-10-13.csv').read_bytes()
        r, _ = self.zadanie('POST', '/api/import?tryb=podglad', tresc)
        self.assertEqual(r.status, 403)
        r, _ = self.zadanie('POST', '/api/import', b'x', {'X-Finanse': '1', 'Content-Length': str(reczny.MAKS_ROZMIAR + 1)})
        self.assertEqual(r.status, 413)

    def test_podglad_zapis_i_historia(self):
        tresc = _ing_bajty(self.srod.katalog, 'Lista.csv', [operacja('2026-10-14', '3,00', '3,00', 'NOWY KLIENT', 'FV7', ref='S1')])
        nagl = {'X-Finanse': '1', 'X-Nazwa-Pliku': 'Lista%20transakcji%20%C5%BC%C3%B3%C5%82ta.csv', 'Content-Type': 'application/octet-stream'}
        r, odp = self.zadanie('POST', '/api/import?tryb=podglad', tresc, nagl)
        odp = json.loads(odp)
        self.assertEqual((r.status, odp['zrodlo'], odp['zapisano'], odp['plik_zrodlowy']), (200, 'ING', False, 'Lista transakcji żółta.csv'))
        self.assertFalse((self.srod.dane / 'ING' / 'ing_historia_2026-10-14_2026-10-14.csv').exists())
        r, odp = self.zadanie('POST', '/api/import?tryb=zapisz', tresc, nagl)
        odp = json.loads(odp)
        self.assertTrue(odp['zapisano'])
        self.assertTrue((self.srod.dane / 'ING' / 'ing_historia_2026-10-14_2026-10-14.csv').exists())
        r, h = self.zadanie('GET', '/api/importy/historia')
        h = json.loads(h)
        self.assertIn(odp['id'], [x['id'] for x in h['importy']])
        r, sz = self.zadanie('GET', '/api/importy/podsumowanie?id=' + odp['id'].replace('/', '%2F'))
        self.assertEqual((r.status, json.loads(sz)['kategoria']), (200, 'dodane'))
        r, odp = self.zadanie('POST', '/api/import?tryb=zapisz', b'%PDF-1.4', nagl)
        self.assertFalse(json.loads(odp)['rozpoznany'])


if __name__ == '__main__':
    unittest.main()
