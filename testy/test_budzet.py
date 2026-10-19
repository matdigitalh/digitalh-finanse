"""Budżet (A–H) na zmyślonych danych: import z eksportu BS, zdarzenia Mata z historią i cofnięciem, plan cykliczny
z częściami, łączenie plan ↔ faktura ↔ bank, etapy z banku, rozjazdy, cash flow, rentowność, eksport, blokada edycji."""
import csv
import io
import json
import unittest

from openpyxl import load_workbook

from finanse import baza, budzet, cashflow, eksport, konfig, odbudowa, rentownosc, zapytania_budzet, zmiany
from testy.pomoc import Srodowisko, faktura, operacja, zapisz_fakturownie, zapisz_ing

KOLUMNY = ['ID', 'Rodzaj', 'Status aktualności', 'Etap płatności', 'Podstawa płatności', 'Numer wewnętrzny', 'Numer dokumentu',
           'Kontrahent', 'Tytuł', 'Netto (wartość)', 'Brutto (wartość)', 'VAT', 'Stawka VAT', 'Miesiąc budżetowy', 'Dzień wystawienia',
           'Dzień płatności', 'Dzień płatności - realny', 'Przesunięcie płatności', 'Miesiąc księgowania', 'Metoda płatności',
           'Księgowane', 'W cash flow', 'Komentarz', 'Komentarz w arkuszu', 'Arkusz', 'Folder', 'Kategoria', 'Sekcja',
           'Przedmiot transakcji', 'Data utworzenia', 'Data modyfikacji']


def bs(id_, rodzaj, etap, podstawa, kontrahent, netto, miesiac, termin, arkusz, folder, kategoria, komentarz='', nr='', status='Aktualne',
       brutto=None, realny=''):
    w = dict.fromkeys(KOLUMNY, '')
    w.update({'ID': id_, 'Rodzaj': rodzaj, 'Status aktualności': status, 'Etap płatności': etap, 'Podstawa płatności': podstawa,
              'Numer dokumentu': nr, 'Kontrahent': kontrahent, 'Netto (wartość)': netto, 'Brutto (wartość)': brutto or netto, 'VAT': '0',
              'Stawka VAT': '0.23', 'Miesiąc budżetowy': f'{miesiac}-01', 'Dzień wystawienia': termin, 'Dzień płatności': termin,
              'Dzień płatności - realny': realny, 'Metoda płatności': 'Przelew', 'Księgowane': 'Tak', 'W cash flow': 'Tak',
              'Komentarz w arkuszu': komentarz, 'Arkusz': arkusz, 'Folder': folder, 'Kategoria': kategoria, 'Sekcja': '-',
              'Przedmiot transakcji': 'Przedmiot transakcji', 'Data utworzenia': '2026-03-01 10:00:00', 'Data modyfikacji': '2026-03-01 10:00:00'})
    return w


def dane(srod):
    zapisz_fakturownie(srod.dane, 'fakturownia_sprzedaz_2026-10-13.csv', [
        faktura(1, 'FV2026/09/1', 'Wystawiona', '2026-09-30', '2026-10-07', 'ALFA SP. Z O.O.', '111', [('Obsługa', 1000)]),
        faktura(2, 'FV2026/08/1', 'Wystawiona', '2026-08-31', '2026-09-07', 'ALFA SP. Z O.O.', '111', [('Obsługa', 1000)]),
    ])
    zapisz_ing(srod.dane, 'ing_historia_2026-01-01_2026-10-13.csv', [
        operacja('2026-10-13', '-10,00', '50000,00', 'ING BANK', 'OPŁATA', ref='X'),
        operacja('2026-09-25', '-600,00', '50010,00', 'URZĄD SKARBOWY', '/OKR/26M08/SFP/VAT-7', ref='V1'),
        operacja('2026-09-25', '-400,00', '50610,00', 'URZĄD SKARBOWY', '/OKR/26M08/SFP/VAT-7 dopłata', ref='V2'),
        operacja('2026-09-10', '-615,00', '51010,00', 'DRUKARNIA SP. Z O.O.', 'FS 1/09', ref='K1'),
        operacja('2026-09-05', '1230,00', '51625,00', 'ALFA SP. Z O.O.', 'FV2026/08/1', ref='W1'),
    ])
    p = srod.dane / 'Budget Studio' / 'budget_studio_2026-10-13.csv'
    p.parent.mkdir(parents=True)
    with open(p, 'w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=KOLUMNY, delimiter='\t')
        w.writeheader()
        w.writerows([
            bs('1', '1', 'Opłacona', 'Faktura ', 'Alfa', '1000,00', '2026-08', '2026-09-07', 'Obsługa stała Alfa', 'Alfa', 'Obsługa stała',
               'Obsługa', nr='FV2026/08/1', brutto='1230,00'),
            bs('2', '1', 'Planowana', 'Faktura ', 'Alfa', '1000,00', '2026-09', '2026-10-07', 'Obsługa stała Alfa', 'Alfa', 'Obsługa stała',
               'Obsługa', brutto='1230,00'),
            bs('3', '1', 'Planowana', 'Faktura ', 'Alfa', '1000,00', '2026-10', '2026-11-07', 'Obsługa stała Alfa', 'Alfa', 'Obsługa stała',
               'Obsługa', brutto='1230,00'),
            bs('4', '1', 'Planowana', 'Faktura ', 'Alfa', '500,00', '2026-07', '2026-07-20', 'Obsługa stała Alfa', 'Alfa', 'Obsługa stała',
               'Dodatki', brutto='615,00'),
            bs('5', '-1', 'Opłacona', 'Faktura ', 'Drukarnia', '-500,00', '2026-09', '2026-09-10', 'Obsługa stała Alfa', 'Alfa', 'Obsługa stała',
               'Druk', brutto='-615,00', realny='2026-09-10'),
            bs('6', '-1', 'Opłacona', 'Podatek VAT', 'Urząd Skarbowy', '-500,00', '2026-08', '2026-09-25', 'Opłaty 2026', 'Zobowiązania fiskalne',
               'Podatki i daniny', realny='2026-09-25'),
            bs('6', '-1', 'Opłacona', 'Podatek VAT', 'Urząd Skarbowy', '-500,00', '2026-08', '2026-09-25', 'Opłaty 2026', 'Zobowiązania fiskalne',
               'Podatki i daniny', realny='2026-09-25'),
        ])
    (srod.katalog / 'Finanse' / 'Aplikacja').mkdir(parents=True, exist_ok=True)
    (srod.katalog / 'Finanse' / 'Aplikacja' / 'arkusze_bs.json').write_text(json.dumps({'arkusze': [
        {'nazwa': 'Obsługa stała Alfa', 'folder': 'Alfa', 'kategoria': 'Obsługa stała', 'od': '2026-07', 'do': '2026-12'},
        {'nazwa': 'Pusty arkusz', 'folder': 'Beta', 'kategoria': 'Projekty dodatkowe', 'od': '2026-10', 'do': '2027-09'}]}), encoding='utf-8')


class TestBudzet(unittest.TestCase):
    def setUp(self):
        self.srod = Srodowisko().__enter__()
        dane(self.srod)
        self.raport = odbudowa.odbuduj()
        self.con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)

    def tearDown(self):
        self.con.close()
        self.srod.__exit__()

    def przelicz(self):
        self.con.close()
        self.raport = odbudowa.odbuduj()
        self.con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)

    def arkusz_id(self, nazwa='Obsługa stała Alfa'):
        return self.con.execute('SELECT id FROM bud_arkusz WHERE nazwa = ?', (nazwa,)).fetchone()[0]

    def etap(self, bs_id, nr=''):
        return self.con.execute('SELECT s.etap FROM bud_platnosc p JOIN bud_stan s ON s.platnosc = p.id WHERE p.bs_id = ? ORDER BY p.id',
                                (bs_id,)).fetchall()

    def test_import_struktura_i_odtwarzalnosc(self):
        self.assertEqual([b for b in self.raport['bledy'] if 'Biuro' not in b], [])
        self.assertEqual(self.raport['budzet'], {'arkuszy': 3, 'wierszy': 4, 'platnosci': 7})   # VAT: to samo ID dwa razy, oba zostają
        self.assertEqual(self.con.execute("SELECT od, do FROM bud_arkusz WHERE nazwa = 'Pusty arkusz'").fetchone()[:], ('2026-10', '2027-09'))
        pierwszy = budzet.plik_importu().read_text(encoding='utf-8')
        self.przelicz()
        self.assertEqual(budzet.plik_importu().read_text(encoding='utf-8'), pierwszy)
        self.assertEqual(json.loads(pierwszy)['zrodlo']['plik'], 'budget_studio_2026-10-13.csv')

    def test_etapy_z_banku_i_po_terminie(self):
        self.assertEqual(self.etap('1')[0][0], 'Opłacona')            # faktura FV2026/08/1 zapłacona w banku
        self.assertEqual(self.etap('5')[0][0], 'Opłacona')            # wydatek 615 z banku
        self.assertEqual(self.etap('4')[0][0], 'Po terminie')         # plan z lipca bez faktury i zapłaty
        self.assertEqual(self.etap('3')[0][0], 'Planowana')
        self.assertEqual({r[0] for r in self.etap('6')}, {'Opłacona'})   # VAT: w BS 2 × 500 (to samo ID), w banku 600 + 400

    def test_plan_powiazany_z_faktura_klienta(self):
        r = self.con.execute("SELECT b.ref, b.pewnosc FROM bud_powiazanie b JOIN bud_platnosc p ON p.id = b.platnosc"
                             " WHERE p.bs_id = '2' AND b.typ = 'faktura'").fetchone()
        self.assertEqual(tuple(r), ('FV2026/09/1', 'prawdopodobne'))   # klient (folder) + kwota netto + miesiąc

    def test_widok_arkusza_i_listy(self):
        d = zapytania_budzet.arkusz(self.con, self.arkusz_id())
        self.assertEqual(d['miesiace'][0], '2026-07')
        self.assertEqual(d['podsumowanie'][-1]['narastajaco'], 100000 + 100000 + 100000 + 50000 - 50000)
        po = zapytania_budzet.lista(self.con, 'po_terminie')
        self.assertEqual(po['razem'], 2)   # plan z lipca bez faktury i plan z września, którego faktura jest po terminie

    def test_edycja_wylaczona_do_wlaczenia(self):
        self.assertFalse(zapytania_budzet.planowanie(self.con))
        zmiany.zapisz('ustawienie', zmiany.sprawdz('ustawienie', {'klucz': 'planowanie', 'wartosc': 'tak'}))
        self.przelicz()
        self.assertTrue(zapytania_budzet.planowanie(self.con))

    def test_plan_cykliczny_z_czesciami_historia_i_cofniecie(self):
        aid = self.arkusz_id()
        wpis = zmiany.zapisz('bud', zmiany.sprawdz('bud', {'op': 'wiersz', 'obiekt': {
            'id': 'w-pensja', 'arkusz': aid, 'rodzaj': -1, 'podstawa': 'Umowa o pracę - dla wykonawcy', 'kontrahent': 'Anna',
            'plan': {'netto': 1000000, 'stawka_vat': '0', 'cykl': 'miesiac', 'od': '2026-11', 'do': '2026-12',
                     'czesci': [{'nazwa': 'netto', 'kwota': 700000, 'dzien_platnosci': 10}, {'nazwa': 'ZUS', 'kwota': 300000, 'dzien_platnosci': 15}]}}}))
        self.przelicz()
        p = self.con.execute("SELECT miesiac, termin, netto FROM bud_platnosc WHERE wiersz = 'w-pensja' ORDER BY miesiac, termin").fetchall()
        self.assertEqual([tuple(x) for x in p], [('2026-11', '2026-12-10', 700000), ('2026-11', '2026-12-15', 300000),
                                                ('2026-12', '2027-01-10', 700000), ('2026-12', '2027-01-15', 300000)])
        self.assertEqual(len(zapytania_budzet.historia(self.con)['zmiany']), 1)
        zmiany.zapisz('cofnij', {'wpis': wpis['id']})
        self.przelicz()
        self.assertIsNone(self.con.execute("SELECT 1 FROM bud_wiersz WHERE id = 'w-pensja'").fetchone())

    def test_decyzja_przesun_i_przeniesienie(self):
        pid = self.con.execute("SELECT id FROM bud_platnosc WHERE bs_id = '4'").fetchone()[0]
        wid = self.con.execute("SELECT wiersz FROM bud_platnosc WHERE bs_id = '4'").fetchone()[0]
        zmiany.zapisz('bud', zmiany.sprawdz('bud', {'op': 'decyzja_planu', 'platnosc': pid, 'decyzja': 'przesun', 'termin': '2026-11-30'}))
        zmiany.zapisz('bud', zmiany.sprawdz('bud', {'op': 'przenies', 'wiersze': [wid], 'do_arkusza': self.arkusz_id('Pusty arkusz')}))
        self.przelicz()
        self.assertEqual(self.etap('4')[0][0], 'Planowana')
        self.assertEqual(self.con.execute('SELECT arkusz FROM bud_wiersz WHERE id = ?', (wid,)).fetchone()[0], self.arkusz_id('Pusty arkusz'))
        h = zapytania_budzet.historia(self.con)['zmiany']
        self.assertEqual({x['op'] for x in h}, {'decyzja_planu', 'przenies'})

    def test_walidacja_zdarzen(self):
        for zle in ({'op': 'usun_wszystko'}, {'op': 'wiersz', 'obiekt': {}}, {'op': 'decyzja_planu', 'platnosc': 'x', 'decyzja': 'przesun'}):
            with self.assertRaises(ValueError):
                zmiany.sprawdz('bud', zle)

    def test_rozjazdy(self):
        rodzaje = {r[0] for r in self.con.execute('SELECT rodzaj FROM bud_rozjazd')}
        self.assertNotIn('Kwota w BS inna niż na fakturze', rodzaje)

    def test_cashflow(self):
        c = cashflow.prognoza(self.con)
        self.assertEqual(c['saldo_start'], 5000000)
        self.assertEqual(len(c['tygodnie']), 13)
        self.assertGreaterEqual(c['koniec']['pelny'], c['koniec']['ostrozny'])   # pełny ma dodatkowo zaległości i plany po terminie
        self.assertEqual(len(c['po_terminie']), 1)

    def test_rentownosc_i_eksport(self):
        zmiany.zapisz('estymata', zmiany.sprawdz('estymata', {'klient': 'Alfa', 'udzial': '25'}))
        self.przelicz()
        self.assertEqual(rentownosc.estymaty(self.con)['Alfa']['udzial'], 25.0)
        for nazwa, tresc in (eksport.arkusz(self.con, self.arkusz_id()), eksport.cash_flow(self.con), eksport.rentownosc_xls(self.con, '2026')):
            self.assertTrue(nazwa.endswith('.xlsx'))
            self.assertTrue(load_workbook(io.BytesIO(tresc)).sheetnames)


class TestBlokadaPrzezSerwer(unittest.TestCase):
    def test_zmiana_budzetu_odrzucona_w_trybie_podgladu(self):
        import http.client
        import threading
        from http.server import ThreadingHTTPServer
        from finanse import serwer
        with Srodowisko() as srod:
            dane(srod)
            odbudowa.odbuduj()
            stary = konfig.PORT
            h = ThreadingHTTPServer((konfig.HOST, 0), serwer.Obsluga)
            konfig.PORT = h.server_address[1]
            threading.Thread(target=h.serve_forever, daemon=True).start()
            try:
                def wyslij(typ, dane_):
                    c = http.client.HTTPConnection('127.0.0.1', konfig.PORT, timeout=10)
                    c.request('POST', '/api/zmiana', body=json.dumps({'typ': typ, 'dane': dane_}),
                              headers={'Host': f'127.0.0.1:{konfig.PORT}', 'X-Finanse': '1', 'Content-Type': 'application/json'})
                    r = c.getresponse()
                    return r.status, json.loads(r.read() or b'{}')
                con = baza.polacz(konfig.plik_bazy(), tylko_odczyt=True)
                aid = con.execute("SELECT id FROM bud_arkusz WHERE nazwa = 'Pusty arkusz'").fetchone()[0]
                con.close()
                wiersz = {'op': 'wiersz', 'obiekt': {'id': 'w1', 'arkusz': aid, 'rodzaj': 1, 'podstawa': 'Faktura (jedna)'}}
                status, odp = wyslij('bud', wiersz)
                self.assertEqual(status, 400)
                self.assertIn('podglądu', odp['blad'])
                self.assertEqual(wyslij('ustawienie', {'klucz': 'planowanie', 'wartosc': 'tak'})[0], 200)
                self.assertEqual(wyslij('bud', {**wiersz, 'obiekt': {**wiersz['obiekt'], 'arkusz': 'nie-ma'}})[0], 400)
                self.assertEqual(wyslij('bud', wiersz)[0], 200)
            finally:
                h.shutdown()
                h.server_close()
                konfig.PORT = stary


if __name__ == '__main__':
    unittest.main()
